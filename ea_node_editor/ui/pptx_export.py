from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from PyQt6.QtGui import QImage

from ea_node_editor.ui.canvas_view_export import safe_filename_component


class CanvasViewPptxExportError(RuntimeError):
    """Raised when a canvas view deck cannot be written."""


class CanvasViewPptxDependencyError(CanvasViewPptxExportError):
    """Raised when python-pptx is not installed."""


@dataclass(frozen=True, slots=True)
class CanvasViewPptxSlide:
    title: str
    image_path: Path


@dataclass(frozen=True, slots=True)
class ProjectReviewPptxSlide:
    title: str
    subtitle: str = ""
    image_path: Path | None = None
    body_lines: tuple[str, ...] = ()
    footer: str = ""


SLIDE_SIZE_PRESETS: dict[str, tuple[float, float]] = {
    "16:9 landscape": (13.333333, 7.5),
    "4:3 landscape": (10.0, 7.5),
    "16:9 portrait": (7.5, 13.333333),
    "4:3 portrait": (7.5, 10.0),
}


def default_canvas_views_deck_path(output_dir: Path | str, workspace_name: object) -> Path:
    stem = safe_filename_component(f"{workspace_name}-canvas-views", fallback="canvas-views")
    return Path(output_dir) / f"{stem}.pptx"


def _pptx_imports() -> tuple[Any, Any, Any, Any]:
    try:
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - exercised when optional extra is absent.
        raise CanvasViewPptxDependencyError(
            "PowerPoint export requires the optional presentation dependency. "
            "Install with the 'presentation' extra."
        ) from exc
    return Presentation, Inches, Pt, PP_ALIGN


def _slide_size_preset(slide_size: str) -> tuple[float, float]:
    preset = SLIDE_SIZE_PRESETS.get(str(slide_size or "").strip().lower())
    if preset is None:
        raise CanvasViewPptxExportError(f"Unknown slide size preset: {slide_size}")
    return preset


def _presentation(*, slide_size: str, template_path: Path | str | None = None):
    Presentation, Inches, _Pt, _PP_ALIGN = _pptx_imports()
    template_text = str(template_path or "").strip()
    if template_text:
        template = Path(template_text).expanduser()
        if not template.exists() or not template.is_file():
            raise CanvasViewPptxExportError(f"PowerPoint template was not found: {template}")
        presentation = Presentation(str(template))
    else:
        presentation = Presentation()
    preset = _slide_size_preset(slide_size)
    presentation.slide_width = Inches(preset[0])
    presentation.slide_height = Inches(preset[1])
    return presentation


def _blank_layout(presentation):  # noqa: ANN001
    try:
        return presentation.slide_layouts[6]
    except IndexError:
        return presentation.slide_layouts[0]


def _image_size(path: Path) -> tuple[int, int]:
    image = QImage(str(path))
    if image.isNull() or image.width() <= 0 or image.height() <= 0:
        raise CanvasViewPptxExportError(f"Could not read exported PNG for slide: {path}")
    return int(image.width()), int(image.height())


def _add_fitted_picture(
    slide,  # noqa: ANN001
    presentation,  # noqa: ANN001
    image_path: Path,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    shape_name: str,
):
    image_width, image_height = _image_size(image_path)
    width_ratio = float(width) / float(image_width)
    height_ratio = float(height) / float(image_height)
    ratio = min(width_ratio, height_ratio)
    picture_width = int(round(image_width * ratio))
    picture_height = int(round(image_height * ratio))
    picture_left = int(round(left + (width - picture_width) / 2))
    picture_top = int(round(top + (height - picture_height) / 2))
    picture = slide.shapes.add_picture(
        str(image_path),
        picture_left,
        picture_top,
        width=picture_width,
        height=picture_height,
    )
    picture.name = shape_name
    return picture


def _add_textbox(
    slide,  # noqa: ANN001
    *,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    font_size: int,
    bold: bool = False,
    align: Any | None = None,
) -> None:
    _Presentation, _Inches, Pt, PP_ALIGN = _pptx_imports()
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = str(text or "")
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bool(bold)
    paragraph.alignment = align if align is not None else PP_ALIGN.LEFT


def _add_body_lines(
    slide,  # noqa: ANN001
    *,
    lines: tuple[str, ...],
    left: int,
    top: int,
    width: int,
    height: int,
    font_size: int,
) -> None:
    _Presentation, _Inches, Pt, PP_ALIGN = _pptx_imports()
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.clear()
    for index, line in enumerate(lines or ("",)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(line or "")
        paragraph.font.size = Pt(font_size)
        paragraph.alignment = PP_ALIGN.LEFT


def create_canvas_views_pptx(
    *,
    slides: list[CanvasViewPptxSlide],
    output_path: Path | str,
    slide_size: str = "16:9 landscape",
) -> Path:
    if not slides:
        raise CanvasViewPptxExportError("At least one exported PNG is required for PowerPoint export.")

    presentation = _presentation(slide_size=slide_size)
    blank_layout = presentation.slide_layouts[6]
    for slide_spec in slides:
        image_path = Path(slide_spec.image_path)
        slide = presentation.slides.add_slide(blank_layout)
        shape_name = safe_filename_component(slide_spec.title, fallback="Canvas View", max_length=80)
        _add_fitted_picture(
            slide,
            presentation,
            image_path,
            left=0,
            top=0,
            width=int(presentation.slide_width),
            height=int(presentation.slide_height),
            shape_name=f"Canvas View - {shape_name}",
        )

    deck_path = Path(output_path)
    deck_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(deck_path))
    return deck_path


def create_project_review_pptx(
    *,
    slides: list[ProjectReviewPptxSlide],
    output_path: Path | str,
    slide_size: str = "16:9 landscape",
    template_path: Path | str | None = None,
) -> Path:
    if not slides:
        raise CanvasViewPptxExportError("At least one review slide is required for PowerPoint export.")

    _Presentation, Inches, _Pt, PP_ALIGN = _pptx_imports()
    presentation = _presentation(slide_size=slide_size, template_path=template_path)
    blank_layout = _blank_layout(presentation)
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    margin = int(Inches(0.45))
    title_height = int(Inches(0.5))
    subtitle_height = int(Inches(0.32))
    footer_height = int(Inches(0.25))

    for slide_spec in slides:
        slide = presentation.slides.add_slide(blank_layout)
        title_text = str(slide_spec.title or "Project Review")
        _add_textbox(
            slide,
            text=title_text,
            left=margin,
            top=int(Inches(0.28)),
            width=slide_width - (2 * margin),
            height=title_height,
            font_size=24,
            bold=True,
            align=PP_ALIGN.LEFT,
        )
        subtitle_text = str(slide_spec.subtitle or "")
        if subtitle_text:
            _add_textbox(
                slide,
                text=subtitle_text,
                left=margin,
                top=int(Inches(0.82)),
                width=slide_width - (2 * margin),
                height=subtitle_height,
                font_size=12,
            )

        content_top = int(Inches(1.25 if subtitle_text else 1.0))
        content_bottom = slide_height - margin - (footer_height if slide_spec.footer else 0)
        content_height = max(1, content_bottom - content_top)
        body_lines = tuple(str(line) for line in slide_spec.body_lines if str(line or "").strip())
        image_path = Path(slide_spec.image_path) if slide_spec.image_path is not None else None
        if image_path is not None:
            shape_name = safe_filename_component(title_text, fallback="Review Evidence", max_length=80)
            _add_fitted_picture(
                slide,
                presentation,
                image_path,
                left=margin,
                top=content_top,
                width=slide_width - (2 * margin),
                height=content_height,
                shape_name=f"Project Review - {shape_name}",
            )
        elif body_lines:
            _add_body_lines(
                slide,
                lines=body_lines,
                left=margin,
                top=content_top,
                width=slide_width - (2 * margin),
                height=content_height,
                font_size=13,
            )

        if slide_spec.footer:
            _add_textbox(
                slide,
                text=str(slide_spec.footer),
                left=margin,
                top=slide_height - margin - footer_height,
                width=slide_width - (2 * margin),
                height=footer_height,
                font_size=9,
            )

    deck_path = Path(output_path)
    deck_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(deck_path))
    return deck_path


__all__ = [
    "CanvasViewPptxDependencyError",
    "CanvasViewPptxExportError",
    "CanvasViewPptxSlide",
    "ProjectReviewPptxSlide",
    "SLIDE_SIZE_PRESETS",
    "create_canvas_views_pptx",
    "create_project_review_pptx",
    "default_canvas_views_deck_path",
]
