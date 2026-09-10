from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

import pandas as pd


@dataclass(frozen=True)
class AnalysisTable:
    table_id: str
    title: str
    frame: pd.DataFrame


@dataclass(frozen=True)
class AnalysisTableBundle:
    scope: str
    summary: Mapping[str, Any]
    tables: tuple[AnalysisTable, ...]

    def table(self, table_id: str) -> AnalysisTable | None:
        for table in self.tables:
            if table.table_id == table_id:
                return table
        return None


def records_frame(records: Sequence[Mapping[str, Any]]) -> pd.DataFrame:
    return pd.DataFrame([dict(record) for record in records])


CERTIFICATION_RANKING_COLUMNS = (
    "Channel",
    "Evidence Grade",
    "Certification Eligible",
    "Certification Score",
    "Exclusion Reason",
    "Validation Role",
    "Structural Relevance",
    "Peak Strain Test",
    "Peak Strain FEA",
    "Peak Error (%)",
    "Slope Error (%)",
    "Offset (signal units)",
    "Envelope NMAE (%)",
    "Pearson Correlation",
    "Sign Agreement (%)",
    "Data Quality Warnings",
)


def certification_ranking_frame(metrics: Sequence[Mapping[str, Any]]) -> pd.DataFrame:
    frame = records_frame(metrics)
    if frame.empty:
        return pd.DataFrame(columns=list(CERTIFICATION_RANKING_COLUMNS))
    columns = [column for column in CERTIFICATION_RANKING_COLUMNS if column in frame.columns]
    ranking = frame.loc[:, columns].copy()
    grade_order = {"A": 0, "B": 1, "C": 2, "Reject": 3}
    grade_source = ranking["Evidence Grade"] if "Evidence Grade" in ranking else pd.Series([""] * len(ranking))
    score_source = ranking["Certification Score"] if "Certification Score" in ranking else pd.Series([-1.0] * len(ranking))
    ranking["_grade_order"] = grade_source.map(grade_order).fillna(4)
    ranking["_score_order"] = pd.to_numeric(score_source, errors="coerce").fillna(-1.0)
    return (
        ranking.sort_values(["_grade_order", "_score_order", "Channel"], ascending=[True, False, True])
        .drop(columns=["_grade_order", "_score_order"])
        .reset_index(drop=True)
    )


def build_analysis_table_bundle(
    *,
    scope: str,
    summary: Mapping[str, Any],
    metrics: Sequence[Mapping[str, Any]],
    scale_offset_metrics: Sequence[Mapping[str, Any]],
    sign_metrics: Sequence[Mapping[str, Any]],
    sign_status_frame: pd.DataFrame,
    residual_metrics: Sequence[Mapping[str, Any]],
    residual_frame: pd.DataFrame,
    rolling_frame: pd.DataFrame,
    lag_metrics: Sequence[Mapping[str, Any]],
    lag_correlations: pd.DataFrame,
    event_metrics: Sequence[Mapping[str, Any]],
    quality_rows: Sequence[Mapping[str, Any]],
    calibration_metrics: Sequence[Mapping[str, Any]],
    frequency_metrics: Sequence[Mapping[str, Any]],
    frequency_spectrum: pd.DataFrame,
    overlay_frames: Mapping[str, pd.DataFrame],
) -> AnalysisTableBundle:
    tables: list[AnalysisTable] = [
        AnalysisTable("metrics", "Statistical Metrics", records_frame(metrics)),
        AnalysisTable("certification_ranking", "Certification Ranking", certification_ranking_frame(metrics)),
        AnalysisTable("scale_offset", "Scale and Offset", records_frame(scale_offset_metrics)),
        AnalysisTable("calibration", "Calibration Fit Metrics", records_frame(calibration_metrics)),
        AnalysisTable("residual_summary", "Residual Summary", records_frame(residual_metrics)),
        AnalysisTable("residual_time_series", "Residual Time Series", residual_frame.copy()),
        AnalysisTable("rolling_metrics", "Rolling Metrics", rolling_frame.copy()),
        AnalysisTable("lag_metrics", "Lag Metrics", records_frame(lag_metrics)),
        AnalysisTable("lag_correlations", "Lag Correlations", lag_correlations.copy()),
        AnalysisTable("sign_agreement_summary", "Sign Agreement Summary", records_frame(sign_metrics)),
        AnalysisTable("sign_agreement_status", "Sign Agreement Status", sign_status_frame.copy()),
        AnalysisTable("events", "Events", records_frame(event_metrics)),
        AnalysisTable("data_quality", "Data Quality", records_frame(quality_rows)),
        AnalysisTable("frequency_metrics", "Frequency Metrics", records_frame(frequency_metrics)),
        AnalysisTable("frequency_spectrum", "Frequency Spectrum", frequency_spectrum.copy()),
    ]
    for table_id, title in [
        ("overlay_reference", "Overlay Reference"),
        ("overlay_candidate_original", "Overlay Candidate Original"),
        ("overlay_candidate_scaled", "Overlay Candidate Scaled"),
        ("overlay_candidate_offset", "Overlay Candidate Offset"),
        ("overlay_candidate_scaled_offset", "Overlay Candidate Scaled + Offset"),
    ]:
        frame = overlay_frames.get(table_id)
        if frame is not None:
            tables.append(AnalysisTable(table_id, title, frame.copy()))
    return AnalysisTableBundle(scope=scope, summary=dict(summary), tables=tuple(tables))


def single_table_bundle(table: AnalysisTable, *, summary: Mapping[str, Any], scope: str = "Current Table") -> AnalysisTableBundle:
    return AnalysisTableBundle(scope=scope, summary=dict(summary), tables=(table,))


def write_analysis_workbook(path: str | Path, bundle: AnalysisTableBundle) -> Path:
    try:
        import openpyxl
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("Excel export requires openpyxl. Install with: pip install openpyxl") from exc

    output_path = Path(path)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    try:
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        _write_sheet(summary_sheet, _summary_frame(bundle))
        used_names = {"Summary"}
        for table in bundle.tables:
            sheet_name = _unique_sheet_name(table.title, used_names)
            used_names.add(sheet_name)
            sheet = workbook.create_sheet(sheet_name)
            _write_sheet(sheet, table.frame)
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path


def _summary_frame(bundle: AnalysisTableBundle) -> pd.DataFrame:
    rows = [{"Property": "Export Scope", "Value": bundle.scope}]
    rows.extend({"Property": str(key), "Value": value} for key, value in bundle.summary.items())
    rows.append({"Property": "Table Count", "Value": len(bundle.tables)})
    return pd.DataFrame(rows)


def _write_sheet(sheet: Any, frame: pd.DataFrame) -> None:
    from openpyxl.styles import Font

    columns = [str(column) for column in frame.columns]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in frame.itertuples(index=False, name=None):
        sheet.append([_excel_value(value) for value in row])
    sheet.freeze_panes = "A2"
    if columns:
        sheet.auto_filter.ref = sheet.dimensions
    _fit_column_widths(sheet, columns, frame)


def _excel_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        try:
            value = value.item()
        except Exception:  # noqa: BLE001
            pass
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def _fit_column_widths(sheet: Any, columns: Sequence[str], frame: pd.DataFrame) -> None:
    sample = frame.head(200)
    for index, column in enumerate(columns, start=1):
        values = [column]
        if column in sample:
            values.extend("" if pd.isna(value) else str(value) for value in sample[column].tolist())
        width = min(60, max(10, max(len(value) for value in values) + 2))
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width


def _unique_sheet_name(title: str, used_names: set[str]) -> str:
    base = _safe_sheet_name(title)
    if base not in used_names:
        return base
    suffix = 2
    while True:
        candidate = f"{base[: 31 - len(str(suffix)) - 1]}_{suffix}"
        if candidate not in used_names:
            return candidate
        suffix += 1


def _safe_sheet_name(title: str) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", " ", title).strip()
    cleaned = re.sub(r"\s+", " ", cleaned) or "Table"
    return cleaned[:31]
