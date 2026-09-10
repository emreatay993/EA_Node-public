# Purpose: Section-geometry, visualization payload, overlay, history, and nodal export helpers for the MCF DPF section resultants tool.
# Map: subsystems/packaging_generated_assets
# Tests: tests/test_mcf_dpf_section_resultants_gui.py
# Landmarks: load_static_animation_session; static_animation_frame_state; build_static_animation_render_frame; build_section_visualization_payload
"""Visualization, geometry, and export payload helpers."""

from __future__ import annotations

import argparse
import csv
import errno
import json
import math
import os
import re
import sys
import tempfile
import traceback
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from time import perf_counter
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from scripts.mcf_dpf_section_resultants.core import *
from scripts.mcf_dpf_section_resultants.core import _VISUALIZATION_MESH_SNAPSHOT_CACHE
from scripts.mcf_dpf_section_resultants.dpf_io import *

def clear_visualization_mesh_snapshot_cache() -> None:
    _VISUALIZATION_MESH_SNAPSHOT_CACHE.clear()


FOLLOW_GEOMETRY_REFERENCE_PREVIEW_REASON = "follow_geometry_tracking_unavailable"


def follow_geometry_reference_preview_warning(tracking_node_count: int) -> str:
    return (
        "Follow Geometry cannot track this section because the initial plane provides "
        f"fewer than 3 usable tracking nodes ({int(tracking_node_count)} found). "
        "Showing the initial/reference state only; other result sets cannot be "
        "previewed until the plane cuts the selected named selection or an explicit "
        "frame attachment is selected."
    )


def visualization_mesh_snapshot_cache_key(
    signature: ModalRstSignature,
    element_named_selection: str,
    external_named_selection_path: str = "",
    analysis_mode: str = "modal",
    result_set_id: Optional[int] = None,
    reference_frame_motion: str = "follow-geometry",
    reference_frame_attachment_selection: str = "",
) -> VisualizationMeshSnapshotCacheKey:
    external_path, external_size, external_mtime = optional_file_signature(
        external_named_selection_path
    )
    return (
        signature.path,
        signature.size_bytes,
        signature.mtime_ns,
        str(element_named_selection or "").strip().lower(),
        external_path,
        external_size,
        external_mtime,
        str(analysis_mode or "modal").strip().lower(),
        int(result_set_id) if result_set_id is not None else 0,
        str(reference_frame_motion or "follow-geometry").strip().lower(),
        str(reference_frame_attachment_selection or "").strip().lower(),
    )


def static_animation_topology_from_mesh(
    config: SectionConfig,
    *,
    mesh: Any,
    element_scoping: Any,
    element_name: str,
    available_named: Sequence[str],
    displacement_node_ids: Sequence[int],
    tracking_node_ids: Sequence[int],
    tracking_reference_coordinates: Dict[int, Vector],
    attachment: Dict[str, Any],
    origin_unit_conversion: Optional[Dict[str, Any]] = None,
) -> StaticAnimationTopology:
    """Capture reference topology once; no per-set coordinates live here."""

    import numpy as np

    node_ids = tuple(int(value) for value in selected_mesh_node_ids(mesh))
    displacement_ids = tuple(sorted(set(int(value) for value in displacement_node_ids)))
    displacement_index = {node_id: index for index, node_id in enumerate(displacement_ids)}
    missing_mesh_ids = [node_id for node_id in node_ids if node_id not in displacement_index]
    if missing_mesh_ids:
        raise RuntimeError(
            "Animation displacement scope does not include selected-mesh node IDs: "
            f"{missing_mesh_ids[:10]}."
        )
    reference_points = np.asarray(
        [node_xyz(mesh, node_id) for node_id in node_ids], dtype=np.float64
    )
    displacement_reference_points = np.asarray(
        [
            (
                tracking_reference_coordinates[node_id]
                if node_id in tracking_reference_coordinates
                else node_xyz(mesh, node_id)
            )
            for node_id in displacement_ids
        ],
        dtype=np.float64,
    )
    scoped_element_ids = tuple(int(value) for value in object_ids(element_scoping))
    grid = mesh_grid_payload(mesh)
    grid_element_ids = tuple(int(value) for value in grid.get("element_ids", []))
    if grid_element_ids:
        if set(grid_element_ids) != set(scoped_element_ids):
            raise RuntimeError(
                "Animation grid cells do not represent the selected element scope."
            )
        element_ids = grid_element_ids
    else:
        element_ids = scoped_element_ids
    if len(grid.get("celltypes") or []) != len(element_ids):
        raise RuntimeError(
            "Animation requires one renderable grid cell for each selected element."
        )
    point_index = {node_id: index for index, node_id in enumerate(node_ids)}
    element_nodes: Dict[int, Tuple[int, ...]] = {}
    connectivity: List[int] = []
    offsets: List[int] = []
    for element_id in element_ids:
        offsets.append(len(connectivity))
        nodes = tuple(int(value) for value in element_node_ids(mesh, element_id))
        element_nodes[element_id] = nodes
        try:
            connectivity.extend(point_index[node_id] for node_id in nodes)
        except KeyError as exc:
            raise RuntimeError(
                f"Animation topology element {element_id} references node {exc.args[0]} "
                "outside the selected mesh."
            ) from exc
    tracking_ids = tuple(int(value) for value in tracking_node_ids)
    missing_tracking_ids = [
        node_id for node_id in tracking_ids if node_id not in displacement_index
    ]
    if missing_tracking_ids:
        raise RuntimeError(
            "Animation displacement scope does not include tracking node IDs: "
            f"{missing_tracking_ids[:10]}."
        )
    arrays = (
        reference_points,
        displacement_reference_points,
        np.asarray(grid.get("cells") or [], dtype=np.int64),
        np.asarray(grid.get("celltypes") or [], dtype=np.uint8),
        np.asarray(connectivity, dtype=np.int64),
        np.asarray(offsets, dtype=np.int64),
        np.asarray([displacement_index[node_id] for node_id in node_ids], dtype=np.int64),
        np.asarray([displacement_index[node_id] for node_id in tracking_ids], dtype=np.int64),
    )
    for array in arrays:
        array.setflags(write=False)
    return StaticAnimationTopology(
        node_ids=node_ids,
        displacement_node_ids=displacement_ids,
        reference_points=reference_points,
        displacement_reference_points=displacement_reference_points,
        cells=arrays[2],
        celltypes=arrays[3],
        element_ids=element_ids,
        element_node_ids=element_nodes,
        connectivity_point_indices=arrays[4],
        element_offsets=arrays[5],
        mesh_displacement_indices=arrays[6],
        tracking_node_ids=tracking_ids,
        tracking_displacement_indices=arrays[7],
        element_name=str(element_name),
        available_named_selections=tuple(str(value) for value in available_named),
        mesh_unit=str(getattr(mesh, "unit", "") or "") or None,
        mesh_config=asdict(config),
        origin_unit_conversion=(
            dict(origin_unit_conversion) if origin_unit_conversion is not None else None
        ),
        attachment=dict(attachment),
    )


def create_static_animation_session(
    config: SectionConfig,
    selected_result_sets: Sequence[Dict[str, Any]],
    *,
    mesh: Any,
    element_scoping: Any,
    element_name: str,
    available_named: Sequence[str],
    displacement_node_ids: Sequence[int],
    tracking_node_ids: Sequence[int],
    tracking_reference_coordinates: Dict[int, Vector],
    attachment: Dict[str, Any],
    origin_unit_conversion: Optional[Dict[str, Any]] = None,
    signature_config: Optional[SectionConfig] = None,
    generation: int = 0,
) -> StaticAnimationSession:
    topology = static_animation_topology_from_mesh(
        config,
        mesh=mesh,
        element_scoping=element_scoping,
        element_name=element_name,
        available_named=available_named,
        displacement_node_ids=displacement_node_ids,
        tracking_node_ids=tracking_node_ids,
        tracking_reference_coordinates=tracking_reference_coordinates,
        attachment=attachment,
        origin_unit_conversion=origin_unit_conversion,
    )
    return StaticAnimationSession(
        static_animation_session_signature(signature_config or config, selected_result_sets),
        topology,
        selected_result_sets,
        generation=generation,
    )


def _animation_reference_frame_from_displacements(
    session: StaticAnimationSession,
    frame_index: int,
    displacement_values: Any,
    *,
    deformation_scale: float = 1.0,
    result_value: Optional[float] = None,
) -> ResolvedReferenceFrame:
    import numpy as np

    topology = session.topology
    config = config_from_mapping(topology.mesh_config)
    record = session.records[int(frame_index)]
    value = record.result_value if result_value is None else float(result_value)
    if (
        float(deformation_scale) == 1.0
        and value == record.result_value
        and record.resolved_reference_frame is not None
    ):
        return ResolvedReferenceFrame(**record.resolved_reference_frame)
    if config.reference_frame_motion != "follow-geometry":
        return fixed_reference_frame(
            config.coordinate_system_origin,
            config.coordinate_system_axes,
            result_set_id=record.result_set_id,
            result_value=value,
            result_unit=record.result_unit,
        )
    tracking_indices = topology.tracking_displacement_indices
    if len(tracking_indices) < 3:
        raise ValueError("Follow-geometry animation requires at least 3 tracking nodes.")
    reference = topology.displacement_reference_points[tracking_indices]
    current = reference + float(deformation_scale) * np.asarray(
        displacement_values, dtype=np.float64
    )[tracking_indices]
    attachment = topology.attachment
    return fit_geometry_following_frame(
        reference,
        current,
        config.coordinate_system_origin,
        config.coordinate_system_axes,
        result_set_id=record.result_set_id,
        result_value=value,
        result_unit=record.result_unit,
        attachment_source=str(attachment.get("source") or "local_section_cut_neighborhood"),
        attachment_selection=str(attachment.get("selection") or ""),
        warning_ratio=config.reference_frame_fit_warning_ratio,
        tracking_node_ids=topology.tracking_node_ids,
    )


def load_static_animation_session(
    config: SectionConfig,
    selected_result_sets: Sequence[Dict[str, Any]],
    *,
    generation: int = 0,
    expected_signature: Optional[StaticAnimationSessionSignature] = None,
    priority_index: int = 0,
    direction: int = 1,
    cancel_requested: Optional[Callable[[], bool]] = None,
    session_is_current: Optional[
        Callable[[int, StaticAnimationSessionSignature], bool]
    ] = None,
    on_batch: Optional[Callable[[StaticAnimationSession, List[int]], None]] = None,
    log: LogFn = None,
) -> StaticAnimationSession:
    """Stream preview displacements through one DPF model/streams lifetime."""

    import ansys.dpf.core as dpf

    source_config = config_from_mapping(asdict(config))
    if source_config.analysis_mode != "static":
        raise ValueError("Animation streaming is available only for Static Structural results.")
    if len(selected_result_sets) < 2:
        raise ValueError("Static animation requires at least two selected result sets.")
    signature = static_animation_session_signature(source_config, selected_result_sets)
    if expected_signature is not None and signature != expected_signature:
        raise ValueError("Animation request signature no longer matches the current inputs.")
    preflight_dpf_open(dpf, source_config.modal_rst, log=log)
    data_sources = dpf.DataSources(source_config.modal_rst)
    streams_container = create_streams_container(dpf, data_sources)
    session: Optional[StaticAnimationSession] = None
    try:
        model = dpf.Model(data_sources)
        if source_config.external_named_selection_path:
            element_scoping, element_name, available_named = resolve_external_element_scoping(
                dpf,
                model,
                source_config.external_named_selection_path,
                source_config.element_named_selection,
            )
        else:
            element_scoping, element_name, available_named = resolve_element_named_selection_scoping(
                dpf,
                model,
                data_sources,
                streams_container,
                source_config.element_named_selection,
            )
        reference_mesh = selected_element_mesh(model, element_scoping)
        mesh_config, origin_unit_conversion = config_with_origin_in_mesh_units(
            source_config, reference_mesh, log=log
        )
        mesh_unit = str(getattr(reference_mesh, "unit", "") or "").strip()
        if not mesh_unit:
            raise RuntimeError("Static animation requires an RST mesh length unit.")
        initial_axes = validate_axes(mesh_config.coordinate_system_axes)
        _element_scope, _node_scope, initial_side_info = construction_surface_scoping(
            dpf,
            model,
            element_scoping,
            mesh_config,
            initial_axes,
            mesh=reference_mesh,
            allow_empty_cut=True,
        )
        section_node_ids = selected_mesh_node_ids(reference_mesh)
        tracking_node_ids: List[int] = []
        tracking_reference_coordinates: Dict[int, Vector] = {}
        if mesh_config.reference_frame_motion == "follow-geometry":
            override_name = mesh_config.reference_frame_attachment_selection.strip()
            if override_name:
                tracking_node_ids, attachment = resolve_reference_frame_attachment_nodes(
                    dpf,
                    model,
                    name=override_name,
                    external_named_selection_path=mesh_config.external_named_selection_path,
                )
                full_mesh = getattr(model.metadata, "meshed_region", None)
                full_mesh = full_mesh() if callable(full_mesh) else full_mesh
                tracking_reference_coordinates = mesh_node_coordinates_for_ids(
                    full_mesh, tracking_node_ids
                )
            else:
                tracking_node_ids = unique_element_node_ids(
                    reference_mesh, initial_side_info.get("cut_element_ids", [])
                )
                tracking_reference_coordinates = mesh_node_coordinates_for_ids(
                    reference_mesh, tracking_node_ids
                )
                attachment = {
                    "source": "local_section_cut_neighborhood",
                    "selection": element_name,
                    "entity": "INITIAL_CUT_ELEMENTS",
                    "entity_count": len(initial_side_info.get("cut_element_ids", [])),
                    "node_count": len(tracking_node_ids),
                    "node_ids_hash": stable_id_hash(tracking_node_ids),
                }
        else:
            attachment = {
                "source": "none",
                "selection": "",
                "entity": "NONE",
                "entity_count": 0,
                "node_count": 0,
                "node_ids_hash": stable_id_hash([]),
            }
        displacement_node_ids = sorted(set(section_node_ids).union(tracking_node_ids))
        session = create_static_animation_session(
            mesh_config,
            selected_result_sets,
            mesh=reference_mesh,
            element_scoping=element_scoping,
            element_name=element_name,
            available_named=available_named,
            displacement_node_ids=displacement_node_ids,
            tracking_node_ids=tracking_node_ids,
            tracking_reference_coordinates=tracking_reference_coordinates,
            attachment=attachment,
            origin_unit_conversion=origin_unit_conversion,
            signature_config=source_config,
            generation=generation,
        )
        order = static_animation_load_order(
            len(selected_result_sets),
            priority_index=priority_index,
            direction=direction,
        )
        batch_size = static_animation_batch_size(len(displacement_node_ids))
        emit(
            log,
            "Streaming static animation displacement frames in batches of at most "
            f"{batch_size} set(s); backing store={session.storage_kind}, "
            f"{format_bytes(session.store_bytes)}.",
        )
        for batch_start in range(0, len(order), batch_size):
            if cancel_requested is not None and cancel_requested():
                session.cancelled = True
                session.warnings.append("Animation loading was cancelled after the current batch.")
                break
            if session_is_current is not None and not session_is_current(
                generation, signature
            ):
                session.cancelled = True
                session.warnings.append("Animation loading result was rejected as stale.")
                break
            batch_indices = order[batch_start : batch_start + batch_size]
            batch_displacements = nodal_displacement_batch_for_ids(
                dpf,
                data_sources,
                streams_container,
                displacement_node_ids,
                [session.records[index].result_set_id for index in batch_indices],
                mesh_unit,
                log=log,
            )
            for index in batch_indices:
                record = session.records[index]
                displacements, deformation = batch_displacements[record.result_set_id]
                displacement_values = [
                    displacements[node_id] for node_id in displacement_node_ids
                ]
                frame = _animation_reference_frame_from_displacements(
                    session,
                    index,
                    displacement_values,
                )
                session.put_displacements(
                    index,
                    displacement_values,
                    deformation=deformation,
                    resolved_reference_frame=asdict(frame),
                )
                if on_batch is not None:
                    on_batch(session, [index])
            if session_is_current is not None and not session_is_current(
                generation, signature
            ):
                session.cancelled = True
                session.warnings.append("Animation loading result was rejected as stale.")
                break
        session.mark_complete()
        return session
    except Exception:
        if session is not None:
            session.close()
        raise
    finally:
        try:
            streams_container.release_handles()
            emit(log, "Released DPF animation stream handles.")
        except Exception:
            pass

def copy_grid_payload(payload: Dict[str, Any]) -> Dict[str, Any]:
    copied = {
        "points": [
            [float(value) for value in point[:3]]
            for point in (payload.get("points") or [])
        ],
        "cells": [int(value) for value in (payload.get("cells") or [])],
        "celltypes": [int(value) for value in (payload.get("celltypes") or [])],
    }
    if "warning" in payload:
        copied["warning"] = payload.get("warning")
    if "element_ids" in payload:
        copied["element_ids"] = [int(value) for value in (payload.get("element_ids") or [])]
    return copied


class _SnapshotNode:
    def __init__(self, coordinates: Sequence[float]) -> None:
        self.coordinates = [float(value) for value in coordinates[:3]]


class _SnapshotNodes:
    def __init__(self, coordinates: Dict[int, Vector], node_ids: Sequence[int]) -> None:
        self._coordinates = coordinates
        self.scoping = type(
            "_SnapshotScoping",
            (),
            {"ids": [int(value) for value in node_ids]},
        )()

    def node_by_id(self, node_id: int) -> _SnapshotNode:
        return _SnapshotNode(self._coordinates[int(node_id)])


class _SnapshotElement:
    def __init__(self, node_ids: Sequence[int]) -> None:
        self.node_ids = [int(value) for value in node_ids]


class _SnapshotElements:
    def __init__(self, element_node_ids: Dict[int, List[int]], element_ids: Sequence[int]) -> None:
        self._element_node_ids = element_node_ids
        self.scoping = type("_SnapshotScoping", (), {"ids": [int(value) for value in element_ids]})()

    def element_by_id(self, element_id: int) -> _SnapshotElement:
        return _SnapshotElement(self._element_node_ids[int(element_id)])


class _SnapshotGrid:
    def __init__(self, payload: Dict[str, Any]) -> None:
        self.points = copy_grid_payload(payload).get("points", [])
        self.cells = [int(value) for value in (payload.get("cells") or [])]
        self.celltypes = [int(value) for value in (payload.get("celltypes") or [])]


class _SnapshotMesh:
    def __init__(self, snapshot: VisualizationMeshSnapshot) -> None:
        self.unit = snapshot.mesh_unit
        self._grid_payload = copy_grid_payload(snapshot.grid_payload)
        self.nodes = _SnapshotNodes(snapshot.node_coordinates, snapshot.node_ids)
        self.elements = _SnapshotElements(snapshot.element_node_ids, snapshot.element_ids)
        self.grid = _SnapshotGrid(snapshot.grid_payload)
        self.deformation = dict(snapshot.deformation or {})


class _CoordinateOverrideMesh:
    def __init__(
        self,
        mesh: Any,
        node_ids: Sequence[int],
        node_coordinates: Dict[int, Vector],
        grid_payload: Dict[str, Any],
        deformation: Dict[str, Any],
    ) -> None:
        self.unit = getattr(mesh, "unit", None)
        self.nodes = _SnapshotNodes(node_coordinates, node_ids)
        self.elements = mesh.elements
        self._grid_payload = copy_grid_payload(grid_payload)
        self.grid = _SnapshotGrid(grid_payload)
        self.deformation = dict(deformation)


def mesh_with_nodal_displacements(
    mesh: Any,
    nodal_displacements: Dict[int, Sequence[float]],
    deformation: Dict[str, Any],
    *,
    include_grid: bool = True,
) -> Any:
    node_ids = selected_mesh_node_ids(mesh)
    missing_node_ids = [node_id for node_id in node_ids if node_id not in nodal_displacements]
    if missing_node_ids:
        raise RuntimeError(
            "Cannot construct deformed mesh coordinates because displacement is missing "
            f"for {len(missing_node_ids)} node(s); first missing IDs: {missing_node_ids[:10]}."
        )
    deformed_coordinates = {
        node_id: add_vectors(node_xyz(mesh, node_id), nodal_displacements[node_id])
        for node_id in node_ids
    }
    grid_payload = (
        mesh_grid_payload(mesh)
        if include_grid
        else {
            "points": [],
            "cells": [],
            "celltypes": [],
            "element_ids": mesh_element_ids(mesh),
        }
    )
    grid_points = grid_payload.get("points") or []
    if grid_points:
        if len(grid_points) != len(node_ids):
            raise RuntimeError(
                "Selected mesh grid-point order cannot be aligned to DPF node IDs: "
                f"points={len(grid_points)}, node_ids={len(node_ids)}."
            )
        grid_payload["points"] = [
            [float(value) for value in deformed_coordinates[node_id]]
            for node_id in node_ids
        ]
    return _CoordinateOverrideMesh(
        mesh,
        node_ids,
        deformed_coordinates,
        grid_payload,
        deformation,
    )


def visualization_mesh_snapshot_from_mesh(
    key: VisualizationMeshSnapshotCacheKey,
    *,
    mesh: Any,
    element_name: str,
    available_named: Sequence[str],
    element_scoping: Any,
    nodal_displacements: Optional[Dict[int, Vector]] = None,
    deformation: Optional[Dict[str, Any]] = None,
    resolved_reference_frame: Optional[Dict[str, Any]] = None,
) -> VisualizationMeshSnapshot:
    element_ids = object_ids(element_scoping)
    mesh_node_ids = selected_mesh_node_ids(mesh)
    element_nodes: Dict[int, List[int]] = {}
    node_ids: set[int] = set()
    for element_id in element_ids:
        nodes = element_node_ids(mesh, int(element_id))
        element_nodes[int(element_id)] = nodes
        node_ids.update(nodes)
    missing_topology_node_ids = sorted(node_ids.difference(mesh_node_ids))
    if missing_topology_node_ids:
        raise RuntimeError(
            "Selected mesh topology references node IDs outside its nodal scoping: "
            f"{missing_topology_node_ids[:10]}."
        )
    return VisualizationMeshSnapshot(
        key=key,
        element_name=str(element_name),
        available_named_selections=[str(value) for value in available_named],
        mesh_unit=getattr(mesh, "unit", None),
        element_ids=[int(value) for value in element_ids],
        element_node_ids=element_nodes,
        node_ids=[int(value) for value in mesh_node_ids],
        node_coordinates={
            int(node_id): node_xyz(mesh, int(node_id))
            for node_id in mesh_node_ids
        },
        grid_payload=copy_grid_payload(mesh_grid_payload(mesh)),
        nodal_displacements={
            int(node_id): [float(value) for value in values[:3]]
            for node_id, values in (nodal_displacements or {}).items()
        },
        deformation=dict(deformation) if deformation is not None else None,
        resolved_reference_frame=(
            dict(resolved_reference_frame)
            if resolved_reference_frame is not None
            else None
        ),
    )


def node_xyz(mesh: Any, node_id: int) -> Vector:
    node = mesh.nodes.node_by_id(int(node_id))
    return [float(value) for value in node.coordinates]


def element_node_ids(mesh: Any, element_id: int) -> List[int]:
    element = mesh.elements.element_by_id(int(element_id))
    return [int(node_id) for node_id in element.node_ids]


_SUPPORTED_SOLID_FACE_INDICES: Dict[int, Tuple[Tuple[int, ...], ...]] = {
    4: ((0, 1, 2), (0, 1, 3), (1, 2, 3), (0, 2, 3)),
    5: (
        (0, 1, 2, 3),
        (0, 1, 4),
        (1, 2, 4),
        (2, 3, 4),
        (3, 0, 4),
    ),
    6: (
        (0, 1, 2),
        (3, 4, 5),
        (0, 1, 4, 3),
        (1, 2, 5, 4),
        (2, 0, 3, 5),
    ),
    8: (
        (0, 1, 2, 3),
        (4, 5, 6, 7),
        (0, 1, 5, 4),
        (1, 2, 6, 5),
        (2, 3, 7, 6),
        (3, 0, 4, 7),
    ),
    10: (
        (0, 1, 2, 4, 5, 6),
        (0, 1, 3, 4, 8, 7),
        (1, 2, 3, 5, 9, 8),
        (2, 0, 3, 6, 7, 9),
    ),
    13: (
        (0, 1, 2, 3, 5, 6, 7, 8),
        (0, 1, 4, 5, 10, 9),
        (1, 2, 4, 6, 11, 10),
        (2, 3, 4, 7, 12, 11),
        (3, 0, 4, 8, 9, 12),
    ),
    15: (
        (0, 1, 2, 6, 7, 8),
        (3, 4, 5, 9, 10, 11),
        (0, 1, 4, 3, 6, 13, 9, 12),
        (1, 2, 5, 4, 7, 14, 10, 13),
        (2, 0, 3, 5, 8, 12, 11, 14),
    ),
    20: (
        (0, 1, 2, 3, 8, 9, 10, 11),
        (4, 5, 6, 7, 12, 13, 14, 15),
        (0, 1, 5, 4, 8, 17, 12, 16),
        (1, 2, 6, 5, 9, 18, 13, 17),
        (2, 3, 7, 6, 10, 19, 14, 18),
        (3, 0, 4, 7, 11, 16, 15, 19),
    ),
}

DUPLICATE_INTERFACE_RELATIVE_MATCH_EPSILON = 1.0e-9
DUPLICATE_INTERFACE_MIN_MATCH_EPSILON = 1.0e-12
DUPLICATE_INTERFACE_MAX_MATCH_EPSILON = 1.0e-7


def _complete_supported_solid_face(
    element_node_ids: Sequence[int],
    on_plane_node_ids: Sequence[int],
    coordinates: Dict[int, Vector],
) -> Tuple[Optional[Dict[str, Any]], str]:
    templates = _SUPPORTED_SOLID_FACE_INDICES.get(len(element_node_ids))
    if templates is None:
        return None, "unsupported_element_connectivity"
    on_plane = set(int(value) for value in on_plane_node_ids)
    for face_index, indices in enumerate(templates):
        face_node_ids = [int(element_node_ids[index]) for index in indices]
        if set(face_node_ids) != on_plane:
            continue
        points = [coordinates[node_id] for node_id in face_node_ids]
        maximum_span = max(
            (math.dist(first, second) for first in points for second in points),
            default=0.0,
        )
        maximum_double_area = max(
            (
                math.sqrt(
                    sum(
                        value * value
                        for value in cross(
                            [second[index] - first[index] for index in range(3)],
                            [third[index] - first[index] for index in range(3)],
                        )
                    )
                )
                for first in points
                for second in points
                for third in points
            ),
            default=0.0,
        )
        area_floor = max(maximum_span * maximum_span * 1.0e-12, 1.0e-24)
        if maximum_double_area <= area_floor:
            return None, "zero_area_face"
        return (
            {
                "node_ids": face_node_ids,
                "points": points,
                "face_index": face_index,
                "element_node_count": len(element_node_ids),
            },
            "complete_supported_solid_face",
        )
    return None, "partial_face_edge_or_vertex_touch"


def _unclassified_tangent_face_candidate(
    element_node_ids: Sequence[int],
    on_plane_node_ids: Sequence[int],
    coordinates: Dict[int, Vector],
    face_status: str,
) -> Optional[Dict[str, Any]]:
    """Retain plausible faces whose element connectivity is not yet supported."""

    if face_status != "unsupported_element_connectivity":
        return None
    face_node_ids = [int(value) for value in on_plane_node_ids]
    if len(face_node_ids) < 3:
        return None
    points = [coordinates[node_id] for node_id in face_node_ids]
    maximum_span = max(
        (math.dist(first, second) for first in points for second in points),
        default=0.0,
    )
    maximum_double_area = max(
        (
            math.sqrt(
                sum(
                    value * value
                    for value in cross(
                        [second[index] - first[index] for index in range(3)],
                        [third[index] - first[index] for index in range(3)],
                    )
                )
            )
            for first in points
            for second in points
            for third in points
        ),
        default=0.0,
    )
    area_floor = max(maximum_span * maximum_span * 1.0e-12, 1.0e-24)
    if maximum_double_area <= area_floor:
        return None
    return {
        "node_ids": face_node_ids,
        "points": points,
        "element_node_count": len(element_node_ids),
        "classification_status": face_status,
    }


def _coordinate_matched_tangent_pairs(
    negative_faces: Sequence[Dict[str, Any]],
    positive_faces: Sequence[Dict[str, Any]],
) -> Tuple[List[Tuple[Dict[str, Any], Dict[str, Any]]], Dict[str, Any]]:
    """Pair opposite tangent faces that duplicate the same geometric interface."""

    all_faces = list(negative_faces) + list(positive_faces)
    face_scale = max(
        (
            math.dist(first, second)
            for face in all_faces
            for first in face["points"]
            for second in face["points"]
        ),
        default=0.0,
    )
    match_tolerance = min(
        DUPLICATE_INTERFACE_MAX_MATCH_EPSILON,
        max(
            DUPLICATE_INTERFACE_MIN_MATCH_EPSILON,
            face_scale * DUPLICATE_INTERFACE_RELATIVE_MATCH_EPSILON,
        ),
    )

    def face_centroid(face: Dict[str, Any]) -> Tuple[float, float, float]:
        points = face["points"]
        count = float(len(points))
        return tuple(sum(point[index] for point in points) / count for index in range(3))

    def bucket_key(face: Dict[str, Any]) -> Tuple[int, int, int, int]:
        center = face_centroid(face)
        return (
            len(face["node_ids"]),
            *(math.floor(value / match_tolerance) for value in center),
        )

    def match_error(
        negative_face: Dict[str, Any], positive_face: Dict[str, Any]
    ) -> Optional[float]:
        negative_ids = set(negative_face["node_ids"])
        positive_ids = set(positive_face["node_ids"])
        if len(negative_ids) != len(positive_ids) or negative_ids & positive_ids:
            return None
        negative_points = sorted(tuple(point) for point in negative_face["points"])
        positive_points = sorted(tuple(point) for point in positive_face["points"])
        maximum_error = max(
            math.dist(negative_point, positive_point)
            for negative_point, positive_point in zip(negative_points, positive_points)
        )
        return maximum_error if maximum_error <= match_tolerance else None

    positive_buckets: Dict[Tuple[int, int, int, int], List[Dict[str, Any]]] = {}
    for face in positive_faces:
        positive_buckets.setdefault(bucket_key(face), []).append(face)

    candidate_edges: List[Tuple[int, int, float]] = []
    negative_by_id = {int(face["element_id"]): face for face in negative_faces}
    positive_by_id = {int(face["element_id"]): face for face in positive_faces}
    for negative_face in sorted(negative_faces, key=lambda item: int(item["element_id"])):
        node_count, cell_x, cell_y, cell_z = bucket_key(negative_face)
        for delta_x in (-1, 0, 1):
            for delta_y in (-1, 0, 1):
                for delta_z in (-1, 0, 1):
                    key = (
                        node_count,
                        cell_x + delta_x,
                        cell_y + delta_y,
                        cell_z + delta_z,
                    )
                    for positive_face in positive_buckets.get(key, []):
                        positive_element_id = int(positive_face["element_id"])
                        error = match_error(negative_face, positive_face)
                        if error is not None:
                            candidate_edges.append(
                                (
                                    int(negative_face["element_id"]),
                                    positive_element_id,
                                    float(error),
                                )
                            )
    negative_degree: Dict[int, int] = {}
    positive_degree: Dict[int, int] = {}
    for negative_id, positive_id, _error in candidate_edges:
        negative_degree[negative_id] = negative_degree.get(negative_id, 0) + 1
        positive_degree[positive_id] = positive_degree.get(positive_id, 0) + 1
    ambiguous_edges = [
        edge
        for edge in candidate_edges
        if negative_degree[edge[0]] != 1 or positive_degree[edge[1]] != 1
    ]
    if ambiguous_edges:
        return [], {
            "status": "ambiguous_coordinate_matches_fail_closed",
            "coordinate_match_epsilon": match_tolerance,
            "face_scale": face_scale,
            "candidate_match_count": len(candidate_edges),
            "ambiguous_match_count": len(ambiguous_edges),
            "ambiguous_element_ids_preview": sorted(
                {value for edge in ambiguous_edges for value in edge[:2]}
            )[:30],
            "candidate_edges": candidate_edges,
        }
    pairs = [
        (negative_by_id[negative_id], positive_by_id[positive_id])
        for negative_id, positive_id, _error in candidate_edges
    ]
    return pairs, {
        "status": "unique_one_to_one" if pairs else "no_coordinate_match",
        "coordinate_match_epsilon": match_tolerance,
        "face_scale": face_scale,
        "candidate_match_count": len(candidate_edges),
        "ambiguous_match_count": 0,
        "ambiguous_element_ids_preview": [],
        "candidate_edges": candidate_edges,
    }


def construction_surface_scoping(
    dpf: Any,
    model: Any,
    element_scoping: Any,
    config: SectionConfig,
    axes: Matrix3,
    mesh: Optional[Any] = None,
    *,
    allow_empty_cut: bool = False,
) -> Tuple[Any, Any, Dict[str, Any]]:
    mesh = mesh or model.metadata.meshed_region
    normal = axis_from_coordinate_system(axes, config.section_normal_axis)
    origin = vector3(config.coordinate_system_origin, "coordinate_system_origin")
    raw_ids = object_ids(element_scoping)
    tolerance = float(config.side_filter_tolerance)
    requested_positive = config.extraction_side.lower().startswith("pos")
    requested_negative = config.extraction_side.lower().startswith("neg")

    cut_element_ids: List[int] = []
    not_cut_element_ids: List[int] = []
    selected_node_ids: set[int] = set()
    positive_node_ids: set[int] = set()
    negative_node_ids: set[int] = set()
    on_plane_node_ids: set[int] = set()
    element_failures: List[Dict[str, Any]] = []
    element_nodes_cache: Dict[int, List[int]] = {}
    node_xyz_cache: Dict[int, Vector] = {}
    negative_tangent_faces: List[Dict[str, Any]] = []
    positive_tangent_faces: List[Dict[str, Any]] = []
    negative_unclassified_tangent_faces: List[Dict[str, Any]] = []
    positive_unclassified_tangent_faces: List[Dict[str, Any]] = []
    tangent_face_rejection_counts: Dict[str, int] = {}

    def cached_element_node_ids(element_id: int) -> List[int]:
        cached = element_nodes_cache.get(element_id)
        if cached is None:
            cached = element_node_ids(mesh, element_id)
            element_nodes_cache[element_id] = cached
        return cached

    def cached_node_xyz(node_id: int) -> Vector:
        cached = node_xyz_cache.get(node_id)
        if cached is None:
            cached = node_xyz(mesh, node_id)
            node_xyz_cache[node_id] = cached
        return cached

    for element_id in raw_ids:
        try:
            node_ids = cached_element_node_ids(element_id)
            signed_distances = []
            for node_id in node_ids:
                xyz = cached_node_xyz(node_id)
                signed_distances.append(
                    dot(
                        [xyz[0] - origin[0], xyz[1] - origin[1], xyz[2] - origin[2]],
                        normal,
                    )
                )
        except Exception as exc:
            element_failures.append({"element_id": element_id, "error": repr(exc)})
            continue

        if not signed_distances:
            element_failures.append({"element_id": element_id, "error": "Element has no nodes."})
            continue

        if min(signed_distances) > tolerance or max(signed_distances) < -tolerance:
            not_cut_element_ids.append(element_id)
            continue

        cut_element_ids.append(element_id)
        has_positive = any(value > tolerance for value in signed_distances)
        has_negative = any(value < -tolerance for value in signed_distances)
        plane_node_ids = [
            node_id
            for node_id, signed_distance in zip(node_ids, signed_distances)
            if abs(signed_distance) <= tolerance
        ]
        if plane_node_ids and (has_negative != has_positive):
            element_coordinates = {
                node_id: cached_node_xyz(node_id) for node_id in node_ids
            }
            tangent_face, face_status = _complete_supported_solid_face(
                node_ids,
                plane_node_ids,
                element_coordinates,
            )
            if tangent_face is None:
                tangent_face_rejection_counts[face_status] = (
                    tangent_face_rejection_counts.get(face_status, 0) + 1
                )
                unclassified_face = _unclassified_tangent_face_candidate(
                    node_ids,
                    plane_node_ids,
                    element_coordinates,
                    face_status,
                )
                if unclassified_face is not None:
                    unclassified_face["element_id"] = int(element_id)
                    if has_negative:
                        negative_unclassified_tangent_faces.append(unclassified_face)
                    else:
                        positive_unclassified_tangent_faces.append(unclassified_face)
            else:
                tangent_face["element_id"] = int(element_id)
                if has_negative:
                    negative_tangent_faces.append(tangent_face)
                else:
                    positive_tangent_faces.append(tangent_face)
        for node_id, signed_distance in zip(node_ids, signed_distances):
            if signed_distance > tolerance:
                positive_node_ids.add(node_id)
            elif signed_distance < -tolerance:
                negative_node_ids.add(node_id)
            else:
                on_plane_node_ids.add(node_id)

            if requested_positive:
                if signed_distance >= -tolerance:
                    selected_node_ids.add(node_id)
            elif requested_negative:
                if signed_distance <= tolerance:
                    selected_node_ids.add(node_id)
            else:
                selected_node_ids.add(node_id)

    raw_node_ids_sorted = sorted(
        {
            int(node_id)
            for node_ids in element_nodes_cache.values()
            for node_id in node_ids
        }
    )
    raw_node_points = [cached_node_xyz(node_id) for node_id in raw_node_ids_sorted]
    fallback_centroid = centroid(raw_node_points) if raw_node_points else origin

    if not cut_element_ids and not allow_empty_cut:
        raise ValueError(
            "The construction surface cut selected zero elements. "
            f"side={config.extraction_side}, raw_count={len(raw_ids)}, "
            f"not_cut={len(not_cut_element_ids)}"
        )
    force_summation_terms: List[Dict[str, Any]] = []
    duplicate_pairs: List[Tuple[Dict[str, Any], Dict[str, Any]]] = []
    duplicate_match_evidence: Dict[str, Any] = {
        "status": "not_applicable_for_both_or_unknown_side",
        "coordinate_match_epsilon": None,
        "face_scale": None,
        "candidate_match_count": 0,
        "ambiguous_match_count": 0,
        "ambiguous_element_ids_preview": [],
    }
    if requested_positive or requested_negative:
        duplicate_pairs, duplicate_match_evidence = _coordinate_matched_tangent_pairs(
            negative_tangent_faces,
            positive_tangent_faces,
        )
        if duplicate_match_evidence["status"] == "ambiguous_coordinate_matches_fail_closed":
            raise ValueError(
                "The construction surface intersects an ambiguous stack of "
                "coordinate-coincident, node-disconnected tangent faces. MCF cannot "
                "safely determine the Mechanical-equivalent duplicate-interface "
                "summation signs. Candidate elements: "
                f"{duplicate_match_evidence['ambiguous_element_ids_preview']}. "
                "Move the section plane away from the coincident interface, narrow "
                "the named selection, or remove/remesh the duplicated face sheets."
            )

        all_negative_tangent_faces = (
            negative_tangent_faces + negative_unclassified_tangent_faces
        )
        all_positive_tangent_faces = (
            positive_tangent_faces + positive_unclassified_tangent_faces
        )
        if negative_unclassified_tangent_faces or positive_unclassified_tangent_faces:
            _all_pairs, all_match_evidence = _coordinate_matched_tangent_pairs(
                all_negative_tangent_faces,
                all_positive_tangent_faces,
            )
            unclassified_negative_ids = {
                int(face["element_id"])
                for face in negative_unclassified_tangent_faces
            }
            unclassified_positive_ids = {
                int(face["element_id"])
                for face in positive_unclassified_tangent_faces
            }
            unsafe_edges = [
                edge
                for edge in all_match_evidence["candidate_edges"]
                if edge[0] in unclassified_negative_ids
                or edge[1] in unclassified_positive_ids
            ]
            if unsafe_edges:
                unsafe_element_ids = sorted(
                    {int(value) for edge in unsafe_edges for value in edge[:2]}
                )
                unsupported_counts = sorted(
                    {
                        int(face["element_node_count"])
                        for face in (
                            negative_unclassified_tangent_faces
                            + positive_unclassified_tangent_faces
                        )
                        if int(face["element_id"]) in unsafe_element_ids
                    }
                )
                raise ValueError(
                    "The construction surface may lie on a duplicated, "
                    "node-disconnected interface, but at least one "
                    "coordinate-coincident tangent face has unsupported solid "
                    "connectivity. MCF cannot safely classify or sign that interface. "
                    f"Candidate elements: {unsafe_element_ids[:30]}; element node "
                    f"counts: {unsupported_counts}. Move the section plane away from "
                    "the coincident interface, narrow the named selection, or remesh "
                    "the interface with a supported tetrahedron, pyramid, wedge, or "
                    "hexahedron topology."
                )

    if duplicate_pairs:
        paired_negative_element_ids = {
            int(negative_face["element_id"])
            for negative_face, _positive_face in duplicate_pairs
        }
        paired_positive_element_ids = {
            int(positive_face["element_id"])
            for _negative_face, positive_face in duplicate_pairs
        }
        paired_element_ids = paired_negative_element_ids | paired_positive_element_ids
        baseline_element_ids = [
            element_id for element_id in cut_element_ids if element_id not in paired_element_ids
        ]
        baseline_node_ids: set[int] = set()
        for element_id in baseline_element_ids:
            for node_id in cached_element_node_ids(element_id):
                xyz = cached_node_xyz(node_id)
                signed_distance = dot(
                    [xyz[0] - origin[0], xyz[1] - origin[1], xyz[2] - origin[2]],
                    normal,
                )
                if (
                    (requested_positive and signed_distance >= -tolerance)
                    or (requested_negative and signed_distance <= tolerance)
                ):
                    baseline_node_ids.add(node_id)
        if baseline_element_ids and baseline_node_ids:
            force_summation_terms.append(
                {
                    "role": "baseline",
                    "sign": 1,
                    "element_ids": sorted(baseline_element_ids),
                    "node_ids": sorted(baseline_node_ids),
                }
            )

        negative_interface_element_ids = sorted(paired_negative_element_ids)
        positive_interface_element_ids = sorted(paired_positive_element_ids)
        negative_interface_node_ids = sorted(
            {
                int(node_id)
                for negative_face, _positive_face in duplicate_pairs
                for node_id in negative_face["node_ids"]
            }
        )
        positive_interface_node_ids = sorted(
            {
                int(node_id)
                for _negative_face, positive_face in duplicate_pairs
                for node_id in positive_face["node_ids"]
            }
        )
        tangent_terms = [
            {
                "role": "negative_tangent",
                "sign": 1 if requested_positive else -1,
                "element_ids": negative_interface_element_ids,
                "node_ids": negative_interface_node_ids,
            },
            {
                "role": "positive_tangent",
                "sign": -1 if requested_positive else 1,
                "element_ids": positive_interface_element_ids,
                "node_ids": positive_interface_node_ids,
            },
        ]
        force_summation_terms.extend(tangent_terms)
        selected_node_ids = {
            int(node_id)
            for term in force_summation_terms
            for node_id in term["node_ids"]
        }
        cut_element_ids = sorted(
            {
                int(element_id)
                for term in force_summation_terms
                for element_id in term["element_ids"]
            }
        )
    else:
        force_summation_terms = [
            {
                "role": "baseline",
                "sign": 1,
                "element_ids": sorted(cut_element_ids),
                "node_ids": sorted(selected_node_ids),
            }
        ]

    if not selected_node_ids and not allow_empty_cut:
        raise ValueError(
            "The construction surface side filter selected zero nodes. "
            f"side={config.extraction_side}, cut_element_count={len(cut_element_ids)}"
        )

    selected_node_ids_sorted = sorted(selected_node_ids)
    positive_node_ids_sorted = sorted(positive_node_ids)
    negative_node_ids_sorted = sorted(negative_node_ids)
    on_plane_node_ids_sorted = sorted(on_plane_node_ids)

    selected_node_points = [cached_node_xyz(node_id) for node_id in selected_node_ids_sorted]
    selected_node_centroid = (
        centroid(selected_node_points) if selected_node_points else fallback_centroid
    )
    selected_signed_distances = [
        dot([point[0] - origin[0], point[1] - origin[1], point[2] - origin[2]], normal)
        for point in selected_node_points
    ]

    if duplicate_pairs and requested_positive:
        reference_node_ids = negative_interface_node_ids
        reference_node_side = "duplicate_interface_negative_sheet_for_positive_probe"
    elif duplicate_pairs and requested_negative:
        reference_node_ids = positive_interface_node_ids
        reference_node_side = "duplicate_interface_positive_sheet_for_negative_probe"
    elif requested_positive:
        reference_node_ids = (
            negative_node_ids_sorted or on_plane_node_ids_sorted or selected_node_ids_sorted
        )
        reference_node_side = "opposite_negative_for_positive_probe"
    elif requested_negative:
        reference_node_ids = (
            positive_node_ids_sorted or on_plane_node_ids_sorted or selected_node_ids_sorted
        )
        reference_node_side = "opposite_positive_for_negative_probe"
    else:
        reference_node_ids = on_plane_node_ids_sorted or selected_node_ids_sorted
        reference_node_side = "on_plane_or_selected_for_all_sides"
    reference_node_ids_sorted = sorted(reference_node_ids)

    if reference_node_ids_sorted:
        moment_reference_centroid = centroid(
            [cached_node_xyz(node_id) for node_id in reference_node_ids_sorted]
        )
    else:
        moment_reference_centroid = fallback_centroid
        reference_node_side = "selected_mesh_centroid_fallback_no_side_nodes"
    element_scope = make_scoping(dpf, cut_element_ids, dpf.locations.elemental)
    node_scope = make_scoping(dpf, selected_node_ids_sorted, dpf.locations.nodal)
    force_summation_terms_preview = [
        {
            "role": term["role"],
            "sign": int(term["sign"]),
            "element_count": len(term["element_ids"]),
            "node_count": len(term["node_ids"]),
            "element_ids_preview": term["element_ids"][:30],
            "node_ids_preview": term["node_ids"][:30],
        }
        for term in force_summation_terms
    ]

    return element_scope, node_scope, {
        "method": "construction_surface_cut_elements_side_nodes",
        "normal_axis": config.section_normal_axis,
        "extraction_side": config.extraction_side,
        "tolerance": tolerance,
        "allow_empty_cut": bool(allow_empty_cut),
        "cut_status": "cut" if cut_element_ids else "no_cut",
        "raw_element_count": len(raw_ids),
        "cut_element_count": len(cut_element_ids),
        "raw_node_count": len(raw_node_ids_sorted),
        "selected_node_count": len(selected_node_ids_sorted),
        "selected_node_centroid": selected_node_centroid,
        "selected_node_signed_distance_min": (
            min(selected_signed_distances) if selected_signed_distances else None
        ),
        "selected_node_signed_distance_max": (
            max(selected_signed_distances) if selected_signed_distances else None
        ),
        "moment_reference_centroid": moment_reference_centroid,
        "moment_reference_node_side": reference_node_side,
        "moment_reference_node_count": len(reference_node_ids_sorted),
        "moment_reference_node_ids": reference_node_ids_sorted,
        "moment_reference_node_ids_preview": reference_node_ids_sorted[:30],
        "positive_node_count": len(positive_node_ids_sorted),
        "negative_node_count": len(negative_node_ids_sorted),
        "on_plane_node_count": len(on_plane_node_ids_sorted),
        "not_cut_element_count": len(not_cut_element_ids),
        "element_failure_count": len(element_failures),
        "force_summation_policy": (
            "coordinate_matched_duplicate_tangent_interface"
            if duplicate_pairs
            else "single_baseline"
        ),
        "force_summation_term_count": len(force_summation_terms),
        "force_summation_terms": force_summation_terms,
        "force_summation_terms_preview": force_summation_terms_preview,
        "duplicate_tangent_pair_count": len(duplicate_pairs),
        "duplicate_tangent_negative_element_count": (
            len({int(pair[0]["element_id"]) for pair in duplicate_pairs})
        ),
        "duplicate_tangent_positive_element_count": (
            len({int(pair[1]["element_id"]) for pair in duplicate_pairs})
        ),
        "duplicate_tangent_candidate_negative_face_count": len(negative_tangent_faces),
        "duplicate_tangent_candidate_positive_face_count": len(positive_tangent_faces),
        "duplicate_tangent_unclassified_negative_face_count": len(
            negative_unclassified_tangent_faces
        ),
        "duplicate_tangent_unclassified_positive_face_count": len(
            positive_unclassified_tangent_faces
        ),
        "duplicate_tangent_face_rejection_counts": tangent_face_rejection_counts,
        "duplicate_tangent_match_status": duplicate_match_evidence["status"],
        "duplicate_tangent_coordinate_tolerance": duplicate_match_evidence[
            "coordinate_match_epsilon"
        ],
        "duplicate_tangent_face_scale": duplicate_match_evidence["face_scale"],
        "duplicate_tangent_candidate_match_count": duplicate_match_evidence[
            "candidate_match_count"
        ],
        "duplicate_tangent_ambiguous_match_count": duplicate_match_evidence[
            "ambiguous_match_count"
        ],
        "duplicate_tangent_ambiguous_element_ids_preview": duplicate_match_evidence[
            "ambiguous_element_ids_preview"
        ],
        "raw_element_ids": raw_ids,
        "raw_node_ids": raw_node_ids_sorted,
        "cut_element_ids": cut_element_ids,
        "selected_node_ids": selected_node_ids_sorted,
        "positive_node_ids": positive_node_ids_sorted,
        "negative_node_ids": negative_node_ids_sorted,
        "on_plane_node_ids": on_plane_node_ids_sorted,
        "not_cut_element_ids": not_cut_element_ids,
        "selected_node_coordinates": {
            str(node_id): cached_node_xyz(node_id) for node_id in selected_node_ids_sorted
        },
        "cut_element_ids_preview": cut_element_ids[:30],
        "selected_node_ids_preview": selected_node_ids_sorted[:30],
        "positive_node_ids_preview": positive_node_ids_sorted[:30],
        "negative_node_ids_preview": negative_node_ids_sorted[:30],
        "on_plane_node_ids_preview": on_plane_node_ids_sorted[:30],
        "not_cut_element_ids_preview": not_cut_element_ids[:30],
        "element_failures_preview": element_failures[:10],
    }


def side_filter_with_mesh_coordinates(
    side_filter_info: Dict[str, Any],
    mesh: Any,
    origin: Sequence[float],
    normal: Sequence[float],
    *,
    scoping_geometry_state: str,
) -> Dict[str, Any]:
    """Keep an established section scope while updating its spatial evidence."""

    result = dict(side_filter_info)
    origin_xyz = vector3(origin, "coordinate_system_origin")
    normal_xyz = vector3(normal, "section_normal")
    selected_node_ids = [
        int(value) for value in side_filter_info.get("selected_node_ids", [])
    ]
    raw_node_ids = [int(value) for value in side_filter_info.get("raw_node_ids", [])]
    reference_node_ids = [
        int(value) for value in side_filter_info.get("moment_reference_node_ids", [])
    ]
    coordinate_ids = set(selected_node_ids)
    coordinate_ids.update(raw_node_ids)
    coordinate_ids.update(reference_node_ids)
    coordinates = {
        node_id: node_xyz(mesh, node_id)
        for node_id in coordinate_ids
    }
    selected_points = [coordinates[node_id] for node_id in selected_node_ids]
    raw_points = [coordinates[node_id] for node_id in raw_node_ids]
    fallback_centroid = centroid(raw_points) if raw_points else origin_xyz
    selected_centroid = (
        centroid(selected_points) if selected_points else fallback_centroid
    )
    selected_signed_distances = [
        dot(
            [
                point[0] - origin_xyz[0],
                point[1] - origin_xyz[1],
                point[2] - origin_xyz[2],
            ],
            normal_xyz,
        )
        for point in selected_points
    ]
    reference_points = [coordinates[node_id] for node_id in reference_node_ids]
    result.update(
        {
            "scoping_geometry_state": str(scoping_geometry_state),
            "coordinate_geometry_state": "deformed",
            "selected_node_coordinates": {
                str(node_id): point
                for node_id, point in zip(selected_node_ids, selected_points)
            },
            "selected_node_centroid": selected_centroid,
            "selected_node_signed_distance_min": (
                min(selected_signed_distances) if selected_signed_distances else None
            ),
            "selected_node_signed_distance_max": (
                max(selected_signed_distances) if selected_signed_distances else None
            ),
            "moment_reference_centroid": (
                centroid(reference_points) if reference_points else fallback_centroid
            ),
        }
    )
    return result


FULL_SIDE_FILTER_KEYS = {
    "raw_element_ids",
    "cut_element_ids",
    "selected_node_ids",
    "positive_node_ids",
    "negative_node_ids",
    "on_plane_node_ids",
    "not_cut_element_ids",
    "moment_reference_node_ids",
    "selected_node_coordinates",
    "raw_node_ids",
    "force_summation_terms",
}


def compact_side_filter_info(side_filter_info: Dict[str, Any]) -> Dict[str, Any]:
    return {
        key: value
        for key, value in side_filter_info.items()
        if key not in FULL_SIDE_FILTER_KEYS
    }


def deformation_scope_change_evidence(
    reference_side_filter: Dict[str, Any],
    deformed_side_filter: Dict[str, Any],
) -> Dict[str, Any]:
    reference_elements = {
        int(value) for value in reference_side_filter.get("cut_element_ids", [])
    }
    deformed_elements = {
        int(value) for value in deformed_side_filter.get("cut_element_ids", [])
    }
    reference_nodes = {
        int(value) for value in reference_side_filter.get("selected_node_ids", [])
    }
    deformed_nodes = {
        int(value) for value in deformed_side_filter.get("selected_node_ids", [])
    }
    added_elements = sorted(deformed_elements - reference_elements)
    removed_elements = sorted(reference_elements - deformed_elements)
    added_nodes = sorted(deformed_nodes - reference_nodes)
    removed_nodes = sorted(reference_nodes - deformed_nodes)
    return {
        "reference_cut_element_count": len(reference_elements),
        "deformed_cut_element_count": len(deformed_elements),
        "added_cut_element_count": len(added_elements),
        "removed_cut_element_count": len(removed_elements),
        "added_cut_element_ids_preview": added_elements[:30],
        "removed_cut_element_ids_preview": removed_elements[:30],
        "reference_selected_node_count": len(reference_nodes),
        "deformed_selected_node_count": len(deformed_nodes),
        "added_selected_node_count": len(added_nodes),
        "removed_selected_node_count": len(removed_nodes),
        "added_selected_node_ids_preview": added_nodes[:30],
        "removed_selected_node_ids_preview": removed_nodes[:30],
        "element_membership_changed": bool(added_elements or removed_elements),
        "node_membership_changed": bool(added_nodes or removed_nodes),
    }


def validate_visualization_paths(config: SectionConfig) -> None:
    if not config.modal_rst:
        raise ValueError("Modal RST path is empty.")
    if not Path(config.modal_rst).is_file():
        raise FileNotFoundError(f"Modal RST does not exist: {config.modal_rst}")
    if (
        config.external_named_selection_path
        and not Path(config.external_named_selection_path).is_file()
    ):
        raise FileNotFoundError(
            "External named-selection file does not exist: "
            f"{config.external_named_selection_path}"
        )


def mesh_element_ids(mesh: Any) -> List[int]:
    elements = getattr(mesh, "elements", None)
    if elements is None:
        return []
    scoping = getattr(elements, "scoping", None)
    if scoping is None:
        return []
    try:
        return object_ids(scoping)
    except Exception:
        return []


def mesh_grid_payload(mesh: Any) -> Dict[str, Any]:
    cached_payload = getattr(mesh, "_grid_payload", None)
    if isinstance(cached_payload, dict):
        return copy_grid_payload(cached_payload)

    try:
        grid = getattr(mesh, "grid")
        if callable(grid):
            grid = grid()
    except Exception as exc:
        return {"points": [], "cells": [], "celltypes": [], "warning": repr(exc)}
    if grid is None:
        return {"points": [], "cells": [], "celltypes": [], "warning": "Mesh grid is unavailable."}

    try:
        points = [
            [float(value) for value in point[:3]]
            for point in getattr(grid, "points", [])
        ]
        cells = [int(value) for value in getattr(grid, "cells", [])]
        celltypes = [int(value) for value in getattr(grid, "celltypes", [])]
    except Exception as exc:
        return {"points": [], "cells": [], "celltypes": [], "warning": repr(exc)}

    return {
        "points": points,
        "cells": cells,
        "celltypes": celltypes,
        "element_ids": mesh_element_ids(mesh),
    }


def unique_element_node_ids(mesh: Any, element_ids: Sequence[int]) -> List[int]:
    node_ids: set[int] = set()
    for element_id in element_ids:
        try:
            node_ids.update(element_node_ids(mesh, int(element_id)))
        except Exception:
            continue
    return sorted(node_ids)


def node_marker_payload(mesh: Any, node_ids: Sequence[int]) -> List[Dict[str, Any]]:
    markers: List[Dict[str, Any]] = []
    for node_id in node_ids:
        try:
            markers.append({"id": int(node_id), "xyz": node_xyz(mesh, int(node_id))})
        except Exception:
            continue
    return markers


def mesh_coordinate_points(mesh: Any) -> List[Vector]:
    grid_payload = mesh_grid_payload(mesh)
    points = [
        [float(value) for value in point[:3]]
        for point in (grid_payload.get("points") or [])
        if len(point) >= 3
    ]
    if points:
        return points
    return [
        [float(marker["xyz"][index]) for index in range(3)]
        for marker in node_marker_payload(
            mesh,
            unique_element_node_ids(mesh, mesh_element_ids(mesh)),
        )
    ]


def point_bounds_distance(
    point: Sequence[float],
    bounds: Sequence[Tuple[float, float]],
) -> float:
    squared = 0.0
    for index, (lower, upper) in enumerate(bounds[:3]):
        value = float(point[index])
        if value < lower:
            squared += (lower - value) ** 2
        elif value > upper:
            squared += (value - upper) ** 2
    return math.sqrt(squared)


def coordinate_bounds(points: Sequence[Sequence[float]]) -> List[Tuple[float, float]]:
    return [
        (
            min(float(point[index]) for point in points),
            max(float(point[index]) for point in points),
        )
        for index in range(3)
    ]


def config_with_origin_in_mesh_units(
    config: SectionConfig,
    mesh: Any,
    log: LogFn = None,
) -> Tuple[SectionConfig, Optional[Dict[str, Any]]]:
    cfg = config_from_mapping(asdict(config))
    mesh_unit = str(getattr(mesh, "unit", "") or "").strip()
    if not mesh_unit:
        return cfg, None

    factor = length_unit_conversion_factor(GUI_ORIGIN_UNIT, mesh_unit)
    if factor is None or abs(factor - 1.0) <= 1.0e-12:
        return cfg, None

    original_origin = vector3(cfg.coordinate_system_origin, "coordinate_system_origin")
    scaled_origin = scale_vector(original_origin, factor)
    conversion = {
        "source_unit": GUI_ORIGIN_UNIT,
        "mesh_unit": mesh_unit,
        "factor": float(factor),
        "source_origin": original_origin,
        "mesh_unit_origin": scaled_origin,
    }
    points = mesh_coordinate_points(mesh)
    if points:
        bounds = coordinate_bounds(points)
        conversion["unscaled_distance_to_selected_mesh"] = float(
            point_bounds_distance(original_origin, bounds)
        )
        conversion["scaled_distance_to_selected_mesh"] = float(
            point_bounds_distance(scaled_origin, bounds)
        )

    cfg.coordinate_system_origin = [float(value) for value in scaled_origin]
    emit(
        log,
        "Coordinate-system origin converted from GUI units to mesh units for DPF geometry: "
        f"{conversion['source_origin']} {conversion['source_unit']} -> "
        f"{conversion['mesh_unit_origin']} {conversion['mesh_unit']}.",
    )
    return cfg, conversion


def selected_mesh_plane_points(
    grid_payload: Dict[str, Any],
    all_node_markers: Sequence[Dict[str, Any]],
) -> List[Vector]:
    points = [
        [float(value) for value in marker["xyz"][:3]]
        for marker in all_node_markers
        if "xyz" in marker
    ]
    if points:
        return points
    return [
        [float(value) for value in point[:3]]
        for point in (grid_payload.get("points") or [])
    ]


def side_filter_node_points(
    mesh: Any,
    node_ids: Sequence[int],
    side_filter_info: Dict[str, Any],
) -> List[Vector]:
    coordinates = side_filter_info.get("selected_node_coordinates") or {}
    points: List[Vector] = []
    for node_id in node_ids:
        value = coordinates.get(str(int(node_id)))
        if value is None:
            value = coordinates.get(int(node_id))
        if value is not None:
            points.append([float(component) for component in value[:3]])
            continue
        try:
            points.append(node_xyz(mesh, int(node_id)))
        except Exception:
            continue
    return points


def visualization_plane_points(
    mesh: Any,
    grid_payload: Dict[str, Any],
    side_filter_info: Dict[str, Any],
    all_node_markers: Sequence[Dict[str, Any]],
) -> Tuple[List[Vector], str]:
    selected_node_ids = [int(value) for value in side_filter_info.get("selected_node_ids", [])]
    selected_points = side_filter_node_points(mesh, selected_node_ids, side_filter_info)
    if len(selected_points) >= 2:
        return selected_points, "selected_side_nodes"

    cut_element_ids = [int(value) for value in side_filter_info.get("cut_element_ids", [])]
    cut_node_ids = unique_element_node_ids(mesh, cut_element_ids)
    cut_points = [marker["xyz"] for marker in node_marker_payload(mesh, cut_node_ids)]
    if len(cut_points) >= 2:
        return cut_points, "cut_element_nodes"

    raw_points = selected_mesh_plane_points(grid_payload, all_node_markers)
    return raw_points, "raw_selected_nodes"


def section_geometry_points_from_mesh(
    mesh: Any,
    side_filter_info: Dict[str, Any],
) -> Tuple[List[Vector], str]:
    selected_node_ids = [int(value) for value in side_filter_info.get("selected_node_ids", [])]
    selected_points = side_filter_node_points(mesh, selected_node_ids, side_filter_info)
    if len(selected_points) >= 2:
        return selected_points, "selected_side_nodes"

    cut_element_ids = [int(value) for value in side_filter_info.get("cut_element_ids", [])]
    cut_node_ids = unique_element_node_ids(mesh, cut_element_ids)
    cut_points = [marker["xyz"] for marker in node_marker_payload(mesh, cut_node_ids)]
    if len(cut_points) >= 2:
        return cut_points, "cut_element_nodes"

    raw_element_ids = [int(value) for value in side_filter_info.get("raw_element_ids", [])]
    raw_node_ids = unique_element_node_ids(mesh, raw_element_ids)
    raw_points = [marker["xyz"] for marker in node_marker_payload(mesh, raw_node_ids)]
    return raw_points, "raw_selected_nodes"


def projected_section_geometry_from_mesh(
    mesh: Any,
    config: SectionConfig,
    axes: Matrix3,
    side_filter_info: Dict[str, Any],
) -> Dict[str, Any]:
    points, point_source = section_geometry_points_from_mesh(mesh, side_filter_info)
    return projected_section_geometry(
        config.coordinate_system_origin,
        axes,
        config.section_normal_axis,
        points,
        source_unit=getattr(mesh, "unit", None),
        point_source=point_source,
    )


def bounded_plane_extent(
    raw_extent: float,
    other_extent: float,
    points: Sequence[Sequence[float]],
    origin: Sequence[float],
) -> float:
    raw = abs(float(raw_extent))
    if raw > VISUALIZATION_PLANE_MIN_ABSOLUTE_EXTENT:
        return raw
    coordinate_scale = max(
        [
            abs(float(value))
            for point in list(points) + [origin]
            for value in point[:3]
        ]
        or [0.0]
    )
    return max(
        abs(float(other_extent)) * VISUALIZATION_PLANE_DEGENERATE_EXTENT_FRACTION,
        coordinate_scale * VISUALIZATION_PLANE_DEGENERATE_EXTENT_FRACTION,
        VISUALIZATION_PLANE_MIN_ABSOLUTE_EXTENT,
    )


def projected_section_geometry(
    origin: Sequence[float],
    axes: Matrix3,
    normal_axis_name: str,
    points: Sequence[Sequence[float]],
    *,
    source_unit: Optional[str],
    point_source: str,
) -> Dict[str, Any]:
    width_axis, height_axis = plane_span_axis_names(normal_axis_name)
    scale_to_mm = length_unit_conversion_factor(source_unit, GUI_ORIGIN_UNIT)
    result: Dict[str, Any] = {
        "width_axis": width_axis,
        "height_axis": height_axis,
        "mesh_unit": source_unit,
        "unit": GUI_ORIGIN_UNIT if scale_to_mm is not None else None,
        "point_source": str(point_source),
        "point_count": len(points),
        "width_mm": None,
        "height_mm": None,
        "area_mm2": None,
    }
    if scale_to_mm is None:
        return result

    origin_vector = vector3(list(origin), "coordinate_system_origin")
    span_u, span_v = plane_span_axes(axes, normal_axis_name)
    if not points:
        points = [origin_vector]
    u_values: List[float] = []
    v_values: List[float] = []
    for point in points:
        relative = [
            float(point[0]) - origin_vector[0],
            float(point[1]) - origin_vector[1],
            float(point[2]) - origin_vector[2],
        ]
        u_values.append(dot(relative, span_u))
        v_values.append(dot(relative, span_v))
    width_mm = abs(max(u_values) - min(u_values)) * scale_to_mm
    height_mm = abs(max(v_values) - min(v_values)) * scale_to_mm
    result.update(
        {
            "width_mm": width_mm,
            "height_mm": height_mm,
            "area_mm2": width_mm * height_mm,
        }
    )
    return result


def format_section_geometry(geometry: Optional[Dict[str, Any]]) -> str:
    if not geometry:
        return ""
    width = geometry.get("width_mm")
    height = geometry.get("height_mm")
    area = geometry.get("area_mm2")
    if width is None or height is None or area is None:
        return ""
    return (
        f"section {geometry.get('width_axis', 'width')}={float(width):.6g} mm, "
        f"{geometry.get('height_axis', 'height')}={float(height):.6g} mm, "
        f"area={float(area):.6g} mm^2"
    )


def section_plane_corners(
    origin: Sequence[float],
    axes: Matrix3,
    normal_axis_name: str,
    points: Sequence[Sequence[float]],
) -> List[Vector]:
    origin_vector = vector3(list(origin), "coordinate_system_origin")
    span_u, span_v = plane_span_axes(axes, normal_axis_name)
    if not points:
        points = [origin_vector]

    u_values: List[float] = []
    v_values: List[float] = []
    for point in points:
        relative = [
            float(point[0]) - origin_vector[0],
            float(point[1]) - origin_vector[1],
            float(point[2]) - origin_vector[2],
        ]
        u_values.append(dot(relative, span_u))
        v_values.append(dot(relative, span_v))

    u_min, u_max = min(u_values), max(u_values)
    v_min, v_max = min(v_values), max(v_values)
    u_raw_extent = abs(u_max - u_min)
    v_raw_extent = abs(v_max - v_min)
    u_extent = bounded_plane_extent(u_raw_extent, v_raw_extent, points, origin_vector)
    v_extent = bounded_plane_extent(v_raw_extent, u_raw_extent, points, origin_vector)
    u_pad = u_extent * VISUALIZATION_PLANE_PADDING_FRACTION
    v_pad = v_extent * VISUALIZATION_PLANE_PADDING_FRACTION
    u_min -= u_pad
    u_max += u_pad
    v_min -= v_pad
    v_max += v_pad

    corners: List[Vector] = []
    for u_value, v_value in (
        (u_min, v_min),
        (u_max, v_min),
        (u_max, v_max),
        (u_min, v_max),
    ):
        corners.append(
            add_vectors(
                add_vectors(origin_vector, scale_vector(span_u, u_value)),
                scale_vector(span_v, v_value),
            )
        )
    return corners


def section_visualization_payload_from_mesh(
    *,
    mesh: Any,
    config: SectionConfig,
    axes: Matrix3,
    element_name: str,
    available_named: Sequence[str],
    side_filter_info: Dict[str, Any],
    dpf_version: Optional[str] = None,
    origin_unit_conversion: Optional[Dict[str, Any]] = None,
    deformation: Optional[Dict[str, Any]] = None,
    resolved_reference_frame: Optional[Dict[str, Any]] = None,
    result_signature: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    mesh_unit = getattr(mesh, "unit", None)
    grid_payload = mesh_grid_payload(mesh)
    raw_element_ids = [int(value) for value in side_filter_info.get("raw_element_ids", [])]
    cut_element_ids = [int(value) for value in side_filter_info.get("cut_element_ids", [])]
    selected_node_ids = [int(value) for value in side_filter_info.get("selected_node_ids", [])]
    all_node_ids = unique_element_node_ids(mesh, raw_element_ids)
    all_node_markers = node_marker_payload(mesh, all_node_ids)
    selected_node_markers = node_marker_payload(mesh, selected_node_ids)
    plane_points, plane_point_source = visualization_plane_points(
        mesh,
        grid_payload,
        side_filter_info,
        all_node_markers,
    )
    plane_corners = section_plane_corners(
        config.coordinate_system_origin,
        axes,
        config.section_normal_axis,
        plane_points,
    )
    section_geometry = projected_section_geometry(
        config.coordinate_system_origin,
        axes,
        config.section_normal_axis,
        plane_points,
        source_unit=mesh_unit,
        point_source=plane_point_source,
    )

    element_ids = [int(value) for value in grid_payload.get("element_ids", [])]
    element_index = {element_id: index for index, element_id in enumerate(element_ids)}
    cut_cell_indices: List[int] = []
    missing_cut_element_ids: List[int] = []
    for element_id in cut_element_ids:
        index = element_index.get(element_id)
        if index is None:
            missing_cut_element_ids.append(element_id)
        else:
            cut_cell_indices.append(index)

    moment_reference_xyz = config.coordinate_system_origin
    if config.moment_reference_mode == "mechanical_probe_mesh_centroid":
        moment_reference_xyz = side_filter_info.get("moment_reference_centroid", moment_reference_xyz)
    elif config.moment_reference_mode == "selected_side_node_centroid":
        moment_reference_xyz = side_filter_info.get("selected_node_centroid", moment_reference_xyz)

    warnings: List[str] = []
    if not cut_element_ids:
        warnings.append(
            "Construction plane cuts zero selected elements; adjust origin, normal, or tolerance."
        )
    if cut_element_ids and not selected_node_ids:
        warnings.append(
            "Construction plane found cut elements but selected zero force-summation nodes."
        )

    return {
        "version": 1,
        "dpf_version": dpf_version,
        "geometry_state": (
            str(deformation.get("geometry_state"))
            if deformation
            else "reference"
        ),
        "result_set_id": (
            int(config.result_set_id) if config.result_set_id is not None else None
        ),
        "deformation": dict(deformation) if deformation is not None else None,
        "mesh_unit": mesh_unit,
        "element_named_selection": str(element_name),
        "available_named_selections": [str(value) for value in available_named],
        "normal_axis": config.section_normal_axis,
        "extraction_side": config.extraction_side,
        "side_filter_tolerance": float(config.side_filter_tolerance),
        "moment_reference_mode": config.moment_reference_mode,
        "moment_reference_label": (
            "Moment reference: "
            f"{moment_reference_choice_label(config.moment_reference_mode)}"
        ),
        "moment_reference_xyz": [float(value) for value in moment_reference_xyz[:3]],
        "selected_node_centroid": [
            float(value)
            for value in side_filter_info.get(
                "selected_node_centroid",
                config.coordinate_system_origin,
            )[:3]
        ],
        "result_signature": result_signature or result_visualization_signature(
            config,
            mesh_unit=mesh_unit,
        ),
        "coordinate_system_origin_unit_conversion": origin_unit_conversion,
        "coordinate_system_origin": [float(value) for value in config.coordinate_system_origin],
        "coordinate_system_axes": {
            key: [float(value) for value in values]
            for key, values in config.coordinate_system_axes.items()
        },
        "resolved_reference_frame": (
            dict(resolved_reference_frame)
            if resolved_reference_frame is not None
            else None
        ),
        "mesh": grid_payload,
        "plane": {
            "origin": [float(value) for value in config.coordinate_system_origin],
            "normal": axis_from_coordinate_system(axes, config.section_normal_axis),
            "corners": plane_corners,
            "point_source": plane_point_source,
        },
        "section_geometry": section_geometry,
        "raw_element_ids": raw_element_ids,
        "cut_element_ids": cut_element_ids,
        "cut_cell_indices": cut_cell_indices,
        "missing_cut_element_ids": missing_cut_element_ids,
        "force_summation_node_ids": selected_node_ids,
        "force_summation_nodes": selected_node_markers,
        "all_selected_mesh_nodes": all_node_markers,
        "side_filter": compact_side_filter_info(side_filter_info),
        "warnings": warnings,
        "counts": {
            "raw_element_count": int(side_filter_info.get("raw_element_count", len(raw_element_ids))),
            "cut_element_count": int(side_filter_info.get("cut_element_count", len(cut_element_ids))),
            "force_summation_node_count": int(
                side_filter_info.get("selected_node_count", len(selected_node_ids))
            ),
            "mesh_node_count": len(grid_payload.get("points") or all_node_markers),
            "mesh_cell_count": len(element_ids) or len(cut_cell_indices),
        },
    }


def classify_static_animation_section(
    topology: StaticAnimationTopology,
    points: Any,
    origin: Sequence[float],
    axes: Matrix3,
    *,
    membership_points: Optional[Any] = None,
    compact: bool = False,
) -> Dict[str, Any]:
    """Vectorized construction-plane membership for an interpolated frame."""

    import numpy as np

    coordinates = np.asarray(points, dtype=np.float64)
    if coordinates.shape != topology.reference_points.shape:
        raise ValueError(
            f"Animation point array has shape {coordinates.shape}; "
            f"expected {topology.reference_points.shape}."
        )
    membership_coordinates = (
        coordinates
        if membership_points is None
        else np.asarray(membership_points, dtype=np.float64)
    )
    if membership_coordinates.shape != topology.reference_points.shape:
        raise ValueError(
            f"Animation membership point array has shape {membership_coordinates.shape}; "
            f"expected {topology.reference_points.shape}."
        )
    config = config_from_mapping(topology.mesh_config)
    normal = np.asarray(
        axis_from_coordinate_system(axes, config.section_normal_axis), dtype=np.float64
    )
    origin_array = np.asarray(origin, dtype=np.float64)
    signed = (membership_coordinates - origin_array) @ normal
    connectivity = topology.connectivity_point_indices
    offsets = topology.element_offsets
    tolerance = float(config.side_filter_tolerance)
    if len(offsets):
        connected_signed = signed[connectivity]
        minimums = np.minimum.reduceat(connected_signed, offsets)
        maximums = np.maximum.reduceat(connected_signed, offsets)
        cut_mask = (minimums <= tolerance) & (maximums >= -tolerance)
        lengths = np.diff(np.append(offsets, len(connectivity)))
        cut_connectivity = connectivity[np.repeat(cut_mask, lengths)]
        cut_point_indices = np.unique(cut_connectivity)
    else:
        cut_mask = np.zeros(0, dtype=bool)
        cut_point_indices = np.zeros(0, dtype=np.int64)
    cut_cell_indices = [int(value) for value in np.flatnonzero(cut_mask)]
    cut_element_ids = [topology.element_ids[index] for index in cut_cell_indices]
    not_cut_indices = np.flatnonzero(~cut_mask)
    not_cut_element_ids = (
        []
        if compact
        else [topology.element_ids[index] for index in not_cut_indices]
    )
    positive_indices = cut_point_indices[signed[cut_point_indices] > tolerance]
    negative_indices = cut_point_indices[signed[cut_point_indices] < -tolerance]
    on_plane_indices = cut_point_indices[np.abs(signed[cut_point_indices]) <= tolerance]
    side = str(config.extraction_side or "both").lower()
    if side.startswith("pos"):
        selected_indices = cut_point_indices[signed[cut_point_indices] >= -tolerance]
        reference_indices = (
            negative_indices
            if len(negative_indices)
            else on_plane_indices
            if len(on_plane_indices)
            else selected_indices
        )
        reference_side = "opposite_negative_for_positive_probe"
    elif side.startswith("neg"):
        selected_indices = cut_point_indices[signed[cut_point_indices] <= tolerance]
        reference_indices = (
            positive_indices
            if len(positive_indices)
            else on_plane_indices
            if len(on_plane_indices)
            else selected_indices
        )
        reference_side = "opposite_positive_for_negative_probe"
    else:
        selected_indices = cut_point_indices
        reference_indices = on_plane_indices if len(on_plane_indices) else selected_indices
        reference_side = "on_plane_or_selected_for_all_sides"
    fallback_centroid = (
        coordinates.mean(axis=0) if len(coordinates) else origin_array
    )
    selected_centroid = (
        coordinates[selected_indices].mean(axis=0)
        if len(selected_indices)
        else fallback_centroid
    )
    if len(reference_indices):
        moment_centroid = coordinates[reference_indices].mean(axis=0)
    else:
        moment_centroid = fallback_centroid
        reference_side = "selected_mesh_centroid_fallback_no_side_nodes"

    def node_ids(indices: Any) -> List[int]:
        return [topology.node_ids[int(index)] for index in indices]

    selected_node_ids = node_ids(selected_indices)
    positive_node_ids = node_ids(positive_indices)
    negative_node_ids = node_ids(negative_indices)
    on_plane_node_ids = node_ids(on_plane_indices)
    reference_node_ids = node_ids(reference_indices)
    selected_signed = signed[selected_indices]
    return {
        "method": "construction_surface_cut_elements_side_nodes",
        "scoping_geometry_state": (
            "deformed" if membership_points is None else "reference"
        ),
        "coordinate_geometry_state": "deformed",
        "normal_axis": config.section_normal_axis,
        "extraction_side": config.extraction_side,
        "tolerance": tolerance,
        "allow_empty_cut": True,
        "cut_status": "cut" if cut_element_ids else "no_cut",
        "raw_element_count": len(topology.element_ids),
        "cut_element_count": len(cut_element_ids),
        "raw_node_count": len(topology.node_ids),
        "selected_node_count": len(selected_node_ids),
        "selected_node_centroid": [float(value) for value in selected_centroid],
        "selected_node_signed_distance_min": (
            float(selected_signed.min()) if len(selected_signed) else None
        ),
        "selected_node_signed_distance_max": (
            float(selected_signed.max()) if len(selected_signed) else None
        ),
        "moment_reference_centroid": [float(value) for value in moment_centroid],
        "moment_reference_node_side": reference_side,
        "moment_reference_node_count": len(reference_node_ids),
        "moment_reference_node_ids": reference_node_ids,
        "moment_reference_node_ids_preview": reference_node_ids[:30],
        "positive_node_count": len(positive_node_ids),
        "negative_node_count": len(negative_node_ids),
        "on_plane_node_count": len(on_plane_node_ids),
        "not_cut_element_count": int(len(not_cut_indices)),
        "element_failure_count": 0,
        "raw_element_ids": [] if compact else list(topology.element_ids),
        "raw_node_ids": [] if compact else list(topology.node_ids),
        "cut_element_ids": cut_element_ids,
        "selected_node_ids": selected_node_ids,
        "positive_node_ids": positive_node_ids,
        "negative_node_ids": negative_node_ids,
        "on_plane_node_ids": on_plane_node_ids,
        "not_cut_element_ids": not_cut_element_ids,
        "selected_node_coordinates": {
            str(topology.node_ids[int(index)]): [
                float(value) for value in coordinates[int(index)]
            ]
            for index in selected_indices
        },
        "cut_element_ids_preview": cut_element_ids[:30],
        "selected_node_ids_preview": selected_node_ids[:30],
        "positive_node_ids_preview": positive_node_ids[:30],
        "negative_node_ids_preview": negative_node_ids[:30],
        "on_plane_node_ids_preview": on_plane_node_ids[:30],
        "not_cut_element_ids_preview": (
            not_cut_element_ids[:30]
            if not compact
            else [topology.element_ids[int(index)] for index in not_cut_indices[:30]]
        ),
        "element_failures_preview": [],
        "cut_cell_indices": cut_cell_indices,
    }


def _interpolated_animation_displacements(
    session: StaticAnimationSession,
    interpolation: StaticAnimationInterpolation,
) -> Any:
    import numpy as np

    lower = session.displacement_frame(interpolation.lower_index)
    if interpolation.lower_index == interpolation.upper_index:
        return lower
    upper = session.displacement_frame(interpolation.upper_index)
    fraction = float(interpolation.fraction)
    return np.add(lower, fraction * (upper - lower))


def _interpolate_animation_row(
    rows: Sequence[Sequence[float]],
    interpolation: StaticAnimationInterpolation,
) -> List[float]:
    lower = [float(value) for value in rows[interpolation.lower_index]]
    if interpolation.lower_index == interpolation.upper_index:
        return lower
    upper = [float(value) for value in rows[interpolation.upper_index]]
    fraction = float(interpolation.fraction)
    return [
        lower[index] + fraction * (upper[index] - lower[index])
        for index in range(min(len(lower), len(upper)))
    ]


def build_static_animation_total_overlay(
    visualization_payload: Dict[str, Any],
    result_data: Optional[Dict[str, Any]],
    result_mode: str,
    interpolation: StaticAnimationInterpolation,
) -> Optional[Dict[str, Any]]:
    """Interpolate primary totals only; nodal arrows remain exact-set evidence."""

    if not result_data or result_data.get("analysis_mode") != "static":
        return None
    mode = "moment" if str(result_mode).lower().startswith("moment") else "force"
    if mode == "moment":
        rows = result_data.get("resultants_reference_global") or []
        component_start = 3
        label = "Moment Reaction (Mechanical-equivalent section resultant)"
        color = VISUALIZATION_TOTAL_MOMENT_COLOR
    else:
        rows = result_data.get("resultants_global_origin") or result_data.get(
            "resultants_global"
        ) or []
        component_start = 0
        label = "Force Reaction (Mechanical-equivalent section resultant)"
        color = VISUALIZATION_TOTAL_FORCE_COLOR
    if interpolation.upper_index >= len(rows):
        return None
    total_row = _interpolate_animation_row(rows, interpolation)
    total_vector = total_row[component_start : component_start + 3]
    anchors = result_data.get("moment_reference_xyz_by_set") or []
    if interpolation.upper_index < len(anchors):
        origin = _interpolate_animation_row(anchors, interpolation)[:3]
    else:
        origin = [
            float(value)
            for value in visualization_payload.get("moment_reference_xyz", [0.0, 0.0, 0.0])[:3]
        ]
    magnitude = vector_magnitude(total_vector)
    display_vector = [0.0, 0.0, 0.0]
    if magnitude > 0.0:
        display_length = total_vector_display_length(
            visualization_payload,
            origin,
            total_vector,
            fallback_extent=visualization_scene_extent(visualization_payload),
        )
        display_vector = scale_vector(total_vector, display_length / magnitude)
    units = result_data.get("result_units") or {}
    unit = units.get(mode)
    selected_sets = result_data.get("selected_result_sets") or []
    lower_record = (
        selected_sets[interpolation.lower_index]
        if interpolation.lower_index < len(selected_sets)
        else {"id": interpolation.lower_index + 1}
    )
    if interpolation.lower_index == interpolation.upper_index:
        state_label = str(lower_record.get("label") or f"Set {lower_record.get('id')}")
    else:
        upper_record = (
            selected_sets[interpolation.upper_index]
            if interpolation.upper_index < len(selected_sets)
            else {"id": interpolation.upper_index + 1}
        )
        state_label = (
            f"Sets {lower_record.get('id')} -> {upper_record.get('id')} | "
            f"visual interpolation {interpolation.fraction:.1%}"
        )
        label += " (visual interpolation)"
    components = ", ".join(f"{value:.6g}" for value in total_vector)
    suffix = f" {unit}" if unit else ""
    return {
        "mode": mode,
        "label": label,
        "unit": unit,
        "time_index": interpolation.lower_index,
        "time": interpolation.axis_value,
        "nodal_vectors": [],
        "total_vector": {
            "origin": origin,
            "vector": total_vector,
            "display_vector": display_vector,
            "magnitude": magnitude,
            "color": color,
        },
        "text": (
            f"{label} | {state_label} | [{components}]{suffix} | "
            f"|V|={magnitude:.6g}{suffix}"
        ),
        "visual_interpolation": interpolation.lower_index != interpolation.upper_index,
    }


def static_animation_frame_state(
    session: StaticAnimationSession,
    progress: float,
    *,
    mode: str = "smooth",
    deformation_scale: float = 1.0,
) -> Dict[str, Any]:
    """Return NumPy-backed moving geometry for the in-place renderer."""

    if session.closed:
        raise RuntimeError("Animation session is closed.")
    interpolation = static_animation_interpolation(
        [
            {
                "id": record.result_set_id,
                "value": record.result_value,
                "unit": record.result_unit,
            }
            for record in session.records
        ],
        progress,
        mode=mode,
    )
    displacement = _interpolated_animation_displacements(session, interpolation)
    scale = min(max(float(deformation_scale), 0.0), 10000.0)
    topology = session.topology
    points = topology.reference_points + scale * displacement[
        topology.mesh_displacement_indices
    ]
    record = session.records[interpolation.lower_index]
    frame = _animation_reference_frame_from_displacements(
        session,
        interpolation.lower_index,
        displacement,
        deformation_scale=scale,
        result_value=interpolation.axis_value,
    )
    resolved_config = config_from_mapping(topology.mesh_config)
    resolved_config.result_set_id = record.result_set_id
    resolved_config.coordinate_system_origin = list(frame.resolved_origin)
    resolved_config.coordinate_system_axes = {
        key: list(value) for key, value in frame.resolved_axes.items()
    }
    resolved_axes = validate_axes(resolved_config.coordinate_system_axes)
    membership_points = (
        None
        if resolved_config.reference_frame_motion == "follow-geometry"
        else topology.reference_points
    )
    side_info = classify_static_animation_section(
        topology,
        points,
        resolved_config.coordinate_system_origin,
        resolved_axes,
        membership_points=membership_points,
        compact=True,
    )
    import numpy as np

    scene_extent = float(np.ptp(points, axis=0).max()) if len(points) else 1.0
    return {
        "interpolation": interpolation,
        "points": points,
        "frame": frame,
        "config": resolved_config,
        "axes": resolved_axes,
        "side_filter": side_info,
        "record": record,
        "upper_record": session.records[interpolation.upper_index],
        "deformation_scale": scale,
        "scene_extent": max(scene_extent, VISUALIZATION_SCENE_MIN_EXTENT),
    }


def _static_animation_plane_points(
    topology: StaticAnimationTopology,
    points: Any,
    side_info: Dict[str, Any],
    origin: Sequence[float],
    axes: Matrix3,
    normal_axis_name: str,
) -> Tuple[List[Vector], str]:
    """Return the same useful plane span without walking the full mesh in Python."""

    import numpy as np

    selected_coordinates = side_info.get("selected_node_coordinates") or {}
    selected_points = [
        [float(value) for value in selected_coordinates[str(node_id)][:3]]
        for node_id in side_info.get("selected_node_ids", [])
        if str(node_id) in selected_coordinates
    ]
    if len(selected_points) >= 2:
        return selected_points, "selected_side_nodes"

    cut_point_indices: List[int] = []
    for cell_index in side_info.get("cut_cell_indices", []):
        start = int(topology.element_offsets[int(cell_index)])
        end = (
            int(topology.element_offsets[int(cell_index) + 1])
            if int(cell_index) + 1 < len(topology.element_offsets)
            else len(topology.connectivity_point_indices)
        )
        cut_point_indices.extend(
            int(value) for value in topology.connectivity_point_indices[start:end]
        )
    if cut_point_indices:
        unique_indices = np.unique(np.asarray(cut_point_indices, dtype=np.int64))
        if len(unique_indices) >= 2:
            return (
                np.asarray(points, dtype=np.float64)[unique_indices].tolist(),
                "cut_element_nodes",
            )

    coordinates = np.asarray(points, dtype=np.float64)
    if not len(coordinates):
        return [[float(value) for value in origin[:3]]], "frame_origin"
    span_u, span_v = plane_span_axes(axes, normal_axis_name)
    relative = coordinates - np.asarray(origin, dtype=np.float64)
    u_values = relative @ np.asarray(span_u, dtype=np.float64)
    v_values = relative @ np.asarray(span_v, dtype=np.float64)
    extreme_indices = np.unique(
        np.asarray(
            [
                int(np.argmin(u_values)),
                int(np.argmax(u_values)),
                int(np.argmin(v_values)),
                int(np.argmax(v_values)),
            ],
            dtype=np.int64,
        )
    )
    return coordinates[extreme_indices].tolist(), "raw_selected_node_extrema"


def _static_animation_plane_geometry(
    origin: Sequence[float],
    axes: Matrix3,
    normal_axis_name: str,
    points: Sequence[Sequence[float]],
    *,
    source_unit: Optional[str],
    point_source: str,
) -> Tuple[List[Vector], Dict[str, Any]]:
    """Project animation plane points once with NumPy for both corners and size."""

    import numpy as np

    origin_array = np.asarray(origin[:3], dtype=np.float64)
    coordinates = np.asarray(points, dtype=np.float64).reshape((-1, 3))
    if not len(coordinates):
        coordinates = origin_array.reshape((1, 3))
    span_u, span_v = plane_span_axes(axes, normal_axis_name)
    span_u_array = np.asarray(span_u, dtype=np.float64)
    span_v_array = np.asarray(span_v, dtype=np.float64)
    relative = coordinates - origin_array
    u_values = relative @ span_u_array
    v_values = relative @ span_v_array
    u_min = float(u_values.min())
    u_max = float(u_values.max())
    v_min = float(v_values.min())
    v_max = float(v_values.max())
    u_raw_extent = abs(u_max - u_min)
    v_raw_extent = abs(v_max - v_min)
    coordinate_scale = float(
        max(np.max(np.abs(coordinates)), np.max(np.abs(origin_array)), 0.0)
    )

    def bounded_extent(raw_extent: float, other_extent: float) -> float:
        if raw_extent > VISUALIZATION_PLANE_MIN_ABSOLUTE_EXTENT:
            return raw_extent
        return max(
            other_extent * VISUALIZATION_PLANE_DEGENERATE_EXTENT_FRACTION,
            coordinate_scale * VISUALIZATION_PLANE_DEGENERATE_EXTENT_FRACTION,
            VISUALIZATION_PLANE_MIN_ABSOLUTE_EXTENT,
        )

    u_pad = (
        bounded_extent(u_raw_extent, v_raw_extent)
        * VISUALIZATION_PLANE_PADDING_FRACTION
    )
    v_pad = (
        bounded_extent(v_raw_extent, u_raw_extent)
        * VISUALIZATION_PLANE_PADDING_FRACTION
    )
    u_min -= u_pad
    u_max += u_pad
    v_min -= v_pad
    v_max += v_pad
    uv = np.asarray(
        [(u_min, v_min), (u_max, v_min), (u_max, v_max), (u_min, v_max)],
        dtype=np.float64,
    )
    corners = (
        origin_array
        + uv[:, :1] * span_u_array
        + uv[:, 1:] * span_v_array
    ).tolist()
    width_axis, height_axis = plane_span_axis_names(normal_axis_name)
    scale_to_mm = length_unit_conversion_factor(source_unit, GUI_ORIGIN_UNIT)
    geometry: Dict[str, Any] = {
        "width_axis": width_axis,
        "height_axis": height_axis,
        "mesh_unit": source_unit,
        "unit": GUI_ORIGIN_UNIT if scale_to_mm is not None else None,
        "point_source": str(point_source),
        "point_count": len(coordinates),
        "width_mm": None,
        "height_mm": None,
        "area_mm2": None,
    }
    if scale_to_mm is not None:
        width_mm = u_raw_extent * scale_to_mm
        height_mm = v_raw_extent * scale_to_mm
        geometry.update(
            {
                "width_mm": width_mm,
                "height_mm": height_mm,
                "area_mm2": width_mm * height_mm,
            }
        )
    return corners, geometry


def build_static_animation_render_frame(
    session: StaticAnimationSession,
    progress: float,
    *,
    mode: str = "smooth",
    deformation_scale: float = 1.0,
    result_data: Optional[Dict[str, Any]] = None,
    result_mode: str = "force",
) -> Dict[str, Any]:
    """Build one GUI payload from shared topology and at most two endpoint frames."""

    state = static_animation_frame_state(
        session,
        progress,
        mode=mode,
        deformation_scale=deformation_scale,
    )
    interpolation = state["interpolation"]
    evidence_interpolation = static_animation_interpolation(
        [
            {
                "id": item.result_set_id,
                "value": item.result_value,
                "unit": item.result_unit,
            }
            for item in session.records
        ],
        interpolation.progress,
        mode="exact",
    )
    evidence_time_index = int(evidence_interpolation.lower_index)
    evidence_record = session.records[evidence_time_index]
    evidence_config = config_from_mapping(session.topology.mesh_config)
    evidence_config.result_set_id = evidence_record.result_set_id
    evidence_result_signature = result_visualization_signature(
        evidence_config,
        mesh_unit=session.topology.mesh_unit,
    )
    scale = state["deformation_scale"]
    topology = session.topology
    points = state["points"]
    record = state["record"]
    frame = state["frame"]
    resolved_config = state["config"]
    resolved_axes = state["axes"]
    side_info = state["side_filter"]
    deformation = {
        "geometry_state": "deformed",
        "coordinate_equation": "X_visual = X_reference + scale * U_interpolated",
        "coordinate_basis": "global_cartesian",
        "result_set_id": record.result_set_id,
        "mesh_unit": topology.mesh_unit,
        "node_count": len(topology.node_ids),
        "visual_scale": scale,
        "visual_interpolation": interpolation.lower_index != interpolation.upper_index,
    }
    plane_points, plane_point_source = _static_animation_plane_points(
        topology,
        points,
        side_info,
        resolved_config.coordinate_system_origin,
        resolved_axes,
        resolved_config.section_normal_axis,
    )
    plane_corners, section_geometry = _static_animation_plane_geometry(
        resolved_config.coordinate_system_origin,
        resolved_axes,
        resolved_config.section_normal_axis,
        plane_points,
        source_unit=topology.mesh_unit,
        point_source=plane_point_source,
    )
    selected_node_ids = [int(value) for value in side_info.get("selected_node_ids", [])]
    selected_coordinates = side_info.get("selected_node_coordinates") or {}
    selected_node_markers = [
        {
            "id": node_id,
            "xyz": [float(value) for value in selected_coordinates[str(node_id)][:3]],
        }
        for node_id in selected_node_ids
        if str(node_id) in selected_coordinates
    ]
    moment_reference_xyz = list(resolved_config.coordinate_system_origin)
    if resolved_config.moment_reference_mode == "mechanical_probe_mesh_centroid":
        moment_reference_xyz = list(
            side_info.get("moment_reference_centroid") or moment_reference_xyz
        )
    elif resolved_config.moment_reference_mode == "selected_side_node_centroid":
        moment_reference_xyz = list(
            side_info.get("selected_node_centroid") or moment_reference_xyz
        )
    cut_element_ids = [int(value) for value in side_info.get("cut_element_ids", [])]
    warnings: List[str] = []
    if not cut_element_ids:
        warnings.append(
            "Construction plane cuts zero selected elements; adjust origin, normal, or tolerance."
        )
    if cut_element_ids and not selected_node_ids:
        warnings.append(
            "Construction plane found cut elements but selected zero force-summation nodes."
        )
    payload: Dict[str, Any] = {
        "version": 1,
        "dpf_version": None,
        "geometry_state": "deformed",
        "result_set_id": record.result_set_id,
        "deformation": deformation,
        "mesh_unit": topology.mesh_unit,
        "element_named_selection": topology.element_name,
        "available_named_selections": list(topology.available_named_selections),
        "normal_axis": resolved_config.section_normal_axis,
        "extraction_side": resolved_config.extraction_side,
        "side_filter_tolerance": float(resolved_config.side_filter_tolerance),
        "moment_reference_mode": resolved_config.moment_reference_mode,
        "moment_reference_label": (
            "Moment reference: "
            f"{moment_reference_choice_label(resolved_config.moment_reference_mode)}"
        ),
        "moment_reference_xyz": [float(value) for value in moment_reference_xyz[:3]],
        "selected_node_centroid": [
            float(value)
            for value in (
                side_info.get("selected_node_centroid")
                or resolved_config.coordinate_system_origin
            )[:3]
        ],
        "result_signature": result_visualization_signature(
            resolved_config, mesh_unit=topology.mesh_unit
        ),
        "coordinate_system_origin_unit_conversion": topology.origin_unit_conversion,
        "coordinate_system_origin": [
            float(value) for value in resolved_config.coordinate_system_origin
        ],
        "coordinate_system_axes": {
            key: [float(value) for value in values]
            for key, values in resolved_config.coordinate_system_axes.items()
        },
        "resolved_reference_frame": asdict(frame),
        "mesh": {
            "points": topology.reference_points,
            "cells": topology.cells,
            "celltypes": topology.celltypes,
            "element_ids": topology.element_ids,
        },
        "plane": {
            "origin": [float(value) for value in resolved_config.coordinate_system_origin],
            "normal": axis_from_coordinate_system(
                resolved_axes, resolved_config.section_normal_axis
            ),
            "corners": plane_corners,
            "point_source": plane_point_source,
        },
        "section_geometry": section_geometry,
        "raw_element_ids": topology.element_ids,
        "cut_element_ids": cut_element_ids,
        "cut_cell_indices": [
            int(value) for value in side_info.get("cut_cell_indices", [])
        ],
        "missing_cut_element_ids": [],
        "force_summation_node_ids": selected_node_ids,
        "force_summation_nodes": selected_node_markers,
        "all_selected_mesh_nodes": [],
        "side_filter": compact_side_filter_info(side_info),
        "warnings": warnings,
        "counts": {
            "raw_element_count": len(topology.element_ids),
            "cut_element_count": len(cut_element_ids),
            "force_summation_node_count": len(selected_node_ids),
            "mesh_node_count": len(topology.node_ids),
            "mesh_cell_count": len(topology.element_ids),
        },
        "_animation_topology_identity": id(topology),
        "_animation_scene_extent": state["scene_extent"],
    }
    upper_record = state["upper_record"]
    payload["animation"] = {
        "lower_index": interpolation.lower_index,
        "upper_index": interpolation.upper_index,
        "lower_result_set_id": record.result_set_id,
        "upper_result_set_id": upper_record.result_set_id,
        "fraction": interpolation.fraction,
        "progress": interpolation.progress,
        "axis_value": interpolation.axis_value,
        "exact": interpolation.lower_index == interpolation.upper_index,
        "mode": str(mode).strip().lower(),
        "deformation_scale": scale,
        "visual_only": True,
        "time_axis_warning": interpolation.warning,
        "evidence_time_index": evidence_time_index,
        "evidence_result_set_id": evidence_record.result_set_id,
        "evidence_result_signature": evidence_result_signature,
    }
    payload["_animation_points_numpy"] = points
    if interpolation.warning:
        payload.setdefault("warnings", []).append(interpolation.warning)
    if scale != 1.0:
        payload.setdefault("warnings", []).append(
            f"Displayed deformation is visual-only at {scale:g}x; extraction evidence remains 1x."
        )
    overlay = build_static_animation_total_overlay(
        payload, result_data, result_mode, interpolation
    )
    if overlay is not None:
        payload["result_overlay"] = overlay
    return payload


def visualization_scene_extent(payload: Dict[str, Any]) -> float:
    animation_extent = payload.get("_animation_scene_extent")
    if animation_extent is not None:
        return max(float(animation_extent), VISUALIZATION_SCENE_MIN_EXTENT)
    points: List[Sequence[float]] = []
    mesh_payload = payload.get("mesh") or {}
    mesh_points = mesh_payload.get("points")
    if mesh_points is not None:
        points.extend(mesh_points)
    plane = payload.get("plane") or {}
    points.extend(plane.get("corners") or [])
    points.extend(
        item["xyz"]
        for item in payload.get("force_summation_nodes", [])
        if isinstance(item, dict) and "xyz" in item
    )
    reference_xyz = payload.get("moment_reference_xyz")
    if reference_xyz:
        points.append(reference_xyz)
    if not points:
        return 1.0

    extents = []
    for index in range(3):
        values = [float(point[index]) for point in points]
        extents.append(max(values) - min(values))
    return max(max(extents), VISUALIZATION_SCENE_MIN_EXTENT)


def visualization_bottom_ruler(
    payload: Dict[str, Any],
) -> Optional[Tuple[Vector, Vector, float]]:
    scale = length_unit_conversion_factor(payload.get("mesh_unit"), GUI_ORIGIN_UNIT)
    if scale is None:
        return None

    corners = (payload.get("plane") or {}).get("corners") or []
    if len(corners) != 4:
        return None

    pointa = [float(value) for value in corners[0][:3]]
    pointb = [float(value) for value in corners[1][:3]]
    try:
        offset_axis = normalize(point_delta(corners[0], corners[3]), "ruler_bottom_axis")
    except ValueError:
        offset_axis = [0.0, 0.0, 0.0]
    offset = visualization_scene_extent(payload) * VISUALIZATION_RULER_OFFSET_FRACTION
    return (
        add_vectors(pointa, scale_vector(offset_axis, offset)),
        add_vectors(pointb, scale_vector(offset_axis, offset)),
        scale,
    )


def add_visualization_bottom_ruler(plotter: Any, payload: Dict[str, Any]) -> None:
    ruler = visualization_bottom_ruler(payload)
    add_ruler = getattr(plotter, "add_ruler", None)
    if ruler is None or not callable(add_ruler):
        return
    pointa, pointb, scale = ruler
    add_ruler(
        pointa,
        pointb,
        title=f"Distance [{GUI_ORIGIN_UNIT}]",
        number_minor_ticks=4,
        label_color="#17212b",
        tick_color="#17212b",
        scale=scale,
    )


def visualization_static_ruler_length(plotter: Any, payload: Dict[str, Any]) -> float:
    scale = length_unit_conversion_factor(payload.get("mesh_unit"), GUI_ORIGIN_UNIT)
    if scale is None:
        return 0.0

    width = 0.0
    camera = getattr(plotter, "camera", None)
    window_size = getattr(plotter, "window_size", None) or (1, 1)
    try:
        aspect = float(window_size[0]) / max(float(window_size[1]), 1.0)
    except (TypeError, ValueError, IndexError):
        aspect = 1.0
    if camera is not None:
        if bool(getattr(camera, "parallel_projection", False)):
            width = 2.0 * float(getattr(camera, "parallel_scale", 0.0) or 0.0) * aspect
        else:
            distance = float(getattr(camera, "distance", 0.0) or 0.0)
            view_angle = math.radians(float(getattr(camera, "view_angle", 30.0) or 30.0))
            width = 2.0 * distance * math.tan(view_angle / 2.0) * aspect
    if width <= 0.0:
        width = visualization_scene_extent(payload)
    return width * 0.64 * scale


def add_visualization_static_bottom_ruler(
    plotter: Any,
    payload: Dict[str, Any],
) -> Optional[Any]:
    length = visualization_static_ruler_length(plotter, payload)
    if length <= 0.0:
        return None
    try:
        import vtk
    except Exception:
        return None

    ruler = vtk.vtkAxisActor2D()
    start = ruler.GetPositionCoordinate()
    end = ruler.GetPosition2Coordinate()
    start.SetCoordinateSystemToNormalizedViewport()
    end.SetCoordinateSystemToNormalizedViewport()
    start.SetValue(0.18, VISUALIZATION_STATIC_RULER_VIEWPORT_Y, 0.0)
    end.SetValue(0.82, VISUALIZATION_STATIC_RULER_VIEWPORT_Y, 0.0)
    ruler.SetRange(0.0, length)
    ruler.SetTitle(f"Distance [{GUI_ORIGIN_UNIT}]")
    ruler.SetNumberOfMinorTicks(4)
    ruler.SetTickVisibility(True)
    ruler.SetTickLength(5)
    ruler.SetMinorTickLength(3)
    ruler.GetProperty().SetColor(0.09, 0.13, 0.17)
    ruler.GetLabelTextProperty().SetColor(0.09, 0.13, 0.17)
    ruler.GetTitleTextProperty().SetColor(0.09, 0.13, 0.17)
    add_actor = getattr(plotter, "add_actor", None)
    if callable(add_actor):
        add_actor(ruler, reset_camera=False, pickable=False)
        return ruler
    renderer = getattr(plotter, "renderer", None)
    add_actor_2d = getattr(renderer, "AddActor2D", None)
    if callable(add_actor_2d):
        add_actor_2d(ruler)
        return ruler
    return None


def nodal_force_scalar_bar_args(unit: Optional[str]) -> Dict[str, Any]:
    return {
        "title": label_with_unit("Nodal force", unit),
        "vertical": True,
        "position_x": 0.03,
        "position_y": 0.22,
        "width": 0.055,
        "height": 0.68,
        "title_font_size": 9,
        "label_font_size": 8,
    }


def update_static_ruler_actor_range(actor: Any, plotter: Any, payload: Dict[str, Any]) -> bool:
    length = visualization_static_ruler_length(plotter, payload)
    if length <= 0.0:
        return False
    actor.SetRange(0.0, length)
    return True


def remove_camera_observer(camera: Any, observer_id: Any) -> None:
    remove_observer = getattr(camera, "RemoveObserver", None)
    if observer_id is not None and callable(remove_observer):
        remove_observer(observer_id)


def capture_plotter_camera_state(plotter: Any) -> Optional[Dict[str, Any]]:
    try:
        camera_position = getattr(plotter, "camera_position")
    except Exception:
        return None
    if camera_position is None:
        return None
    state: Dict[str, Any] = {"camera_position": camera_position}
    camera = getattr(plotter, "camera", None)
    try:
        state["parallel_scale"] = float(getattr(camera, "parallel_scale"))
    except Exception:
        state["parallel_scale"] = None
    return state


def restore_or_reset_plotter_camera(
    plotter: Any,
    camera_state: Optional[Dict[str, Any]],
) -> bool:
    if camera_state:
        try:
            plotter.camera_position = camera_state["camera_position"]
        except Exception:
            pass
        else:
            camera = getattr(plotter, "camera", None)
            parallel_scale = camera_state.get("parallel_scale")
            if camera is not None and parallel_scale is not None:
                try:
                    camera.parallel_scale = float(parallel_scale)
                except Exception:
                    pass
            return True
    reset_camera = getattr(plotter, "reset_camera", None)
    if callable(reset_camera):
        reset_camera()
    return False


def add_visualization_ruler(
    plotter: Any,
    payload: Dict[str, Any],
    *,
    visible: bool = True,
    static_bottom: bool = False,
) -> None:
    if not visible:
        return
    if static_bottom:
        add_visualization_static_bottom_ruler(plotter, payload)
    else:
        add_visualization_bottom_ruler(plotter, payload)


def point_delta(a: Sequence[float], b: Sequence[float]) -> Vector:
    return [float(a[index]) - float(b[index]) for index in range(3)]


def section_plane_corner_span(corners: Sequence[Sequence[float]]) -> float:
    if len(corners) < 2:
        return 0.0
    distances = []
    for first_index, first in enumerate(corners):
        for second in corners[first_index + 1 :]:
            distances.append(vector_magnitude(point_delta(first, second)))
    return max(distances or [0.0])


def total_vector_display_length(
    visualization_payload: Dict[str, Any],
    origin: Sequence[float],
    vector: Sequence[float],
    *,
    fallback_extent: float,
) -> float:
    magnitude = vector_magnitude(vector)
    if magnitude <= 0.0:
        return 0.0

    plane = visualization_payload.get("plane") or {}
    corners = plane.get("corners") or []
    plane_span = section_plane_corner_span(corners)
    if len(corners) != 4 or plane_span <= 0.0:
        return max(1.0, float(fallback_extent) * VISUALIZATION_TOTAL_VECTOR_LENGTH_FRACTION)

    corner0 = [float(value) for value in corners[0][:3]]
    edge_u = point_delta(corners[1], corners[0])
    edge_v = point_delta(corners[3], corners[0])
    try:
        axis_u = normalize(edge_u, "section_plane_edge_u")
        axis_v = normalize(edge_v, "section_plane_edge_v")
    except ValueError:
        return max(
            plane_span * VISUALIZATION_TOTAL_VECTOR_NORMAL_PLANE_LENGTH_FRACTION,
            float(fallback_extent) * VISUALIZATION_TOTAL_VECTOR_LENGTH_FRACTION,
        )

    origin_delta = point_delta(origin, corner0)
    unit_vector = scale_vector(vector, 1.0 / magnitude)
    origin_uv = (dot(origin_delta, axis_u), dot(origin_delta, axis_v))
    direction_uv = (dot(unit_vector, axis_u), dot(unit_vector, axis_v))
    corner_uv = [
        (dot(point_delta(corner, corner0), axis_u), dot(point_delta(corner, corner0), axis_v))
        for corner in corners
    ]
    u_min = min(value[0] for value in corner_uv)
    u_max = max(value[0] for value in corner_uv)
    v_min = min(value[1] for value in corner_uv)
    v_max = max(value[1] for value in corner_uv)

    candidates = []
    for origin_coord, direction_coord, lower, upper in (
        (origin_uv[0], direction_uv[0], u_min, u_max),
        (origin_uv[1], direction_uv[1], v_min, v_max),
    ):
        if abs(direction_coord) <= 1.0e-9:
            continue
        boundary = upper if direction_coord > 0.0 else lower
        distance = (boundary - origin_coord) / direction_coord
        if distance > 0.0:
            candidates.append(distance)

    min_length = plane_span * VISUALIZATION_TOTAL_VECTOR_MIN_PLANE_LENGTH_FRACTION
    if candidates:
        requested_length = min(candidates) * VISUALIZATION_TOTAL_VECTOR_PLANE_EXIT_MARGIN
        max_length = plane_span * VISUALIZATION_TOTAL_VECTOR_PLANE_EXIT_MARGIN
        return min(max(requested_length, min_length), max_length)
    return max(
        plane_span * VISUALIZATION_TOTAL_VECTOR_NORMAL_PLANE_LENGTH_FRACTION,
        min_length,
    )


def combine_modal_nodal_force_vectors(
    nodal_force_modal_coefficients: Dict[Any, Sequence[Sequence[float]]],
    modal_coordinate_row: Sequence[float],
    *,
    nodal_moment_modal_coefficients: Optional[
        Dict[Any, Sequence[Sequence[float]]]
    ] = None,
    node_coordinates: Optional[Dict[Any, Sequence[float]]] = None,
    modes_used: Optional[int] = None,
) -> List[Dict[str, Any]]:
    mode_count = min(
        int(modes_used) if modes_used is not None else len(modal_coordinate_row),
        len(modal_coordinate_row),
    )
    vectors: List[Dict[str, Any]] = []
    for node_key, modal_vectors in nodal_force_modal_coefficients.items():
        total = [0.0, 0.0, 0.0]
        for mode_index in range(min(mode_count, len(modal_vectors))):
            scale = float(modal_coordinate_row[mode_index])
            vector = modal_vectors[mode_index]
            for component_index in range(3):
                total[component_index] += scale * float(vector[component_index])
        node_id = int(node_key)
        item: Dict[str, Any] = {
            "node_id": node_id,
            "vector": total,
            "magnitude": vector_magnitude(total),
        }
        if nodal_moment_modal_coefficients is not None:
            moment_vectors = nodal_moment_modal_coefficients.get(node_key)
            if moment_vectors is None:
                moment_vectors = nodal_moment_modal_coefficients.get(str(node_key))
            if moment_vectors is None:
                moment_vectors = nodal_moment_modal_coefficients.get(node_id, [])
            explicit_couple = [0.0, 0.0, 0.0]
            for mode_index in range(min(mode_count, len(moment_vectors))):
                scale = float(modal_coordinate_row[mode_index])
                for component_index in range(3):
                    explicit_couple[component_index] += (
                        scale * float(moment_vectors[mode_index][component_index])
                    )
            item["explicit_couple"] = explicit_couple
        if node_coordinates is not None:
            xyz = node_coordinates.get(node_key)
            if xyz is None:
                xyz = node_coordinates.get(str(node_key))
            if xyz is None:
                xyz = node_coordinates.get(node_id)
            if xyz is not None:
                item["origin"] = [float(value) for value in xyz[:3]]
        vectors.append(item)
    return vectors


def build_nodal_force_time_export_payload(
    result_data: Optional[Dict[str, Any]],
    time_index: int,
) -> Optional[Dict[str, Any]]:
    if not result_data:
        return None
    times = [float(value) for value in result_data.get("times", [])]
    modal_coordinates = result_data.get("modal_coordinates") or []
    static_mode = result_data.get("analysis_mode") == "static"
    if time_index < 0 or time_index >= len(times):
        return None
    if static_mode:
        detail_index = int(result_data.get("detail_time_index", time_index))
        if time_index != detail_index:
            return None
        modal_coordinate_row = [1.0]
    else:
        if time_index >= len(modal_coordinates):
            return None
        modal_coordinate_row = modal_coordinates[time_index]
    rows = combine_modal_nodal_force_vectors(
        result_data.get("nodal_force_modal_coefficients") or {},
        modal_coordinate_row,
        nodal_moment_modal_coefficients=(
            result_data.get("nodal_moment_modal_coefficients")
            if result_data.get("analysis_mode") == "static"
            else None
        ),
        node_coordinates=result_data.get("selected_node_coordinates") or {},
        modes_used=(
            1
            if static_mode
            else int(result_data.get("modes_used") or len(modal_coordinate_row))
        ),
    )
    rows = [row for row in rows if "origin" in row]
    if not rows:
        return None
    rows.sort(key=lambda row: int(row["node_id"]))
    payload = {
        "analysis_mode": result_data.get("analysis_mode") or "modal",
        "time": times[time_index],
        "time_index": int(time_index),
        "force_unit": (result_data.get("result_units") or {}).get("force"),
        "location_unit": result_data.get("mesh_unit"),
        "rows": rows,
        "nodal_parity": result_data.get("nodal_parity"),
    }


    if payload["analysis_mode"] == "static":
        axes = _axes_from_result_data(result_data, time_index)
        reference_xyz = _moment_reference_from_result_data(result_data, time_index)
        result_units = complete_result_units(
            result_data.get("result_units"),
            mesh_unit=result_data.get("mesh_unit"),
        )
        source_moment_unit = infer_moment_unit(
            result_units.get("force"),
            result_data.get("mesh_unit"),
        )
        moment_unit = result_units.get("moment") or source_moment_unit
        factor = moment_unit_conversion_factor(source_moment_unit, moment_unit) or 1.0
        for row in rows:
            origin = [float(value) for value in row["origin"][:3]]
            force = [float(value) for value in row["vector"][:3]]
            lever_arm = [origin[index] - reference_xyz[index] for index in range(3)]
            orbital = scale_vector(cross(lever_arm, force), factor)
            explicit = [float(value) for value in row.get("explicit_couple", [0.0] * 3)[:3]]
            combined = add_vectors(orbital, explicit)
            row.update(
                {
                    "lever_arm": lever_arm,
                    "orbital_moment_global": orbital,
                    "explicit_couple_global": explicit,
                    "combined_moment_global": combined,
                    "combined_moment_local": [dot(combined, axis) for axis in axes],
                }
            )
        payload["moment_unit"] = moment_unit
        payload["reference_xyz"] = reference_xyz
    return payload


def result_data_signature_for_time(
    result_data: Optional[Dict[str, Any]],
    time_index: int,
) -> Optional[Dict[str, Any]]:
    if not result_data:
        return None
    signatures = result_data.get("signatures") or []
    if 0 <= int(time_index) < len(signatures):
        return signatures[int(time_index)]
    signature = result_data.get("signature")
    return signature if isinstance(signature, dict) else None


def result_signature_matches_visualization(
    result_data: Optional[Dict[str, Any]],
    visualization_payload: Optional[Dict[str, Any]],
    time_index: int,
) -> bool:
    if not result_data or not visualization_payload:
        return False
    return result_data_signature_for_time(result_data, time_index) == (
        visualization_payload.get("result_signature")
    )


def result_history_signature_matches_visualization(
    result_data: Optional[Dict[str, Any]],
    visualization_payload: Optional[Dict[str, Any]],
    time_index: int,
) -> bool:
    """Match history to exact evidence while animation geometry stays visual-only."""

    if not result_data or not visualization_payload:
        return False
    animation = visualization_payload.get("animation")
    if not isinstance(animation, dict):
        return result_signature_matches_visualization(
            result_data,
            visualization_payload,
            time_index,
        )
    try:
        evidence_time_index = int(animation["evidence_time_index"])
        evidence_result_set_id = int(animation["evidence_result_set_id"])
    except (KeyError, TypeError, ValueError):
        return False
    if int(time_index) != evidence_time_index:
        return False
    evidence_signature = animation.get("evidence_result_signature")
    if not isinstance(evidence_signature, dict):
        return False
    if evidence_signature.get("result_set_id") != evidence_result_set_id:
        return False
    return result_data_signature_for_time(
        result_data,
        evidence_time_index,
    ) == evidence_signature


def _axes_from_result_data(
    result_data: Dict[str, Any],
    time_index: int = 0,
) -> Matrix3:
    frames = result_data.get("resolved_reference_frames") or []
    if 0 <= int(time_index) < len(frames):
        frame_axes = (frames[int(time_index)] or {}).get("resolved_axes") or {}
        if frame_axes:
            return validate_axes(
                {key: frame_axes[key] for key in ("x", "y", "z")}
            )
    signature = result_data.get("signature") or {}
    axes_data = signature.get("coordinate_system_axes") or {}
    return validate_axes(
        {
            "x": axes_data.get("x", [1.0, 0.0, 0.0]),
            "y": axes_data.get("y", [0.0, 1.0, 0.0]),
            "z": axes_data.get("z", [0.0, 0.0, 1.0]),
        }
    )


def _moment_reference_from_result_data(
    result_data: Dict[str, Any],
    time_index: int,
) -> Vector:
    references = result_data.get("moment_reference_xyz_by_set") or []
    if 0 <= int(time_index) < len(references):
        return [float(value) for value in references[int(time_index)][:3]]
    return [
        float(value)
        for value in (result_data.get("moment_reference_xyz") or [0.0, 0.0, 0.0])[:3]
    ]


def _node_coordinates_for_key(
    node_coordinates: Dict[Any, Sequence[float]],
    node_key: Any,
) -> Optional[Vector]:
    node_id = int(node_key)
    xyz = node_coordinates.get(node_key)
    if xyz is None:
        xyz = node_coordinates.get(str(node_key))
    if xyz is None:
        xyz = node_coordinates.get(node_id)
    if xyz is None:
        return None
    return [float(value) for value in xyz[:3]]


def nodal_reconstructed_resultant_rows(
    result_data: Optional[Dict[str, Any]],
) -> Optional[Dict[str, Any]]:
    if not result_data:
        return None
    modal_coordinates = result_data.get("modal_coordinates") or []
    detail_time_index = int(result_data.get("detail_time_index", 0))
    if result_data.get("analysis_mode") == "static" and not modal_coordinates:
        modal_coordinates = [[1.0]]
    nodal_force_modal_coefficients = (
        result_data.get("nodal_force_modal_coefficients") or {}
    )
    nodal_moment_modal_coefficients = (
        result_data.get("nodal_moment_modal_coefficients") or {}
    )
    node_coordinates = result_data.get("selected_node_coordinates") or {}
    if not modal_coordinates or not nodal_force_modal_coefficients or not node_coordinates:
        return None
    max_modes = min(
        int(result_data.get("modes_used") or len(modal_coordinates[0])),
        min(len(row) for row in modal_coordinates),
    )
    if max_modes <= 0:
        return None
    axes = _axes_from_result_data(result_data, detail_time_index)
    reference_xyz = _moment_reference_from_result_data(
        result_data,
        detail_time_index,
    )
    result_units = complete_result_units(
        result_data.get("result_units"),
        mesh_unit=result_data.get("mesh_unit"),
    )
    source_moment_unit = infer_moment_unit(
        result_units.get("force"),
        result_data.get("mesh_unit"),
    )
    moment_unit = result_units.get("moment") or source_moment_unit
    moment_unit_factor = moment_unit_conversion_factor(source_moment_unit, moment_unit)
    if moment_unit_factor is None:
        if result_data.get("analysis_mode") == "static":
            raise RuntimeError(
                "Static nodal moment reconstruction cannot convert the orbital "
                f"moment unit {source_moment_unit!r} to {moment_unit!r}."
            )
        moment_unit_factor = 1.0
    force_coefficients = [[0.0, 0.0, 0.0] for _ in range(max_modes)]
    moment_coefficients = [[0.0, 0.0, 0.0] for _ in range(max_modes)]
    contributed = False
    for node_key, modal_vectors in nodal_force_modal_coefficients.items():
        origin = _node_coordinates_for_key(node_coordinates, node_key)
        if origin is None:
            continue
        lever_arm = [origin[index] - reference_xyz[index] for index in range(3)]
        explicit_vectors = nodal_moment_modal_coefficients.get(node_key)
        if explicit_vectors is None:
            explicit_vectors = nodal_moment_modal_coefficients.get(str(node_key))
        if explicit_vectors is None:
            explicit_vectors = nodal_moment_modal_coefficients.get(int(node_key), [])
        for mode_index in range(min(max_modes, len(modal_vectors))):
            force = [float(value) for value in modal_vectors[mode_index][:3]]
            moment = scale_vector(
                cross(lever_arm, force),
                moment_unit_factor,
            )
            if mode_index < len(explicit_vectors):
                moment = add_vectors(moment, explicit_vectors[mode_index][:3])
            for component_index in range(3):
                force_coefficients[mode_index][component_index] += force[
                    component_index
                ]
                moment_coefficients[mode_index][component_index] += moment[
                    component_index
                ]
            contributed = True
    if not contributed:
        return None
    global_rows: List[List[float]] = []
    local_rows: List[List[float]] = []
    for modal_row in modal_coordinates:
        force_total = [0.0, 0.0, 0.0]
        moment_total = [0.0, 0.0, 0.0]
        for mode_index in range(min(max_modes, len(modal_row))):
            scale = float(modal_row[mode_index])
            for component_index in range(3):
                force_total[component_index] += (
                    scale * force_coefficients[mode_index][component_index]
                )
                moment_total[component_index] += (
                    scale * moment_coefficients[mode_index][component_index]
                )
        global_row = force_total + moment_total
        global_rows.append(global_row)
        local_rows.append(rotate_to_local(global_row, axes))
    return {
        "global_rows": global_rows,
        "local_rows": local_rows,
        "moment_unit": moment_unit,
        "source_moment_unit": source_moment_unit,
        "moment_unit_factor": moment_unit_factor,
    }


def build_nodal_force_time_excel_payload(
    result_data: Optional[Dict[str, Any]],
    time_index: int,
) -> Optional[Dict[str, Any]]:
    base_payload = build_nodal_force_time_export_payload(result_data, time_index)
    if not result_data or base_payload is None:
        return None
    axes = _axes_from_result_data(result_data, time_index)
    reference_xyz = _moment_reference_from_result_data(result_data, time_index)
    force_unit = base_payload.get("force_unit")
    location_unit = base_payload.get("location_unit")
    source_moment_unit = infer_moment_unit(force_unit, location_unit)
    result_units = complete_result_units(
        result_data.get("result_units"),
        mesh_unit=location_unit,
    )
    moment_unit = result_units.get("moment") or source_moment_unit
    moment_unit_factor = moment_unit_conversion_factor(source_moment_unit, moment_unit)
    if moment_unit_factor is None:
        if result_data.get("analysis_mode") == "static":
            raise RuntimeError(
                "Static nodal Excel export cannot convert the orbital moment unit "
                f"{source_moment_unit!r} to {moment_unit!r}."
            )
        moment_unit_factor = 1.0
    resultants_global = result_data.get("resultants_global") or []
    resultants_global_origin = result_data.get("resultants_global_origin") or []
    extracted_global_origin = (
        [float(value) for value in resultants_global_origin[time_index][:6]]
        if time_index < len(resultants_global_origin)
        and len(resultants_global_origin[time_index]) >= 6
        else None
    )
    if extracted_global_origin is None:
        extracted_global_origin = (
            [float(value) for value in resultants_global[time_index][:6]]
            if time_index < len(resultants_global)
            and len(resultants_global[time_index]) >= 6
            else None
        )
    resultants_reference_global = result_data.get("resultants_reference_global") or []
    resultants_local = result_data.get("resultants_local") or []
    extracted_global = (
        [float(value) for value in resultants_reference_global[time_index][:6]]
        if time_index < len(resultants_reference_global)
        and len(resultants_reference_global[time_index]) >= 6
        else None
    )
    if extracted_global is None and extracted_global_origin is not None:
        extracted_global = resultant_about_reference(
            extracted_global_origin,
            reference_xyz,
            cross_product_moment_factor=moment_unit_factor,
        )
    extracted_local = (
        [float(value) for value in resultants_local[time_index][:6]]
        if time_index < len(resultants_local) and len(resultants_local[time_index]) >= 6
        else None
    )
    raw_dpf_global_origin_rows = result_data.get("raw_dpf_resultants_global_origin") or []
    raw_dpf_reference_global_rows = (
        result_data.get("raw_dpf_resultants_reference_global") or []
    )
    raw_dpf_local_rows = result_data.get("raw_dpf_resultants_local") or []

    def row_at(rows: Sequence[Sequence[float]]) -> Optional[List[float]]:
        if time_index >= len(rows) or len(rows[time_index]) < 6:
            return None
        return [float(value) for value in rows[time_index][:6]]

    rows = []
    force_total_global = [0.0, 0.0, 0.0]
    moment_total_global = [0.0, 0.0, 0.0]
    for row in base_payload["rows"]:
        origin = [float(value) for value in row["origin"][:3]]
        force_global = [float(value) for value in row["vector"][:3]]
        lever_arm = [
            origin[index] - reference_xyz[index]
            for index in range(3)
        ]
        orbital_moment_global = scale_vector(
            cross(lever_arm, force_global), moment_unit_factor
        )
        explicit_couple_global = [
            float(value)
            for value in row.get("explicit_couple_global", [0.0, 0.0, 0.0])[:3]
        ]
        moment_global = add_vectors(orbital_moment_global, explicit_couple_global)
        force_local = [dot(force_global, axis) for axis in axes]
        moment_local = [dot(moment_global, axis) for axis in axes]
        for index in range(3):
            force_total_global[index] += force_global[index]
            moment_total_global[index] += moment_global[index]
        rows.append(
            {
                "node_id": int(row["node_id"]),
                "origin": origin,
                "lever_arm": lever_arm,
                "force_global": force_global,
                "force_local": force_local,
                "force_magnitude": vector_magnitude(force_global),
                "orbital_moment_global": orbital_moment_global,
                "explicit_couple_global": explicit_couple_global,
                "combined_moment_global": moment_global,
                "combined_moment_local": moment_local,
                "moment_global": moment_global,
                "moment_local": moment_local,
                "moment_magnitude": vector_magnitude(moment_global),
            }
        )
    force_total_local = [dot(force_total_global, axis) for axis in axes]
    moment_total_local = [dot(moment_total_global, axis) for axis in axes]
    return {
        **base_payload,
        "rows": rows,
        "reference_xyz": reference_xyz,
        "local_axes": axes,
        "force_source_frame": "global",
        "local_component_definition": (
            "local_component = dot(global_vector, local_axis_expressed_in_global_components)"
        ),
        "moment_unit": moment_unit,
        "nodal_moment_source_unit": source_moment_unit,
        "nodal_moment_unit_factor": moment_unit_factor,
        "nodal_totals_global": force_total_global + moment_total_global,
        "nodal_totals_local": force_total_local + moment_total_local,
        "section_geometry": result_data.get("section_geometry") or {},
        "extracted_resultants_global_origin": extracted_global_origin,
        "extracted_resultants_global": extracted_global,
        "extracted_resultants_local": extracted_local,
        "raw_dpf_resultants_global_origin": row_at(raw_dpf_global_origin_rows),
        "raw_dpf_resultants_reference_global": row_at(raw_dpf_reference_global_rows),
        "raw_dpf_resultants_local": row_at(raw_dpf_local_rows),
        "analysis_mode": result_data.get("analysis_mode") or "modal",
        "nodal_parity": result_data.get("nodal_parity"),
    }


def write_nodal_force_time_csv(path: str, payload: Dict[str, Any]) -> None:
    target = Path(path)
    ensure_output_csv_can_be_replaced(str(target))
    rows = payload.get("rows") or []
    force_unit = payload.get("force_unit")
    location_unit = payload.get("location_unit")
    static_mode = payload.get("analysis_mode") == "static"
    moment_unit = payload.get("moment_unit")
    parity = payload.get("nodal_parity") or {}
    temp_path: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            newline="",
            encoding="utf-8",
            dir=str(target.parent),
            prefix=f".{target.name}.",
            suffix=".tmp",
            delete=False,
        ) as stream:
            temp_path = Path(stream.name)
            writer = csv.writer(stream)
            header = [
                    "time_s",
                    "node_id",
                    label_with_unit("x", location_unit),
                    label_with_unit("y", location_unit),
                    label_with_unit("z", location_unit),
                    label_with_unit("fx_global", force_unit),
                    label_with_unit("fy_global", force_unit),
                    label_with_unit("fz_global", force_unit),
                    label_with_unit("f_magnitude", force_unit),
                ]
            if static_mode:
                for descriptor in ("orbital", "explicit_couple", "combined"):
                    header.extend(
                        label_with_unit(f"m{axis}_{descriptor}_global", moment_unit)
                        for axis in ("x", "y", "z")
                    )
                for prefix in (
                    "raw_dpf_m",
                    "mechanical_equivalent_m",
                    "mechanical_equivalent_minus_raw_dpf_m_delta_",
                ):
                    header.extend(
                        label_with_unit(f"{prefix}{axis}", moment_unit)
                        for axis in ("x", "y", "z")
                    )
            writer.writerow(header)
            for row in rows:
                origin = [float(value) for value in row.get("origin", [])[:3]]
                vector = [float(value) for value in row.get("vector", [])[:3]]
                values = (
                    [f"{float(payload.get('time', 0.0)):.16g}", int(row["node_id"])]
                    + [f"{value:.16g}" for value in origin]
                    + [f"{value:.16g}" for value in vector]
                    + [f"{float(row.get('magnitude') or 0.0):.16g}"]
                )
                if static_mode:
                    for key in (
                        "orbital_moment_global",
                        "explicit_couple_global",
                        "combined_moment_global",
                    ):
                        values.extend(
                            f"{float(value):.16g}"
                            for value in row.get(key, [0.0, 0.0, 0.0])[:3]
                        )
                    for key in (
                        "raw_dpf_resultant_reference_global",
                        "mechanical_equivalent_resultant_reference_global",
                        "delta_mechanical_equivalent_minus_raw_dpf",
                    ):
                        aggregate = parity.get(key) or [0.0] * 6
                        values.extend(
                            f"{float(value):.16g}" for value in aggregate[3:6]
                        )
                writer.writerow(values)
        ensure_output_csv_can_be_replaced(str(target))
        temp_path.replace(target)
        temp_path = None
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass


def _style_nodal_force_workbook(workbook: Any) -> None:
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    title_fill = PatternFill("solid", fgColor="E7EEF1")
    header_fill = PatternFill("solid", fgColor="236B5F")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="26343D", bold=True, size=14)
    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = title_fill
                cell.font = title_font
        if worksheet.title == "Nodal Forces" and worksheet.max_row >= 2:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            table = Table(displayName="NodalForces", ref=worksheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)
            for column in ("L", "X"):
                worksheet.conditional_formatting.add(
                    f"{column}2:{column}{worksheet.max_row}",
                    DataBarRule(start_type="min", end_type="max", color="63C384"),
                )
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 10),
                28,
            )


def _append_nodal_force_formula_sheet(workbook: Any, payload: Dict[str, Any]) -> None:
    formulas = workbook.create_sheet("Formulas")
    source_moment_unit = payload.get("nodal_moment_source_unit")
    moment_unit = payload.get("moment_unit")
    moment_factor = float(payload.get("nodal_moment_unit_factor") or 1.0)
    static_mode = payload.get("analysis_mode") == "static"
    formulas.append(
        [
            "DPF Static Structural Section Resultants - Formulas"
            if static_mode
            else "MSUP MCF Section Resultants - Formulas"
        ]
    )
    formulas.append([])
    formulas.append(["Quantity", "Formula", "Notes"])
    formulas.append(
        [
            "fx/fy/fz_global",
            (
                "DPF forces_on_nodes at the selected result set"
                if static_mode
                else "sum over modes j: q_j(time) * modal_force_coefficient[node_id,j,component]"
            ),
            (
                "Static Structural uses one cumulative DPF result set."
                if static_mode
                else "q_j comes from the current MCF time row; modal coefficients come from DPF force_summation nodal forces."
            ),
        ]
    )
    formulas.append(
        [
            "f_magnitude",
            "sqrt(fx_global^2 + fy_global^2 + fz_global^2)",
            "",
        ]
    )
    formulas.append(
        [
            "rx/ry/rz_from_reference",
            "Nodal Forces!F:H = node_xyz - moment_reference_xyz",
            "The moment reference is the selected coordinate-system origin, probe centroid, or selected-node centroid.",
        ]
    )
    formulas.append(
        [
            "fx/fy/fz_local",
            "Nodal Forces!M:O = dot(force_global, local_axis)",
            "Local axes rotate components only; the origin shift is handled by the lever arm.",
        ]
    )
    formulas.append(
        [
            "mx/my/mz_about_reference_global",
            "Nodal Forces!Q:S = moment_unit_factor * cross(lever_arm, force_global)",
            (
                f"moment_unit_factor={moment_factor:.16g}; converts "
                f"{source_moment_unit or 'force*length'} to {moment_unit or 'workbook moment unit'}."
            ),
        ]
    )
    formulas.append(
        [
            "m_magnitude",
            "sqrt(mx_global^2 + my_global^2 + mz_global^2)",
            "",
        ]
    )
    formulas.append(
        [
            "mx/my/mz_about_reference_local",
            "Nodal Forces!U:W = dot(moment_global_about_reference, local_axis)",
            "",
        ]
    )
    if static_mode:
        formulas.append(
            [
                "explicit nodal couple",
                "Nodal Forces!Y:AA = DPF moments_on_nodes (zero when the output is empty)",
                "Beam and shell elements can contribute an explicit nodal couple.",
            ]
        )
        formulas.append(
            [
                "combined nodal moment",
                "Nodal Forces!AB:AD = orbital r x F + explicit nodal couple",
                "This Mechanical-equivalent construction-surface moment is the primary static result.",
            ]
        )
    formulas.append(
        [
            "Summary nodal totals",
            "Summary!B:C = sum each Nodal Forces component column; Summary!G = extracted local - nodal local",
            "These are the values in the Summary columns named Nodal sheet sum.",
        ]
    )
    formulas.append(
        [
            (
                "Summary Mechanical-equivalent totals"
                if static_mode
                else "Summary DPF extracted totals"
            ),
            (
                "DPF force_accumulation plus the nodal sum of deformed r x F and explicit couples"
                if static_mode
                else "sum over modes j: q_j(time) * DPF force_summation resultant_coefficient[j]"
            ),
            (
                "Static Summary!D:F are the primary construction-surface values; raw DPF moment_accumulation is retained in H:J."
                if static_mode
                else "Summary!D is global origin; Summary!E/F are shifted to the selected moment reference before local projection."
            ),
        ]
    )


def write_nodal_force_time_excel(path: str, payload: Dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise RuntimeError(
            "Excel export requires openpyxl. Install the project with the excel, "
            "viewer, all, or dev extra."
        ) from exc

    target = Path(path)
    ensure_output_csv_can_be_replaced(str(target))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    nodal = workbook.create_sheet("Nodal Forces")
    force_unit = payload.get("force_unit")
    moment_unit = payload.get("moment_unit")
    location_unit = payload.get("location_unit")
    section_geometry = payload.get("section_geometry") or {}
    nodal_rows_payload = payload.get("rows") or []
    nodal_last_row = len(nodal_rows_payload) + 1
    axes = payload.get("local_axes") or [
        [1.0, 0.0, 0.0],
        [0.0, 1.0, 0.0],
        [0.0, 0.0, 1.0],
    ]
    axes = [[float(value) for value in axis[:3]] for axis in axes[:3]]
    moment_factor = float(payload.get("nodal_moment_unit_factor") or 1.0)
    static_mode = payload.get("analysis_mode") == "static"

    def formula_number(value: float) -> str:
        return f"{float(value):.16g}"

    def nodal_sum_formula(column: str) -> str:
        if nodal_last_row < 2:
            return "=0"
        return f"=SUM('Nodal Forces'!{column}2:{column}{nodal_last_row})"

    def dot_formula(columns: Sequence[str], row_index: int, axis: Sequence[float]) -> str:
        terms = [
            f"{columns[index]}{row_index}*{formula_number(axis[index])}"
            for index in range(3)
        ]
        return "=" + "+".join(terms)

    summary_rows = [
        [
            (
                "DPF Static Structural Section Resultants - Nodal Force, Local Force, and Local Moment Export"
                if static_mode
                else "MSUP MCF Section Resultants - Nodal Force, Local Force, and Local Moment Export"
            )
        ],
        ["Time index", int(payload.get("time_index", 0)) + 1],
        ["Time [s]", float(payload.get("time", 0.0))],
        ["Force unit", force_unit or ""],
        ["Moment unit", moment_unit or ""],
        ["Location unit", location_unit or ""],
        [
            "Local moment summary source",
            (
                "Nodal Forces sheet sum of local (r x F + explicit couple)"
                if static_mode
                else "Nodal Forces sheet sum of local r x F"
            ),
        ],
        ["Moment reference X", payload.get("reference_xyz", [0.0, 0.0, 0.0])[0]],
        ["Moment reference Y", payload.get("reference_xyz", [0.0, 0.0, 0.0])[1]],
        ["Moment reference Z", payload.get("reference_xyz", [0.0, 0.0, 0.0])[2]],
        ["Section width axis", section_geometry.get("width_axis") or ""],
        ["Section width [mm]", section_geometry.get("width_mm")],
        ["Section height axis", section_geometry.get("height_axis") or ""],
        ["Section height [mm]", section_geometry.get("height_mm")],
        ["Projected section area [mm^2]", section_geometry.get("area_mm2")],
        ["Section geometry point source", section_geometry.get("point_source") or ""],
        ["DPF nodal force source frame", payload.get("force_source_frame") or "global"],
        [
            "Local component definition",
            payload.get("local_component_definition")
            or "local_component = dot(global_vector, local_axis_expressed_in_global_components)",
        ],
        ["Local X axis in global components", *axes[0]],
        ["Local Y axis in global components", *axes[1]],
        ["Local Z axis in global components", *axes[2]],
        [],
        (
            [
                "Component",
                "Nodal sheet sum, global about reference",
                "Nodal sheet sum, local about reference",
                "Mechanical-equivalent resultant, global origin",
                "Mechanical-equivalent resultant, global about reference",
                "Mechanical-equivalent resultant, local about reference",
                "Mechanical-equivalent local minus nodal local",
                "Raw DPF resultant, global origin",
                "Raw DPF resultant, global about reference",
                "Raw DPF resultant, local about reference",
                "Mechanical-equivalent global about reference minus raw DPF",
            ]
            if static_mode
            else [
                "Component",
                "Nodal sheet sum, global about reference",
                "Nodal sheet sum, local about reference",
                "DPF moment_accumulation, global origin",
                "DPF moment_accumulation, global about reference",
                "DPF moment_accumulation, local about reference",
                "DPF local minus nodal local",
            ]
        ),
    ]
    labels = ("Fx", "Fy", "Fz", "Mx", "My", "Mz")
    extracted_global_origin = payload.get("extracted_resultants_global_origin") or []
    extracted_global = payload.get("extracted_resultants_global") or []
    extracted_local = payload.get("extracted_resultants_local") or []
    raw_dpf_global_origin = payload.get("raw_dpf_resultants_global_origin") or []
    raw_dpf_global = payload.get("raw_dpf_resultants_reference_global") or []
    raw_dpf_local = payload.get("raw_dpf_resultants_local") or []
    summary_component_start_row = len(summary_rows) + 1
    nodal_global_columns = (
        ("I", "J", "K", "AB", "AC", "AD")
        if static_mode
        else ("I", "J", "K", "Q", "R", "S")
    )
    nodal_local_columns = (
        ("M", "N", "O", "AE", "AF", "AG")
        if static_mode
        else ("M", "N", "O", "U", "V", "W")
    )
    for index, label in enumerate(labels):
        component_row = summary_component_start_row + index
        row = [
            label,
            nodal_sum_formula(nodal_global_columns[index]),
            nodal_sum_formula(nodal_local_columns[index]),
            (
                extracted_global_origin[index]
                if index < len(extracted_global_origin)
                else None
            ),
            extracted_global[index] if index < len(extracted_global) else None,
            extracted_local[index] if index < len(extracted_local) else None,
            f"=F{component_row}-C{component_row}",
        ]
        if static_mode:
            row.extend(
                [
                    (
                        raw_dpf_global_origin[index]
                        if index < len(raw_dpf_global_origin)
                        else None
                    ),
                    raw_dpf_global[index] if index < len(raw_dpf_global) else None,
                    raw_dpf_local[index] if index < len(raw_dpf_local) else None,
                    f"=E{component_row}-I{component_row}",
                ]
            )
        summary_rows.append(row)
    for row in summary_rows:
        summary.append(row)

    nodal_header = [
            "node_id",
            "time_s",
            label_with_unit("x", location_unit),
            label_with_unit("y", location_unit),
            label_with_unit("z", location_unit),
            label_with_unit("rx_from_reference", location_unit),
            label_with_unit("ry_from_reference", location_unit),
            label_with_unit("rz_from_reference", location_unit),
            label_with_unit("fx_global", force_unit),
            label_with_unit("fy_global", force_unit),
            label_with_unit("fz_global", force_unit),
            label_with_unit("f_magnitude", force_unit),
            label_with_unit("fx_local", force_unit),
            label_with_unit("fy_local", force_unit),
            label_with_unit("fz_local", force_unit),
            label_with_unit("f_local_magnitude", force_unit),
            label_with_unit(
                "mx_orbital_about_reference_global" if static_mode else "mx_about_reference_global",
                moment_unit,
            ),
            label_with_unit(
                "my_orbital_about_reference_global" if static_mode else "my_about_reference_global",
                moment_unit,
            ),
            label_with_unit(
                "mz_orbital_about_reference_global" if static_mode else "mz_about_reference_global",
                moment_unit,
            ),
            label_with_unit("m_magnitude", moment_unit),
            label_with_unit("mx_about_reference_local", moment_unit),
            label_with_unit("my_about_reference_local", moment_unit),
            label_with_unit("mz_about_reference_local", moment_unit),
            label_with_unit("m_local_magnitude", moment_unit),
        ]
    if static_mode:
        nodal_header.extend(
            label_with_unit(label, moment_unit)
            for label in (
                "mx_explicit_couple_global",
                "my_explicit_couple_global",
                "mz_explicit_couple_global",
                "mx_combined_about_reference_global",
                "my_combined_about_reference_global",
                "mz_combined_about_reference_global",
                "mx_combined_about_reference_local",
                "my_combined_about_reference_local",
                "mz_combined_about_reference_local",
                "m_combined_local_magnitude",
            )
        )
    nodal.append(nodal_header)
    for excel_row, row in enumerate(nodal_rows_payload, start=2):
        nodal_row = [
                int(row["node_id"]),
                float(payload.get("time", 0.0)),
                *[float(value) for value in row.get("origin", [])[:3]],
                f"=C{excel_row}-Summary!$B$8",
                f"=D{excel_row}-Summary!$B$9",
                f"=E{excel_row}-Summary!$B$10",
                *[float(value) for value in row.get("force_global", [])[:3]],
                f"=SQRT(SUMSQ(I{excel_row}:K{excel_row}))",
                dot_formula(("I", "J", "K"), excel_row, axes[0]),
                dot_formula(("I", "J", "K"), excel_row, axes[1]),
                dot_formula(("I", "J", "K"), excel_row, axes[2]),
                f"=SQRT(SUMSQ(M{excel_row}:O{excel_row}))",
                (
                    f"=(G{excel_row}*K{excel_row}-H{excel_row}*J{excel_row})"
                    f"*{formula_number(moment_factor)}"
                ),
                (
                    f"=(H{excel_row}*I{excel_row}-F{excel_row}*K{excel_row})"
                    f"*{formula_number(moment_factor)}"
                ),
                (
                    f"=(F{excel_row}*J{excel_row}-G{excel_row}*I{excel_row})"
                    f"*{formula_number(moment_factor)}"
                ),
                f"=SQRT(SUMSQ(Q{excel_row}:S{excel_row}))",
                dot_formula(("Q", "R", "S"), excel_row, axes[0]),
                dot_formula(("Q", "R", "S"), excel_row, axes[1]),
                dot_formula(("Q", "R", "S"), excel_row, axes[2]),
                f"=SQRT(SUMSQ(U{excel_row}:W{excel_row}))",
            ]
        if static_mode:
            explicit = [
                float(value)
                for value in row.get("explicit_couple_global", [0.0, 0.0, 0.0])[:3]
            ]
            nodal_row.extend(
                explicit
                + [f"=Q{excel_row}+Y{excel_row}", f"=R{excel_row}+Z{excel_row}", f"=S{excel_row}+AA{excel_row}"]
                + [
                    dot_formula(("AB", "AC", "AD"), excel_row, axes[0]),
                    dot_formula(("AB", "AC", "AD"), excel_row, axes[1]),
                    dot_formula(("AB", "AC", "AD"), excel_row, axes[2]),
                    f"=SQRT(SUMSQ(AE{excel_row}:AG{excel_row}))",
                ]
            )
        nodal.append(nodal_row)
    _append_nodal_force_formula_sheet(workbook, payload)
    _style_nodal_force_workbook(workbook)
    temp_path: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(
            "wb",
            dir=str(target.parent),
            prefix=f".{target.name}.",
            suffix=".tmp.xlsx",
            delete=False,
        ) as stream:
            temp_path = Path(stream.name)
        workbook.save(str(temp_path))
        ensure_output_csv_can_be_replaced(str(target))
        temp_path.replace(target)
        temp_path = None
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass


def scaled_vector_items(
    items: Sequence[Dict[str, Any]],
    *,
    scene_extent: float,
    length_fraction: float,
) -> List[Dict[str, Any]]:
    max_magnitude = max((float(item.get("magnitude") or 0.0) for item in items), default=0.0)
    if max_magnitude <= 0.0:
        return []
    bounded_extent = max(float(scene_extent), VISUALIZATION_SCENE_MIN_EXTENT)
    scale = bounded_extent * float(length_fraction) / max_magnitude
    scaled: List[Dict[str, Any]] = []
    for item in items:
        magnitude = float(item.get("magnitude") or 0.0)
        if magnitude <= 0.0 or "origin" not in item:
            continue
        copied = dict(item)
        copied["display_vector"] = scale_vector(copied["vector"], scale)
        scaled.append(copied)
    return scaled


def nodal_vector_glyph_items(nodal_vectors: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    glyph_items: List[Dict[str, Any]] = []
    for item in nodal_vectors:
        if not isinstance(item, dict):
            continue
        if "origin" not in item or "display_vector" not in item:
            continue
        display_vector = [float(value) for value in item["display_vector"][:3]]
        display_length = vector_magnitude(display_vector)
        if display_length <= 0.0:
            continue
        glyph_items.append(
            {
                "origin": [float(value) for value in item["origin"][:3]],
                "display_vector": display_vector,
                "display_length": display_length,
                "magnitude": float(item.get("magnitude") or 0.0),
            }
        )
    return glyph_items


def build_result_overlay_payload(
    visualization_payload: Dict[str, Any],
    result_data: Optional[Dict[str, Any]],
    result_mode: str,
    time_index: int,
) -> Optional[Dict[str, Any]]:
    if not result_data:
        return None
    if not result_signature_matches_visualization(
        result_data,
        visualization_payload,
        time_index,
    ):
        return None

    times = [float(value) for value in result_data.get("times", [])]
    modal_coordinates = result_data.get("modal_coordinates") or []
    resultants_global = result_data.get("resultants_global") or []
    resultants_reference_global = (
        result_data.get("resultants_reference_global") or resultants_global
    )
    static_mode = result_data.get("analysis_mode") == "static"
    if not times or not resultants_global or (not static_mode and not modal_coordinates):
        return None
    if time_index < 0 or time_index >= len(times):
        raise IndexError(f"Result time index is out of range: {time_index}")

    mode = "moment" if str(result_mode).lower().startswith("moment") else "force"
    scene_extent = visualization_scene_extent(visualization_payload)
    total_row = [float(value) for value in resultants_global[time_index]]
    reference_total_row = [
        float(value) for value in resultants_reference_global[time_index]
    ]
    reference_anchor = _moment_reference_from_result_data(result_data, time_index)
    if mode == "moment":
        nodal_resultants = nodal_reconstructed_resultant_rows(result_data)
        if result_data.get("analysis_mode") == "static":
            total_vector = reference_total_row[3:6]
            label = "Moment Reaction (Mechanical-equivalent section resultant)"
        elif (
            nodal_resultants is not None
            and time_index < len(nodal_resultants["global_rows"])
        ):
            total_vector = nodal_resultants["global_rows"][time_index][3:6]
            label = "Moment Reaction (nodal r x F)"
        else:
            total_vector = reference_total_row[3:6]
            label = "Moment Reaction"
        anchor = reference_anchor
        nodal_vectors: List[Dict[str, Any]] = []
        total_color = VISUALIZATION_TOTAL_MOMENT_COLOR
    else:
        total_vector = total_row[:3]
        label = "Force Reaction"
        anchor = reference_anchor
        detail_index = int(result_data.get("detail_time_index", time_index))
        if static_mode and detail_index != time_index:
            raw_nodal_vectors = []
        else:
            modal_row = [1.0] if static_mode else modal_coordinates[time_index]
            raw_nodal_vectors = combine_modal_nodal_force_vectors(
                result_data.get("nodal_force_modal_coefficients") or {},
                modal_row,
                node_coordinates=result_data.get("selected_node_coordinates") or {},
                modes_used=(
                    1
                    if static_mode
                    else int(result_data.get("modes_used") or len(modal_row))
                ),
            )
        nodal_vectors = scaled_vector_items(
            raw_nodal_vectors,
            scene_extent=scene_extent,
            length_fraction=VISUALIZATION_NODAL_VECTOR_LENGTH_FRACTION,
        )
        total_color = VISUALIZATION_TOTAL_FORCE_COLOR

    unit = (result_data.get("result_units") or {}).get(mode)
    total_magnitude = vector_magnitude(total_vector)
    total_origin = [float(value) for value in (anchor or [0.0, 0.0, 0.0])[:3]]
    total_display_vector = [0.0, 0.0, 0.0]
    if total_magnitude > 0.0:
        display_length = total_vector_display_length(
            visualization_payload,
            total_origin,
            total_vector,
            fallback_extent=scene_extent,
        )
        total_display_vector = scale_vector(
            total_vector,
            display_length / total_magnitude,
        )

    components = ", ".join(f"{value:.6g}" for value in total_vector)
    unit_suffix = f" {unit}" if unit else ""
    selected_sets = result_data.get("selected_result_sets") or []
    set_label = ""
    if static_mode and time_index < len(selected_sets):
        set_label = str(selected_sets[time_index].get("label") or "")
    return {
        "mode": mode,
        "label": label,
        "unit": unit,
        "time_index": int(time_index),
        "time": times[time_index],
        "nodal_vectors": nodal_vectors,
        "total_vector": {
            "origin": total_origin,
            "vector": total_vector,
            "display_vector": total_display_vector,
            "magnitude": total_magnitude,
            "color": total_color,
        },
        "text": (
            f"{label} | {set_label or f't={times[time_index]:.6g} s'} | "
            f"[{components}]{unit_suffix} | |V|={total_magnitude:.6g}{unit_suffix}"
        ),
    }


def visualization_payload_with_result_overlay(
    visualization_payload: Dict[str, Any],
    result_data: Optional[Dict[str, Any]],
    result_mode: str,
    time_index: int,
) -> Dict[str, Any]:
    payload = dict(visualization_payload)
    overlay = build_result_overlay_payload(
        visualization_payload,
        result_data,
        result_mode,
        time_index,
    )
    if overlay is not None:
        payload["result_overlay"] = overlay
    else:
        payload.pop("result_overlay", None)
    return payload


def _result_history_components(
    rows: Sequence[Sequence[float]],
    start_index: int,
) -> Dict[str, List[float]]:
    components = {
        "x": [],
        "y": [],
        "z": [],
        "total": [],
    }
    for row in rows:
        vector = [float(row[start_index + offset]) for offset in range(3)]
        components["x"].append(vector[0])
        components["y"].append(vector[1])
        components["z"].append(vector[2])
        components["total"].append(vector_magnitude(vector))
    return components


def result_history_series_signature(
    result_data: Optional[Dict[str, Any]],
) -> Any:
    if not result_data:
        return None
    signatures = result_data.get("signatures") or []
    if signatures:
        return {
            "analysis_mode": result_data.get("analysis_mode"),
            "result_set_ids": [
                int(value) for value in result_data.get("result_set_ids", [])
            ],
            "signatures": list(signatures),
        }
    return result_data.get("signature")


def result_history_x_axis(
    result_data: Optional[Dict[str, Any]],
    row_count: int,
) -> Dict[str, Any]:
    count = max(0, int(row_count))
    data = result_data or {}
    result_values = [float(value) for value in data.get("times", [])[:count]]
    if data.get("analysis_mode") != "static":
        return {
            "kind": "time",
            "values": result_values,
            "label": "Time [s]",
            "unit": "s",
            "point_labels": [f"t={value:.6g} s" for value in result_values],
            "result_set_ids": [],
            "result_values": result_values,
        }

    selected_sets = list(data.get("selected_result_sets") or [])[:count]
    result_set_ids = [int(value) for value in data.get("result_set_ids", [])[:count]]
    if len(result_set_ids) < count:
        result_set_ids = [
            int(selected_sets[index].get("id", index + 1))
            if index < len(selected_sets)
            else index + 1
            for index in range(count)
        ]
    unit = next(
        (
            str(item.get("unit") or "").strip()
            for item in selected_sets
            if str(item.get("unit") or "").strip()
        ),
        "",
    )
    unit_suffix = f" {unit}" if unit else ""
    point_labels: List[str] = []
    for index, (set_id, value) in enumerate(zip(result_set_ids, result_values)):
        supplied = selected_sets[index].get("label") if index < len(selected_sets) else None
        point_labels.append(
            str(supplied)
            if supplied
            else f"Set {set_id} — {value:.6g}{unit_suffix}"
        )
    strictly_increasing = all(math.isfinite(value) for value in result_values) and all(
        result_values[index] > result_values[index - 1]
        for index in range(1, len(result_values))
    )
    if count == len(result_values) and strictly_increasing:
        time_units = {"s", "sec", "second", "seconds"}
        label = (
            f"Time [{unit}]"
            if unit.lower() in time_units
            else (f"Result value [{unit}]" if unit else "Result value")
        )
        kind = "result_value"
        values: List[float] = result_values
    else:
        label = "Cumulative set ID"
        kind = "result_set_id"
        values = [float(value) for value in result_set_ids]
    return {
        "kind": kind,
        "values": values,
        "label": label,
        "unit": unit or None,
        "point_labels": point_labels,
        "result_set_ids": result_set_ids,
        "result_values": result_values,
    }


def result_component_history(
    result_data: Optional[Dict[str, Any]],
    frame: str = "global",
    *,
    moment_display_unit: Optional[str] = None,
) -> Dict[str, Any]:
    frame_key = "local" if str(frame).strip().lower() == "local" else "global"
    result_units = (result_data or {}).get("result_units") or {}
    source_moment_unit = result_units.get("moment")
    moment_unit = source_moment_unit
    moment_unit_factor = moment_unit_conversion_factor(
        source_moment_unit,
        moment_display_unit,
    )
    if moment_unit_factor is not None:
        target_text = str(moment_display_unit or "").strip()
        if target_text and target_text.lower() != "source":
            moment_unit = target_text
    else:
        moment_unit_factor = 1.0
    history: Dict[str, Any] = {
        "analysis_mode": (result_data or {}).get("analysis_mode") or "modal",
        "frame": frame_key,
        "times": [],
        "x_axis": result_history_x_axis(result_data, 0),
        "source_result_units": {
            "force": result_units.get("force"),
            "moment": source_moment_unit,
        },
        "result_units": {
            "force": result_units.get("force"),
            "moment": moment_unit,
        },
        "force": {
            "unit": result_units.get("force"),
            "labels": {"x": "Fx", "y": "Fy", "z": "Fz", "total": "Ftotal"},
            "components": {"x": [], "y": [], "z": [], "total": []},
        },
        "moment": {
            "unit": moment_unit,
            "labels": {"x": "Mx", "y": "My", "z": "Mz", "total": "Mtotal"},
            "components": {"x": [], "y": [], "z": [], "total": []},
        },
    }
    if not result_data:
        return history

    rows_key = "resultants_local" if frame_key == "local" else "resultants_global"
    times = [float(value) for value in result_data.get("times", [])]
    rows = [
        [float(value) for value in row[:6]]
        for row in result_data.get(rows_key, [])
        if len(row) >= 6
    ]
    nodal_resultants = (
        None
        if result_data.get("analysis_mode") == "static"
        else nodal_reconstructed_resultant_rows(result_data)
    )
    nodal_rows: List[List[float]] = []
    parity = result_data.get("nodal_parity") or {}
    use_nodal_moments = not (
        result_data.get("analysis_mode") == "static"
        and parity.get("moment_matches") is not True
    )
    if frame_key == "local" and nodal_resultants is not None and use_nodal_moments:
        nodal_rows = [
            [float(value) for value in row[:6]]
            for row in nodal_resultants.get("local_rows", [])
            if len(row) >= 6
        ]
    row_count = min(len(times), len(rows))
    if row_count <= 0:
        return history

    rows = rows[:row_count]
    history["times"] = times[:row_count]
    history["x_axis"] = result_history_x_axis(result_data, row_count)
    history["force"]["components"] = _result_history_components(rows, 0)
    moment_rows = nodal_rows[:row_count] if len(nodal_rows) >= row_count else rows
    history["moment"]["components"] = _result_history_components(moment_rows, 3)
    if moment_unit_factor != 1.0:
        history["moment"]["components"] = convert_result_components(
            history["moment"]["components"],
            moment_unit_factor,
        )
    return history


def result_history_default_frame(visualization_payload: Optional[Dict[str, Any]]) -> str:
    if not visualization_payload:
        return "global"
    reference = visualization_payload.get("moment_reference_xyz")
    if reference is None:
        return "global"
    try:
        return "local" if vector_magnitude(reference) > 1.0e-12 else "global"
    except Exception:
        return "global"


def build_result_history_plot_payload(
    visualization_payload: Dict[str, Any],
    result_data: Optional[Dict[str, Any]],
    frame: str = "global",
    *,
    selected_time_index: int = 0,
    moment_display_unit: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    if not result_data:
        return None
    if not result_history_signature_matches_visualization(
        result_data,
        visualization_payload,
        selected_time_index,
    ):
        return None
    history = result_component_history(
        result_data,
        frame,
        moment_display_unit=moment_display_unit,
    )
    if not history.get("times"):
        return None
    history["signature"] = result_data_signature_for_time(
        result_data,
        selected_time_index,
    )
    history["series_signature"] = result_history_series_signature(result_data)
    return history


def nearest_history_point(
    candidates: Sequence[Dict[str, Any]],
    x_pixel: float,
    y_pixel: float,
    *,
    max_distance_pixels: float = 14.0,
) -> Optional[Dict[str, Any]]:
    best: Optional[Dict[str, Any]] = None
    best_distance = float(max_distance_pixels)
    for candidate in candidates:
        try:
            distance = math.hypot(
                float(candidate["x_pixel"]) - float(x_pixel),
                float(candidate["y_pixel"]) - float(y_pixel),
            )
        except (KeyError, TypeError, ValueError):
            continue
        if distance <= best_distance:
            best_distance = distance
            best = dict(candidate)
            best["distance_pixels"] = distance
    return best


def visualization_payload_exceeds_large_scene_limit(payload: Dict[str, Any]) -> bool:
    counts = payload.get("counts") or {}
    return (
        int(counts.get("raw_element_count") or 0) > VISUALIZATION_LARGE_ELEMENT_COUNT
        or int(counts.get("mesh_node_count") or 0) > VISUALIZATION_LARGE_NODE_COUNT
    )


def visualization_large_scene_message(payload: Dict[str, Any]) -> str:
    counts = payload.get("counts") or {}
    return (
        "This visualization contains "
        f"{int(counts.get('raw_element_count') or 0)} selected element(s) and "
        f"{int(counts.get('mesh_node_count') or 0)} mesh node(s). "
        "Rendering may take time or make the window temporarily unresponsive."
    )


def _current_reference_frame_from_snapshot(
    mesh_config: SectionConfig,
    snapshot: VisualizationMeshSnapshot,
    reference_mesh: Any,
    reference_side_filter_info: Dict[str, Any],
) -> ResolvedReferenceFrame:
    cached_frame = dict(snapshot.resolved_reference_frame or {})

    if mesh_config.reference_frame_motion != "follow-geometry":
        if not cached_frame:
            raise RuntimeError(
                "Static visualization snapshot has no reference-frame evidence."
            )
        return rebase_reference_frame(
            cached_frame,
            mesh_config.coordinate_system_origin,
            mesh_config.coordinate_system_axes,
            warning_ratio=mesh_config.reference_frame_fit_warning_ratio,
        )

    attachment_selection = (
        mesh_config.reference_frame_attachment_selection or ""
    ).strip()
    if attachment_selection:
        if not cached_frame:
            raise RuntimeError(
                "Static visualization snapshot has no explicit attachment frame evidence."
            )
        return rebase_reference_frame(
            cached_frame,
            mesh_config.coordinate_system_origin,
            mesh_config.coordinate_system_axes,
            warning_ratio=mesh_config.reference_frame_fit_warning_ratio,
        )

    tracking_node_ids = unique_element_node_ids(
        reference_mesh,
        reference_side_filter_info.get("cut_element_ids", []),
    )
    if len(tracking_node_ids) < 3:
        raise RuntimeError(
            follow_geometry_reference_preview_warning(len(tracking_node_ids))
        )
    cached_tracking_ids = sorted(
        set(int(value) for value in cached_frame.get("tracking_node_ids", []))
    )
    if cached_frame and tracking_node_ids == cached_tracking_ids:
        return rebase_reference_frame(
            cached_frame,
            mesh_config.coordinate_system_origin,
            mesh_config.coordinate_system_axes,
            warning_ratio=mesh_config.reference_frame_fit_warning_ratio,
        )

    missing_coordinates = [
        node_id for node_id in tracking_node_ids if node_id not in snapshot.node_coordinates
    ]
    missing_displacements = [
        node_id for node_id in tracking_node_ids if node_id not in snapshot.nodal_displacements
    ]
    if missing_coordinates or missing_displacements:
        raise RuntimeError(
            "Cached static visualization data cannot resolve the current following-frame "
            "tracking neighborhood: "
            f"missing_coordinates={missing_coordinates[:10]}, "
            f"missing_displacements={missing_displacements[:10]}."
        )
    reference_coordinates = {
        node_id: snapshot.node_coordinates[node_id] for node_id in tracking_node_ids
    }
    current_coordinates = {
        node_id: add_vectors(
            snapshot.node_coordinates[node_id],
            snapshot.nodal_displacements[node_id],
        )
        for node_id in tracking_node_ids
    }
    return fit_geometry_following_frame(
        reference_coordinates,
        current_coordinates,
        mesh_config.coordinate_system_origin,
        mesh_config.coordinate_system_axes,
        result_set_id=int(
            cached_frame.get("result_set_id")
            or (snapshot.deformation or {}).get("result_set_id")
            or mesh_config.result_set_id
        ),
        result_value=cached_frame.get("result_value") if cached_frame else None,
        result_unit=cached_frame.get("result_unit") if cached_frame else None,
        attachment_source="local_section_cut_neighborhood",
        attachment_selection=str(
            (cached_frame.get("attachment_selection") if cached_frame else None)
            or snapshot.element_name
        ),
        warning_ratio=mesh_config.reference_frame_fit_warning_ratio,
    )


def _section_visualization_payload_from_snapshot(
    dpf: Any,
    config: SectionConfig,
    axes: Matrix3,
    snapshot: VisualizationMeshSnapshot,
    *,
    phase_start: float,
    log: LogFn,
) -> Dict[str, Any]:
    reference_mesh = _SnapshotMesh(snapshot)
    mesh_config, origin_unit_conversion = config_with_origin_in_mesh_units(
        config,
        reference_mesh,
        log=log,
    )
    result_signature = result_visualization_signature(
        mesh_config,
        mesh_unit=getattr(reference_mesh, "unit", None),
    )
    element_scoping = make_scoping(dpf, snapshot.element_ids, dpf.locations.elemental)
    mesh = reference_mesh
    deformation = dict(snapshot.deformation or {})
    reference_side_filter_info: Optional[Dict[str, Any]] = None
    resolved_frame: Dict[str, Any] = {}
    resolved_config = mesh_config
    resolved_axes = axes
    if config.analysis_mode == "static":
        if not snapshot.nodal_displacements or not deformation:
            raise RuntimeError(
                "Static visualization requires selected-set displacement; reference-geometry "
                "fallback is disabled."
            )
        _, _, reference_side_filter_info = construction_surface_scoping(
            dpf,
            None,
            element_scoping,
            mesh_config,
            axes,
            mesh=reference_mesh,
            allow_empty_cut=True,
        )
        implicit_tracking_node_ids: List[int] = []
        if (
            mesh_config.reference_frame_motion == "follow-geometry"
            and not mesh_config.reference_frame_attachment_selection.strip()
        ):
            implicit_tracking_node_ids = unique_element_node_ids(
                reference_mesh,
                reference_side_filter_info.get("cut_element_ids", []),
            )
        if (
            mesh_config.reference_frame_motion == "follow-geometry"
            and not mesh_config.reference_frame_attachment_selection.strip()
            and len(implicit_tracking_node_ids) < 3
        ):
            warning = follow_geometry_reference_preview_warning(
                len(implicit_tracking_node_ids)
            )
            payload = section_visualization_payload_from_mesh(
                mesh=reference_mesh,
                config=mesh_config,
                axes=axes,
                element_name=snapshot.element_name,
                available_named=snapshot.available_named_selections,
                side_filter_info=reference_side_filter_info,
                dpf_version=getattr(dpf, "__version__", None),
                origin_unit_conversion=origin_unit_conversion,
                deformation=None,
                resolved_reference_frame=None,
            )
            payload["requested_result_set_id"] = mesh_config.result_set_id
            payload["result_set_id"] = None
            payload["result_signature"] = None
            payload["preview_only_reason"] = FOLLOW_GEOMETRY_REFERENCE_PREVIEW_REASON
            payload.setdefault("warnings", []).append(warning)
            payload["elapsed_seconds"] = perf_counter() - phase_start
            emit(log, warning)
            return payload
        current_frame = _current_reference_frame_from_snapshot(
            mesh_config,
            snapshot,
            reference_mesh,
            reference_side_filter_info,
        )
        resolved_frame = asdict(current_frame)
        resolved_config = config_from_mapping(asdict(mesh_config))
        resolved_config.coordinate_system_origin = list(current_frame.resolved_origin)
        resolved_config.coordinate_system_axes = {
            key: list(current_frame.resolved_axes[key]) for key in ("x", "y", "z")
        }
        resolved_axes = validate_axes(resolved_config.coordinate_system_axes)
        mesh = mesh_with_nodal_displacements(
            reference_mesh,
            snapshot.nodal_displacements,
            deformation,
        )

    follows_geometry = (
        resolved_config.reference_frame_motion == "follow-geometry"
    )
    scoping_mesh = mesh
    if config.analysis_mode == "static" and not follows_geometry:
        scoping_mesh = reference_mesh
    filtered_elem_scope, node_scope, side_filter_info = construction_surface_scoping(
        dpf,
        None,
        element_scoping,
        resolved_config,
        resolved_axes,
        mesh=scoping_mesh,
        allow_empty_cut=True,
    )
    if config.analysis_mode == "static":
        if follows_geometry:
            side_filter_info["scoping_geometry_state"] = "deformed"
            side_filter_info["coordinate_geometry_state"] = "deformed"
        else:
            side_filter_info = side_filter_with_mesh_coordinates(
                side_filter_info,
                mesh,
                resolved_config.coordinate_system_origin,
                axis_from_coordinate_system(
                    resolved_axes,
                    resolved_config.section_normal_axis,
                ),
                scoping_geometry_state="reference",
            )
    if reference_side_filter_info is not None:
        deformation["scope_change"] = deformation_scope_change_evidence(
            reference_side_filter_info,
            side_filter_info,
        )
        deformation["scoping_geometry_state"] = side_filter_info.get(
            "scoping_geometry_state",
            "deformed",
        )
    payload = section_visualization_payload_from_mesh(
        mesh=mesh,
        config=resolved_config,
        axes=resolved_axes,
        element_name=snapshot.element_name,
        available_named=snapshot.available_named_selections,
        side_filter_info=side_filter_info,
        dpf_version=getattr(dpf, "__version__", None),
        origin_unit_conversion=origin_unit_conversion,
        deformation=deformation or None,
        resolved_reference_frame=resolved_frame or None,
        result_signature=result_signature,
    )
    payload["counts"]["cut_element_count"] = int(len(object_ids(filtered_elem_scope)))
    payload["counts"]["force_summation_node_count"] = int(len(object_ids(node_scope)))
    payload["elapsed_seconds"] = perf_counter() - phase_start
    return payload


def build_section_visualization_payload(
    config: SectionConfig,
    log: LogFn = None,
) -> Dict[str, Any]:
    import ansys.dpf.core as dpf

    phase_start = perf_counter()
    config = config_from_mapping(asdict(config))
    validate_visualization_paths(config)
    if config.analysis_mode == "static" and config.result_set_id is None:
        raise ValueError("Static visualization requires a cumulative result-set ID.")
    axes = validate_axes(config.coordinate_system_axes)
    signature = modal_rst_signature(config.modal_rst)
    cache_key = visualization_mesh_snapshot_cache_key(
        signature,
        config.element_named_selection,
        config.external_named_selection_path,
        config.analysis_mode,
        config.result_set_id,
        config.reference_frame_motion,
        config.reference_frame_attachment_selection,
    )
    emit(log, f"Loading modal RST for visualization: {config.modal_rst}")
    emit(log, f"PyDPF version: {getattr(dpf, '__version__', 'unknown')}")
    snapshot = _VISUALIZATION_MESH_SNAPSHOT_CACHE.get(cache_key)
    if snapshot is not None:
        _VISUALIZATION_MESH_SNAPSHOT_CACHE.pop(cache_key, None)
        _VISUALIZATION_MESH_SNAPSHOT_CACHE[cache_key] = snapshot
        emit(log, "Visualization mesh cache hit; reusing selected mesh snapshot.")
        payload = _section_visualization_payload_from_snapshot(
            dpf,
            config,
            axes,
            snapshot,
            phase_start=phase_start,
            log=log,
        )
        emit(
            log,
            "Visualization scope ready from cache: "
            f"{payload['counts']['raw_element_count']} selected element(s), "
            f"{payload['counts']['cut_element_count']} cut element(s), "
            f"{payload['counts']['force_summation_node_count']} force node(s).",
        )
        return payload

    emit(log, "Visualization mesh cache miss; reading selected mesh from DPF.")
    preflight_dpf_open(dpf, config.modal_rst, log=log)
    data_sources = dpf.DataSources(config.modal_rst)
    streams_container = create_streams_container(dpf, data_sources)
    try:
        model = dpf.Model(data_sources)
        selection_source = (
            config.external_named_selection_path or "the RST named selections"
        )
        emit(
            log,
            f"Resolving named selection {config.element_named_selection!r} "
            f"from {selection_source} as elemental scoping for visualization.",
        )
        if config.external_named_selection_path:
            element_scoping, element_name, available_named = (
                resolve_external_element_scoping(
                    dpf,
                    model,
                    config.external_named_selection_path,
                    config.element_named_selection,
                )
            )
        else:
            element_scoping, element_name, available_named = (
                resolve_element_named_selection_scoping(
                    dpf,
                    model,
                    data_sources,
                    streams_container,
                    config.element_named_selection,
                )
            )
        emit(
            log,
            f"Visualization named selection {element_name!r}: "
            f"{scoping_id_count(element_scoping)} elemental ID(s).",
        )
        reference_mesh = selected_element_mesh(model, element_scoping)
        nodal_displacements: Dict[int, Vector] = {}
        deformation: Optional[Dict[str, Any]] = None
        resolved_reference_frame: Optional[Dict[str, Any]] = None
        if config.analysis_mode == "static":
            mesh_config, _origin_conversion = config_with_origin_in_mesh_units(
                config,
                reference_mesh,
                log=log,
            )
            initial_axes = validate_axes(mesh_config.coordinate_system_axes)
            _initial_elements, _initial_nodes, initial_side_info = construction_surface_scoping(
                dpf,
                model,
                element_scoping,
                mesh_config,
                initial_axes,
                mesh=reference_mesh,
                allow_empty_cut=True,
            )
            tracking_node_ids: List[int] = []
            tracking_coordinates: Dict[int, Vector] = {}
            attachment_source = "none"
            attachment_selection = ""
            if mesh_config.reference_frame_motion == "follow-geometry":
                attachment_selection = (
                    mesh_config.reference_frame_attachment_selection.strip()
                )
                if attachment_selection:
                    tracking_node_ids, attachment = resolve_reference_frame_attachment_nodes(
                        dpf,
                        model,
                        name=attachment_selection,
                        external_named_selection_path=mesh_config.external_named_selection_path,
                    )
                    attachment_source = attachment["source"]
                    full_mesh = getattr(model.metadata, "meshed_region", None)
                    full_mesh = full_mesh() if callable(full_mesh) else full_mesh
                    tracking_coordinates = mesh_node_coordinates_for_ids(
                        full_mesh,
                        tracking_node_ids,
                    )
                else:
                    tracking_node_ids = unique_element_node_ids(
                        reference_mesh,
                        initial_side_info.get("cut_element_ids", []),
                    )
                    tracking_coordinates = mesh_node_coordinates_for_ids(
                        reference_mesh,
                        tracking_node_ids,
                    )
                    attachment_source = "local_section_cut_neighborhood"
                    attachment_selection = element_name
            section_node_ids = selected_mesh_node_ids(reference_mesh)
            displacement_ids = sorted(set(section_node_ids).union(tracking_node_ids))
            if tracking_node_ids:
                all_displacements, deformation = nodal_displacements_for_ids(
                    dpf,
                    data_sources,
                    streams_container,
                    displacement_ids,
                    int(config.result_set_id),
                    str(getattr(reference_mesh, "unit", "") or ""),
                    log=log,
                )
            else:
                all_displacements, deformation = selected_set_nodal_displacements(
                    dpf,
                    data_sources,
                    streams_container,
                    reference_mesh,
                    int(config.result_set_id),
                    log=log,
                )
            nodal_displacements = {
                node_id: all_displacements[node_id] for node_id in section_node_ids
            }
            try:
                available_result_sets = result_set_options(
                    model.metadata.time_freq_support
                )
            except Exception:
                available_result_sets = []
            selected_option = next(
                (
                    option
                    for option in available_result_sets
                    if int(option["id"]) == int(config.result_set_id)
                ),
                {"value": None, "unit": None},
            )
            if mesh_config.reference_frame_motion == "follow-geometry":
                current_tracking = {
                    node_id: add_vectors(
                        tracking_coordinates[node_id],
                        all_displacements[node_id],
                    )
                    for node_id in tracking_node_ids
                }
                if (
                    mesh_config.reference_frame_attachment_selection.strip()
                    or len(tracking_node_ids) >= 3
                ):
                    frame = fit_geometry_following_frame(
                        tracking_coordinates,
                        current_tracking,
                        mesh_config.coordinate_system_origin,
                        mesh_config.coordinate_system_axes,
                        result_set_id=int(config.result_set_id),
                        result_value=selected_option.get("value"),
                        result_unit=selected_option.get("unit"),
                        attachment_source=attachment_source,
                        attachment_selection=attachment_selection,
                        warning_ratio=mesh_config.reference_frame_fit_warning_ratio,
                    )
                else:
                    frame = None
            else:
                frame = fixed_reference_frame(
                    mesh_config.coordinate_system_origin,
                    mesh_config.coordinate_system_axes,
                    result_set_id=int(config.result_set_id),
                    result_value=selected_option.get("value"),
                    result_unit=selected_option.get("unit"),
                )
            resolved_reference_frame = asdict(frame) if frame is not None else None
        snapshot = visualization_mesh_snapshot_from_mesh(
            cache_key,
            mesh=reference_mesh,
            element_name=element_name,
            available_named=available_named,
            element_scoping=element_scoping,
            nodal_displacements=nodal_displacements,
            deformation=deformation,
            resolved_reference_frame=resolved_reference_frame,
        )
        _VISUALIZATION_MESH_SNAPSHOT_CACHE[cache_key] = snapshot
        while len(_VISUALIZATION_MESH_SNAPSHOT_CACHE) > 3:
            oldest_key = next(iter(_VISUALIZATION_MESH_SNAPSHOT_CACHE))
            _VISUALIZATION_MESH_SNAPSHOT_CACHE.pop(oldest_key, None)
        payload = _section_visualization_payload_from_snapshot(
            dpf,
            config,
            axes,
            snapshot,
            phase_start=phase_start,
            log=log,
        )
        emit(
            log,
            "Visualization scope ready: "
            f"{payload['counts']['raw_element_count']} selected element(s), "
            f"{payload['counts']['cut_element_count']} cut element(s), "
            f"{payload['counts']['force_summation_node_count']} force node(s).",
        )
        return payload
    finally:
        try:
            streams_container.release_handles()
            emit(log, "Released DPF stream handles.")
        except Exception:
            pass
