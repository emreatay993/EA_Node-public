from __future__ import annotations

import csv
import json
import math
from pathlib import Path
import textwrap

import h5py
import numpy as np
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

# Import the add-on catalog before bootstrap to avoid the current lazy import cycle.
import ea_node_editor.addons.catalog  # noqa: F401
from ea_node_editor.nodes.bootstrap import build_default_registry
from ea_node_editor.common.artifact_refs import format_managed_artifact_ref
from ea_node_editor.persistence.artifact_store import format_node_artifact_folder
from ea_node_editor.persistence.serializer import JsonProjectSerializer
from ea_node_editor.settings import SCHEMA_VERSION


ROOT = Path(__file__).resolve().parents[1]
PROJECT_PATH = ROOT / "examples" / "tabular_plot_showcase.cxproj"
DIRECT_PROJECT_PATH = ROOT / "examples" / "tabular_plot_showcase_direct.cxproj"
DATA_ROOT = PROJECT_PATH.with_name("tabular_plot_showcase.data")
DIRECT_DATA_ROOT = DIRECT_PROJECT_PATH.with_name("tabular_plot_showcase_direct.data")


PLOT_DEFINITIONS = (
    {
        "key": "line",
        "title": "Line Plot",
        "plot_type_id": "plot.scatter",
        "source_name": "line_series.csv",
        "artifact_id": "tabular_source.line_csv",
        "source_output": "table_data",
        "adapter_title": "CSV To Line Series",
        "script": "line",
        "selected_columns": ("month", "north", "south"),
        "tabular_mapping": {"x": "month", "y": ("north", "south")},
        "x_label": "Month",
        "y_label": "Revenue",
    },
    {
        "key": "scatter",
        "title": "Scatter Plot",
        "plot_type_id": "plot.scatter",
        "source_name": "scatter_points.tsv",
        "artifact_id": "tabular_source.scatter_tsv",
        "source_output": "table_data",
        "adapter_title": "TSV To Scatter Series",
        "script": "scatter",
        "selected_columns": ("load", "efficiency"),
        "tabular_mapping": {"x": "load", "y": "efficiency"},
        "x_label": "Load",
        "y_label": "Efficiency",
    },
    {
        "key": "bar",
        "title": "Bar Plot",
        "plot_type_id": "plot.bar",
        "source_name": "bar_channels.xlsx",
        "artifact_id": "tabular_source.bar_xlsx",
        "source_output": "table_data",
        "selected_object": "Quarterly Channels",
        "adapter_title": "XLSX To Bar Series",
        "script": "bar",
        "selected_columns": ("channel", "units"),
        "tabular_mapping": {"category": "channel", "y": "units"},
        "x_label": "Channel",
        "y_label": "Units",
    },
    {
        "key": "histogram",
        "title": "Histogram Plot",
        "plot_type_id": "plot.histogram",
        "source_name": "histogram_latency.txt",
        "artifact_id": "tabular_source.histogram_txt",
        "source_output": "table_data",
        "adapter_title": "TXT To Histogram Series",
        "script": "histogram",
        "selected_columns": ("latency_ms",),
        "tabular_mapping": {"values": ("latency_ms",)},
        "x_label": "Latency ms",
        "y_label": "Count",
        "plot_options": {"bins": 12},
    },
    {
        "key": "heatmap",
        "title": "Heatmap Plot",
        "plot_type_id": "plot.heatmap",
        "source_name": "heatmap_grid.npy",
        "artifact_id": "tabular_source.heatmap_npy",
        "source_output": "array_data",
        "adapter_title": "NPY To Heatmap Grid",
        "script": "heatmap",
        "x_label": "Column",
        "y_label": "Row",
        "colormap": "magma",
    },
    {
        "key": "contour",
        "title": "Contour Plot",
        "plot_type_id": "plot.contour",
        "source_name": "contour_field.npz",
        "artifact_id": "tabular_source.contour_npz",
        "source_output": "array_data",
        "selected_object": "contour",
        "adapter_title": "NPZ To Contour Grid",
        "script": "contour",
        "x_label": "X",
        "y_label": "Y",
        "colormap": "viridis",
    },
    {
        "key": "surface",
        "title": "Surface Plot",
        "plot_type_id": "plot.surface",
        "source_name": "surface_elevation.parquet",
        "artifact_id": "tabular_source.surface_parquet",
        "source_output": "table_data",
        "adapter_title": "Parquet To Surface Grid",
        "script": "surface",
        "selected_columns": ("x", "y", "z"),
        "tabular_mapping": {"values": ("x", "y", "z")},
        "x_label": "X",
        "y_label": "Y",
        "z_label": "Elevation",
        "colormap": "terrain",
        "plot_options": {"show_edges": True},
    },
    {
        "key": "point_cloud",
        "title": "Point Cloud Plot",
        "plot_type_id": "plot.point_cloud",
        "source_name": "point_cloud.csv",
        "artifact_id": "tabular_source.point_cloud_csv",
        "source_output": "table_data",
        "adapter_title": "CSV To Point Cloud",
        "script": "point_cloud",
        "selected_columns": ("x", "y", "z", "intensity"),
        "tabular_mapping": {"x": "x", "y": "y", "z": "z"},
        "x_label": "X",
        "y_label": "Y",
        "z_label": "Z",
        "colormap": "turbo",
        "plot_options": {"point_size": 7},
    },
    {
        "key": "streamlines",
        "title": "Streamlines Plot",
        "plot_type_id": "plot.streamlines",
        "source_name": "streamline_path.h5",
        "artifact_id": "tabular_source.streamlines_hdf5",
        "source_output": "array_data",
        "selected_object": "streamline_path",
        "adapter_title": "HDF5 To Streamline Path",
        "script": "streamlines",
        "array_slice_2d": {"row_limit": 80, "column_limit": 3},
        "x_label": "X",
        "y_label": "Y",
        "z_label": "Z",
        "colormap": "cividis",
        "plot_options": {"line_width": 4, "render_lines_as_tubes": True},
    },
)
PLOT_DEFINITIONS_BY_SOURCE_NAME = {str(definition["source_name"]): definition for definition in PLOT_DEFINITIONS}


def _source_relative_path(definition: dict[str, object], *, workspace_id: str) -> Path:
    node_folder = format_node_artifact_folder(
        workspace_id=workspace_id,
        node_id=f"node_{definition['key']}_input",
        node_title=str(definition["source_name"]),
        node_type="tabular.input",
    )
    return Path("nodes") / node_folder / "in" / "tabular" / "source" / str(definition["source_name"])


def _source_path(data_root: Path, source_name: str, *, workspace_id: str) -> Path:
    return data_root / _source_relative_path(PLOT_DEFINITIONS_BY_SOURCE_NAME[source_name], workspace_id=workspace_id)


SCRIPT_PREAMBLE = r'''
from pathlib import Path
from urllib.parse import unquote, urlparse


def _field(value, key, default=""):
    if isinstance(value, dict):
        return value.get(key, default)
    return getattr(value, key, default)


def _source_path(ref):
    uri = str(_field(ref, "source_uri", "") or "").strip()
    if not uri:
        raise ValueError("Tabular ref did not include a source_uri.")
    parsed = urlparse(uri)
    if parsed.scheme == "file":
        raw_path = unquote(parsed.path or parsed.netloc)
        if raw_path.startswith("/") and len(raw_path) >= 3 and raw_path[2] == ":":
            raw_path = raw_path[1:]
        return Path(raw_path)
    return Path(uri)


_source_path.__globals__.update(
    {"_field": _field, "Path": Path, "unquote": unquote, "urlparse": urlparse}
)
path = _source_path(payload)
object_id = str(_field(payload, "object_id", "") or "")
'''


SCRIPT_BODIES = {
    "line": r'''
import csv

with path.open("r", encoding="utf-8", newline="") as handle:
    rows = list(csv.DictReader(handle))

months = []
north = []
south = []
for row in rows:
    months.append(float(row["month"]))
    north.append(float(row["north"]))
    south.append(float(row["south"]))

output_data = [
    {"label": "North", "x": months, "y": north},
    {"label": "South", "x": months, "y": south},
]
''',
    "scatter": r'''
import csv

with path.open("r", encoding="utf-8", newline="") as handle:
    rows = list(csv.DictReader(handle, delimiter="\t"))

x_values = []
y_values = []
for row in rows:
    x_values.append(float(row["load"]))
    y_values.append(float(row["efficiency"]))

output_data = {
    "label": "Sensor lots",
    "x": x_values,
    "y": y_values,
}
''',
    "bar": r'''
import openpyxl

workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
try:
    sheet = workbook[object_id or "Quarterly Channels"]
    rows = list(sheet.iter_rows(values_only=True))
finally:
    workbook.close()

headers = []
for value in rows[0]:
    headers.append(str(value))
x_values = []
y_values = []
for row in rows[1:]:
    record = dict(zip(headers, row))
    x_values.append(str(record["channel"]))
    y_values.append(float(record["units"]))

output_data = {
    "label": "Q4 units",
    "x": x_values,
    "y": y_values,
}
''',
    "histogram": r'''
import csv

with path.open("r", encoding="utf-8", newline="") as handle:
    rows = list(csv.DictReader(handle))

values = []
for row in rows:
    values.append(float(row["latency_ms"]))

output_data = {
    "label": "Latency",
    "values": values,
    "bins": 12,
}
''',
    "heatmap": r'''
import numpy as np

grid = np.load(path, allow_pickle=False)
output_data = {"label": "Thermal grid", "values": grid.tolist()}
''',
    "contour": r'''
import numpy as np

with np.load(path, allow_pickle=False) as archive:
    key = object_id or "contour"
    output_data = {"label": key, "z": archive[key].tolist()}
''',
    "surface": r'''
import pyarrow.parquet as pq

columns = pq.read_table(path).to_pydict()
x_set = set()
y_set = set()
for value in columns["x"]:
    x_set.add(float(value))
for value in columns["y"]:
    y_set.add(float(value))
xs = sorted(x_set)
ys = sorted(y_set)
lookup = {}
for x_value, y_value, z_value in zip(columns["x"], columns["y"], columns["z"]):
    lookup[(float(x_value), float(y_value))] = float(z_value)

z_grid = []
for y_value in ys:
    row_values = []
    for x_value in xs:
        row_values.append(lookup[(x_value, y_value)])
    z_grid.append(row_values)

output_data = {
    "label": "Elevation",
    "x": xs,
    "y": ys,
    "z": z_grid,
}
''',
    "point_cloud": r'''
import csv

with path.open("r", encoding="utf-8", newline="") as handle:
    rows = list(csv.DictReader(handle))

x_values = []
y_values = []
z_values = []
scalars = []
for row in rows:
    x_values.append(float(row["x"]))
    y_values.append(float(row["y"]))
    z_values.append(float(row["z"]))
    scalars.append(float(row["intensity"]))

output_data = {
    "label": "Scan points",
    "x": x_values,
    "y": y_values,
    "z": z_values,
    "scalars": scalars,
    "point_size": 7,
}
''',
    "streamlines": r'''
import h5py

with h5py.File(path, "r") as handle:
    points = handle[object_id or "streamline_path"][...]

output_data = {
    "label": "Flow path",
    "points": points.tolist(),
    "line_width": 4,
    "color": "cyan",
}
''',
}


def _adapter_script(name: str) -> str:
    body = textwrap.indent((SCRIPT_PREAMBLE + SCRIPT_BODIES[name]).strip(), "    ")
    return (
        "import corex\n\n"
        "@corex.node\n"
        '@corex.input("payload", value_type=corex.Any)\n'
        '@corex.output("result", value_type=corex.Any)\n'
        "def run(ctx, payload):\n"
        f"{body}\n"
        '    return {"result": output_data}\n'
    )


def _write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str], *, delimiter: str = ",") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


def write_sources(data_root: Path = DATA_ROOT, *, workspace_id: str = "ws_tabular_plot_showcase") -> None:
    data_root.mkdir(parents=True, exist_ok=True)
    _write_csv(
        _source_path(data_root, "line_series.csv", workspace_id=workspace_id),
        [
            {
                "month": month,
                "north": round(112 + month * 6.4 + math.sin(month / 2) * 8, 2),
                "south": round(96 + month * 5.2 + math.cos(month / 3) * 6, 2),
            }
            for month in range(1, 13)
        ],
        ["month", "north", "south"],
    )
    _write_csv(
        _source_path(data_root, "scatter_points.tsv", workspace_id=workspace_id),
        [
            {
                "load": round(18 + index * 2.5, 2),
                "efficiency": round(74 + math.sin(index / 2.7) * 9 + index * 0.32, 2),
            }
            for index in range(28)
        ],
        ["load", "efficiency"],
        delimiter="\t",
    )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Quarterly Channels"
    sheet.append(["channel", "units"])
    for channel, units in (
        ("Direct", 1540),
        ("Partner", 1210),
        ("Marketplace", 1840),
        ("Services", 960),
        ("Trial", 520),
    ):
        sheet.append([channel, units])
    bar_channels_path = _source_path(data_root, "bar_channels.xlsx", workspace_id=workspace_id)
    bar_channels_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(bar_channels_path)
    workbook.close()
    _write_csv(
        _source_path(data_root, "histogram_latency.txt", workspace_id=workspace_id),
        [
            {
                "latency_ms": round(
                    42 + math.sin(index * 0.7) * 12 + (index % 9) * 3.8 + (index % 5) * 1.4,
                    2,
                )
            }
            for index in range(96)
        ],
        ["latency_ms"],
    )
    heatmap = np.array(
        [
            [
                round(18 + math.sin(row / 1.6) * 6 + math.cos(col / 1.9) * 5 + row * 0.7 + col * 0.3, 3)
                for col in range(12)
            ]
            for row in range(9)
        ],
        dtype=np.float64,
    )
    heatmap_path = _source_path(data_root, "heatmap_grid.npy", workspace_id=workspace_id)
    heatmap_path.parent.mkdir(parents=True, exist_ok=True)
    np.save(heatmap_path, heatmap, allow_pickle=False)
    xs = np.linspace(-3.0, 3.0, 25)
    ys = np.linspace(-2.5, 2.5, 21)
    contour = np.array(
        [
            [
                math.sin(x * 1.4) * math.cos(y * 1.2) + math.exp(-((x - 0.8) ** 2 + (y + 0.4) ** 2))
                for x in xs
            ]
            for y in ys
        ],
        dtype=np.float64,
    )
    contour_path = _source_path(data_root, "contour_field.npz", workspace_id=workspace_id)
    contour_path.parent.mkdir(parents=True, exist_ok=True)
    np.savez(contour_path, contour=contour, comparison=contour * 0.75)
    surface_rows = []
    for y in np.linspace(-3.0, 3.0, 17):
        for x in np.linspace(-3.0, 3.0, 17):
            z = math.sin(x) * math.cos(y) + 0.18 * x
            surface_rows.append({"x": float(x), "y": float(y), "z": float(z)})
    surface_path = _source_path(data_root, "surface_elevation.parquet", workspace_id=workspace_id)
    surface_path.parent.mkdir(parents=True, exist_ok=True)
    pq.write_table(pa.Table.from_pylist(surface_rows), surface_path)
    _write_csv(
        _source_path(data_root, "point_cloud.csv", workspace_id=workspace_id),
        [
            {
                "x": round(math.cos(index * 0.35) * (1.5 + index / 38), 5),
                "y": round(math.sin(index * 0.35) * (1.5 + index / 38), 5),
                "z": round((index % 18) / 3.0 + math.sin(index / 5) * 0.7, 5),
                "intensity": round(20 + index * 1.8, 5),
            }
            for index in range(72)
        ],
        ["x", "y", "z", "intensity"],
    )
    streamline = np.array(
        [
            [
                math.cos(t) * (1.0 + t / 18),
                math.sin(t) * (1.0 + t / 18),
                t / 2.8,
            ]
            for t in np.linspace(0.0, 9.0, 80)
        ],
        dtype=np.float64,
    )
    streamline_path = _source_path(data_root, "streamline_path.h5", workspace_id=workspace_id)
    streamline_path.parent.mkdir(parents=True, exist_ok=True)
    with h5py.File(streamline_path, "w") as handle:
        handle.create_dataset("streamline_path", data=streamline)


def _tabular_properties(definition: dict[str, object], *, direct: bool = False) -> dict[str, object]:
    array_slice = {
        "row_offset": 0,
        "column_offset": 0,
        "row_limit": 50,
        "column_limit": 50,
    }
    if direct:
        array_slice.update(dict(definition.get("array_slice_2d", {})))
    properties = {
        "path": format_managed_artifact_ref(str(definition["artifact_id"])),
        "delimiter": "",
        "encoding": "utf-8",
        "header_row": 0,
        "skip_rows": 0,
        "schema_hints": {},
        "selected_object": str(definition.get("selected_object", "")),
        "array_slice_2d": array_slice,
        "cache_policy": "source_direct",
        "project_managed_source": True,
        "project_managed_cache": False,
        "allow_npz_archive_preview": True,
        "tabular_table_view_state": {},
    }
    if direct:
        properties["tabular_selected_columns"] = list(definition.get("selected_columns", ()))
    return properties


def _plot_properties(definition: dict[str, object], *, direct: bool = False) -> dict[str, object]:
    properties = {
        "backend": "auto",
        "title": str(definition["title"]),
        "x_label": str(definition.get("x_label", "")),
        "y_label": str(definition.get("y_label", "")),
        "z_label": str(definition.get("z_label", "")),
        "axis_limits": {"x": [None, None], "y": [None, None], "z": [None, None]},
        "log_scales": {"x": False, "y": False, "z": False},
        "grid": True,
        "legend": True,
        "render_in_canvas": True,
        "archive_export_on_run": False,
        "static_export_format": "png",
        "data_export_format": "csv",
        "plot_options": dict(definition.get("plot_options", {})),
    }
    if "colormap" in definition:
        properties["colormap"] = str(definition["colormap"])
    if direct and "tabular_mapping" in definition:
        properties["tabular_mapping"] = dict(definition["tabular_mapping"])  # type: ignore[arg-type]
    return properties


def _node(
    *,
    node_id: str,
    type_id: str,
    title: str,
    x: float,
    y: float,
    properties: dict[str, object] | None = None,
    exposed_ports: dict[str, bool] | None = None,
    custom_width: float | None = None,
    custom_height: float | None = None,
) -> dict[str, object]:
    return {
        "node_id": node_id,
        "type_id": type_id,
        "title": title,
        "x": x,
        "y": y,
        "collapsed": False,
        "properties": dict(properties or {}),
        "exposed_ports": dict(exposed_ports or {}),
        "port_labels": {},
        "visual_style": {},
        "parent_node_id": None,
        "custom_width": custom_width,
        "custom_height": custom_height,
    }


def _edge(
    edge_id: str,
    source_node_id: str,
    source_port_key: str,
    target_node_id: str,
    target_port_key: str,
) -> dict[str, object]:
    return {
        "edge_id": edge_id,
        "source_node_id": source_node_id,
        "source_port_key": source_port_key,
        "target_node_id": target_node_id,
        "target_port_key": target_port_key,
        "enabled": True,
        "input_order": 0,
        "label": "",
        "visual_style": {},
    }


def build_document(*, direct: bool = False) -> dict[str, object]:
    nodes: list[dict[str, object]] = []
    edges: list[dict[str, object]] = []
    workspace_id = "ws_tabular_plot_showcase_direct" if direct else "ws_tabular_plot_showcase"
    view_id = "view_tabular_plot_showcase_direct" if direct else "view_tabular_plot_showcase"
    project_id = "proj_tabular_plot_showcase_direct" if direct else "proj_tabular_plot_showcase"
    project_name = "Tabular Plot Showcase Direct" if direct else "Tabular Plot Showcase"
    description = (
        "Loads CSV, TSV, XLSX, TXT, NPY, NPZ, Parquet, CSV point cloud, "
        "and HDF5 streamline sources directly into the generic plot node family "
        "without adapter nodes."
        if direct
        else (
            "Loads CSV, TSV, XLSX, TXT, NPY, NPZ, Parquet, CSV point cloud, "
            "and HDF5 streamline sources into the generic plot node family."
        )
    )
    source_x = -820.0 if direct else -860.0
    plot_x = -300.0 if direct else 60.0

    for index, definition in enumerate(PLOT_DEFINITIONS):
        key = str(definition["key"])
        y = index * 220.0
        source_id = f"node_{key}_input"
        adapter_id = f"node_{key}_adapter"
        plot_id = f"node_{key}_plot"
        nodes.append(
            _node(
                node_id=source_id,
                type_id="tabular.input",
                title=f"{definition['source_name']}",
                x=source_x,
                y=y,
                properties=_tabular_properties(definition, direct=direct),
                exposed_ports={
                    "path": True,
                    "table_data": True,
                    "array_data": True,
                },
                custom_width=360.0,
                custom_height=220.0,
            )
        )
        if not direct:
            nodes.append(
                _node(
                    node_id=adapter_id,
                    type_id="core.python_script",
                    title=str(definition["adapter_title"]),
                    x=-390.0,
                    y=y + 16.0,
                    properties={"script": _adapter_script(str(definition["script"]))},
                    exposed_ports={"payload": True, "result": True},
                    custom_width=310.0,
                )
            )
        nodes.append(
            _node(
                node_id=plot_id,
                type_id=str(definition["plot_type_id"]),
                title=str(definition["title"]),
                x=plot_x,
                y=y,
                properties=_plot_properties(definition, direct=direct),
                exposed_ports={
                    "series": True,
                    "static_export": True,
                    "data_export": True,
                    "exports": True,
                },
                custom_width=380.0,
                custom_height=210.0,
            )
        )
        if direct:
            edges.append(
                _edge(
                    f"edge_data_{key}_ref",
                    source_id,
                    str(definition["source_output"]),
                    plot_id,
                    "series",
                )
            )
        else:
            edges.extend(
                [
                    _edge(
                        f"edge_data_{key}_ref",
                        source_id,
                        str(definition["source_output"]),
                        adapter_id,
                        "payload",
                    ),
                    _edge(
                        f"edge_data_{key}_series",
                        adapter_id,
                        "result",
                        plot_id,
                        "series",
                    ),
                ]
            )

    nodes.append(
        _node(
            node_id="node_showcase_note",
            type_id="passive.annotation.group_backdrop",
            title=project_name,
            x=-940.0,
            y=-190.0,
            properties={"title": project_name},
            custom_width=1480.0,
            custom_height=2180.0,
        )
    )

    artifact_entries = {
        str(definition["artifact_id"]): {
            "artifact_kind": "tabular_source",
            "relative_path": _source_relative_path(definition, workspace_id=workspace_id).as_posix(),
        }
        for definition in PLOT_DEFINITIONS
    }
    return {
        "schema_version": SCHEMA_VERSION,
        "project_id": project_id,
        "name": project_name,
        "active_workspace_id": workspace_id,
        "workspace_order": [workspace_id],
        "workspaces": [
            {
                "workspace_id": workspace_id,
                "name": project_name,
                "dirty": False,
                "active_view_id": view_id,
                "views": [
                    {
                        "view_id": view_id,
                        "name": "Showcase",
                        "zoom": 0.72,
                        "pan_x": 1020.0,
                        "pan_y": 220.0,
                        "scope_path": [],
                        "hide_optional_ports": True,
                    }
                ],
                "nodes": nodes,
                "edges": edges,
            }
        ],
        "metadata": {
            "artifact_store": {
                "artifacts": artifact_entries,
                "staged": {},
            },
            "workflow_settings": {
                "general": {
                    "project_name": project_name,
                    "description": description,
                    "author": "",
                },
                "environment": {"python_path": "", "working_directory": ""},
                "logging": {"level": "info", "capture_console": True},
                "plugins": {"enabled": []},
                "solver_config": {
                    "thread_count": 8,
                    "memory_limit_gb": 12,
                    "enable_parallel": True,
                },
            },
            "ui": {
                "script_editor": {"visible": False, "floating": False},
                "passive_style_presets": {"node_presets": [], "edge_presets": []},
            },
            "custom_workflows": [],
            "workspace_order": [workspace_id],
        },
    }


def main() -> None:
    write_sources()
    write_sources(DIRECT_DATA_ROOT, workspace_id="ws_tabular_plot_showcase_direct")
    registry = build_default_registry()
    missing = [
        type_id
        for type_id in {"tabular.input", "core.python_script", *(item["plot_type_id"] for item in PLOT_DEFINITIONS)}
        if registry.spec_or_none(str(type_id)) is None
    ]
    if missing:
        raise RuntimeError(f"Required node types are not registered: {', '.join(sorted(map(str, missing)))}")
    serializer = JsonProjectSerializer(registry)
    for path, raw_document in (
        (PROJECT_PATH, build_document(direct=False)),
        (DIRECT_PROJECT_PATH, build_document(direct=True)),
    ):
        project = serializer.from_document(json.loads(json.dumps(raw_document)))
        persistent_document = serializer.to_persistent_document(project)
        serializer.save_document(str(path), persistent_document)


if __name__ == "__main__":
    main()
