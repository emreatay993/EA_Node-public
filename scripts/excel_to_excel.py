from __future__ import annotations

import argparse
import hashlib
import math
import os
import sys
import time
import traceback
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from PyQt6.QtCore import QSize, QThread, pyqtSignal
    from PyQt6.QtWidgets import (
        QApplication,
        QComboBox,
        QFileDialog,
        QGridLayout,
        QGroupBox,
        QHBoxLayout,
        QLabel,
        QLineEdit,
        QMessageBox,
        QPushButton,
        QSizePolicy,
        QSpinBox,
        QStyle,
        QTextEdit,
        QVBoxLayout,
        QWidget,
    )

    QT_MAJOR = 6
except ImportError:  # pragma: no cover - fallback for machines with PyQt5 only.
    from PyQt5.QtCore import QSize, QThread, pyqtSignal
    from PyQt5.QtWidgets import (
        QApplication,
        QComboBox,
        QFileDialog,
        QGridLayout,
        QGroupBox,
        QHBoxLayout,
        QLabel,
        QLineEdit,
        QMessageBox,
        QPushButton,
        QSizePolicy,
        QSpinBox,
        QStyle,
        QTextEdit,
        QVBoxLayout,
        QWidget,
    )

    QT_MAJOR = 5


BATCH_SIZE = 6
MAX_BATCH_SIZE = 10000
MOCK_FIRST_VISIBLE_INPUT_ROWS = 4000
SMALL_MOCK_FIRST_VISIBLE_INPUT_ROWS = 13
MOCK_VISIBLE_LOAD_COLS = 20
MOCK_VISIBLE_LOAD_ESTIMATED_BYTES_PER_ROW = 930

XL_UP = -4162
XL_CALCULATION_MANUAL = -4135
XL_CALCULATION_DONE = 0

MOCK_TARGET_SIZES_MB = {
    "mock_first_input.xlsx": 30,
    "mock_second_processor.xlsx": 17,
}

MOCK_FIRST_SHEET = "GirdiVerisi"
MOCK_SECOND_SHEET = "İşlemci"

FIRST_DATA_START_ROW = 3
FIRST_INPUT_START_COL = 6  # F
FIRST_INPUT_END_COL = 46  # AT
FIRST_OUTPUT_START_COL = 47  # AU
FIRST_OUTPUT_END_COL = 69  # BQ

SECOND_INPUT_START_ROW = 9
SECOND_INPUT_START_COL = 2  # B
SECOND_INPUT_END_COL = 42  # AP
SECOND_OUTPUT_ROW = 9
SECOND_OUTPUT_START_COL = 43  # AQ
SECOND_OUTPUT_END_COL = 65  # BM


def col_range(start_col: int, end_col: int) -> range:
    return range(start_col, end_col + 1)


def iter_cells_2d(
    ws,
    row1: int,
    col1: int,
    row2: int,
    col2: int,
) -> Iterable[tuple[int, int]]:
    for row in range(row1, row2 + 1):
        for col in range(col1, col2 + 1):
            yield row, col


def clear_range(ws, row1: int, col1: int, row2: int, col2: int) -> None:
    for row, col in iter_cells_2d(ws, row1, col1, row2, col2):
        ws.cell(row=row, column=col).value = None


def read_values(ws, row1: int, col1: int, row2: int, col2: int) -> list[list[object]]:
    return [
        [ws.cell(row=row, column=col).value for col in col_range(col1, col2)]
        for row in range(row1, row2 + 1)
    ]


def write_values(ws, start_row: int, start_col: int, values: Sequence[Sequence[object]]) -> None:
    for row_offset, row_values in enumerate(values):
        for col_offset, value in enumerate(row_values):
            ws.cell(row=start_row + row_offset, column=start_col + col_offset).value = value


def sheet_names(file_path: Path) -> list[str]:
    wb = load_workbook(file_path, read_only=True, keep_vba=file_path.suffix.lower() == ".xlsm")
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def visible_load_value(workbook_key: str, row: int, col: int) -> str:
    digest = hashlib.sha256(f"{workbook_key}:{row}:{col}:gorunur-yuk".encode("ascii")).hexdigest()
    return f"{workbook_key}_gorunur_yuk_{row:06d}_{col:02d}_{digest}"


def ensure_visible_load_sheet(wb: Workbook, sheet_name: str):
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        ws.append([f"yuk_kolonu_{index:02d}" for index in range(1, MOCK_VISIBLE_LOAD_COLS + 1)])
        style_header(ws, 1, 1, MOCK_VISIBLE_LOAD_COLS, "FFF2CC")
        set_column_widths(ws, 1, MOCK_VISIBLE_LOAD_COLS, 24)
        ws.freeze_panes = "A2"

    return ws


def append_visible_load_rows(ws, workbook_key: str, row_count: int) -> None:
    start_row = ws.max_row + 1
    for row in range(start_row, start_row + row_count):
        ws.append(
            [
                visible_load_value(workbook_key, row, col)
                for col in range(1, MOCK_VISIBLE_LOAD_COLS + 1)
            ]
        )


def inflate_workbook_with_visible_rows(
    wb: Workbook,
    file_path: Path,
    target_mb: int,
    sheet_name: str,
    workbook_key: str,
) -> None:
    target_bytes = target_mb * 1024 * 1024
    ws = ensure_visible_load_sheet(wb, sheet_name)

    wb.save(file_path)
    current_bytes = file_path.stat().st_size

    attempts = 0
    while current_bytes < target_bytes and attempts < 6:
        remaining_bytes = target_bytes - current_bytes
        rows_to_add = max(
            50,
            math.ceil(remaining_bytes / MOCK_VISIBLE_LOAD_ESTIMATED_BYTES_PER_ROW),
        )
        append_visible_load_rows(ws, workbook_key, rows_to_add)
        wb.save(file_path)
        current_bytes = file_path.stat().st_size
        attempts += 1


def excel_range(ws, row1: int, col1: int, row2: int, col2: int):
    return ws.Range(ws.Cells(row1, col1), ws.Cells(row2, col2))


def normalize_com_2d(value, rows: int, cols: int) -> tuple[tuple[object, ...], ...]:
    if rows == 1 and cols == 1:
        return ((value,),)

    if rows == 1:
        if isinstance(value, tuple) and len(value) == 1 and isinstance(value[0], tuple):
            return value
        if isinstance(value, tuple):
            return (value,)
        return ((value,),)

    if cols == 1:
        if isinstance(value, tuple):
            return tuple(row if isinstance(row, tuple) else (row,) for row in value)
        return tuple((value,) for _ in range(rows))

    return value


def find_last_used_row_com(ws, start_col: int, end_col: int) -> int:
    last_row = FIRST_DATA_START_ROW - 1
    for col in col_range(start_col, end_col):
        row = ws.Cells(ws.Rows.Count, col).End(XL_UP).Row
        if row > last_row:
            last_row = row
    return last_row


def wait_for_excel_calculation(excel, timeout_seconds: float = 3600) -> None:
    started = time.monotonic()
    while excel.CalculationState != XL_CALCULATION_DONE:
        if time.monotonic() - started > timeout_seconds:
            raise TimeoutError("Excel hesaplaması zaman aşımı dolmadan bitmedi.")
        time.sleep(0.2)


def press_f9_on_selected_second_sheet(excel, wb_second, ws_second, timeout_seconds: float = 3600) -> None:
    wb_second.Activate()
    ws_second.Activate()

    # Excel keyboard F9 recalculates open workbooks. Activating the second sheet first
    # gives Worksheet_Calculate handlers and sheet-scoped logic the same active context.
    excel.Calculate()
    wait_for_excel_calculation(excel, timeout_seconds=timeout_seconds)


def process_workbooks(
    first_file: Path,
    first_sheet: str,
    second_file: Path,
    second_sheet: str,
    log: Callable[[str], None],
    max_batches: int | None = None,
    batch_size: int = BATCH_SIZE,
) -> None:
    import pythoncom
    import win32com.client as win32

    if batch_size < 1:
        raise ValueError("Grup satir sayisi en az 1 olmali.")

    first_file = first_file.resolve()
    second_file = second_file.resolve()

    pythoncom.CoInitialize()
    excel = None
    wb_first = None
    wb_second = None
    previous_calculation = None

    try:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = True

        wb_first = excel.Workbooks.Open(str(first_file), 0, False)
        wb_second = excel.Workbooks.Open(str(second_file), 0, False)

        try:
            previous_calculation = excel.Calculation
            excel.Calculation = XL_CALCULATION_MANUAL
        except Exception as exc:
            log(f"Excel elle hesaplama moduna alınamadı; mevcut modla devam ediliyor: {exc}")

        ws_first = wb_first.Worksheets(first_sheet)
        ws_second = wb_second.Worksheets(second_sheet)

        last_row = find_last_used_row_com(
            ws_first,
            FIRST_INPUT_START_COL,
            FIRST_INPUT_END_COL,
        )

        if last_row < FIRST_DATA_START_ROW:
            raise ValueError("F3:AT aralığında input verisi bulunamadı.")

        batch_index = 1
        for source_start_row in range(FIRST_DATA_START_ROW, last_row + 1, batch_size):
            source_end_row = min(source_start_row + batch_size - 1, last_row)
            row_count = source_end_row - source_start_row + 1

            log(f"Grup {batch_index}: ilk Excel satırları okunuyor: {source_start_row}-{source_end_row}")
            source_range = excel_range(
                ws_first,
                source_start_row,
                FIRST_INPUT_START_COL,
                source_end_row,
                FIRST_INPUT_END_COL,
            )
            source_values = normalize_com_2d(
                source_range.Value,
                row_count,
                FIRST_INPUT_END_COL - FIRST_INPUT_START_COL + 1,
            )

            input_end_row = SECOND_INPUT_START_ROW + batch_size - 1
            log(f"  input ikinci Excel'e B9:AP{input_end_row} aralığına yazılıyor")
            input_clear_range = excel_range(
                ws_second,
                SECOND_INPUT_START_ROW,
                SECOND_INPUT_START_COL,
                input_end_row,
                SECOND_INPUT_END_COL,
            )
            input_clear_range.ClearContents()

            target_input_range = excel_range(
                ws_second,
                SECOND_INPUT_START_ROW,
                SECOND_INPUT_START_COL,
                SECOND_INPUT_START_ROW + row_count - 1,
                SECOND_INPUT_END_COL,
            )
            target_input_range.Value = source_values

            log("  seçili ikinci Excel sayfasında gerçek F9 hesaplaması çalıştırılıyor")
            press_f9_on_selected_second_sheet(excel, wb_second, ws_second)

            log(
                "  ikinci Excel'deki "
                f"AQ{SECOND_OUTPUT_ROW}:BM{SECOND_OUTPUT_ROW + row_count - 1} "
                f"aralığı ilk Excel'de AU{source_start_row}:BQ{source_end_row} aralığına kopyalanıyor"
            )
            second_output_range = excel_range(
                ws_second,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_START_COL,
                SECOND_OUTPUT_ROW + row_count - 1,
                SECOND_OUTPUT_END_COL,
            )
            output_values = normalize_com_2d(
                second_output_range.Value,
                row_count,
                SECOND_OUTPUT_END_COL - SECOND_OUTPUT_START_COL + 1,
            )

            first_output_range = excel_range(
                ws_first,
                source_start_row,
                FIRST_OUTPUT_START_COL,
                source_end_row,
                FIRST_OUTPUT_END_COL,
            )
            first_output_range.Value = output_values

            batch_index += 1
            if max_batches is not None and batch_index > max_batches:
                log(f"{max_batches} grup işlendiği için durduruldu.")
                break

        if previous_calculation is not None:
            try:
                excel.Calculation = previous_calculation
            except Exception as exc:
                log(f"Excel hesaplama modu geri alınamadı: {exc}")
        wb_first.Save()
        wb_second.Save()
        log(f"İlk Excel kaydedildi: {first_file}")
        log(f"İkinci Excel kaydedildi: {second_file}")
    finally:
        if wb_second is not None:
            wb_second.Close(SaveChanges=False)
        if wb_first is not None:
            wb_first.Close(SaveChanges=False)
        if excel is not None:
            excel.DisplayAlerts = True
            excel.Quit()
        pythoncom.CoUninitialize()


def style_header(ws, row: int, start_col: int, end_col: int, fill_color: str) -> None:
    fill = PatternFill("solid", fgColor=fill_color)
    for col in col_range(start_col, end_col):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(bold=True)


def set_column_widths(ws, start_col: int, end_col: int, width: float) -> None:
    for col in col_range(start_col, end_col):
        ws.column_dimensions[get_column_letter(col)].width = width


def standard_help_icon(widget: QWidget):
    if QT_MAJOR == 6:
        pixmap_id = QStyle.StandardPixmap.SP_MessageBoxQuestion
    else:
        pixmap_id = QStyle.SP_MessageBoxQuestion
    return widget.style().standardIcon(pixmap_id)


def question_yes_no(parent: QWidget, title: str, text: str) -> bool:
    if QT_MAJOR == 6:
        yes = QMessageBox.StandardButton.Yes
        no = QMessageBox.StandardButton.No
    else:
        yes = QMessageBox.Yes
        no = QMessageBox.No

    result = QMessageBox.question(parent, title, text, yes | no, no)
    return result == yes


def create_mock_workbooks(output_dir: Path, small: bool = False, batch_size: int = BATCH_SIZE) -> tuple[Path, Path]:
    if batch_size < 1:
        raise ValueError("Grup satir sayisi en az 1 olmali.")

    output_dir.mkdir(parents=True, exist_ok=True)

    first_file = output_dir / "mock_first_input.xlsx"
    second_file = output_dir / "mock_second_processor.xlsx"
    input_rows = SMALL_MOCK_FIRST_VISIBLE_INPUT_ROWS if small else MOCK_FIRST_VISIBLE_INPUT_ROWS

    wb_first = Workbook()
    ws_first = wb_first.active
    ws_first.title = MOCK_FIRST_SHEET
    ws_first["A1"] = "Örnek ilk Excel. Inputlar F3:AT aralığında. Outputlar AU:BQ aralığına yazılır."

    for index, col in enumerate(col_range(FIRST_INPUT_START_COL, FIRST_INPUT_END_COL), start=1):
        ws_first.cell(row=2, column=col).value = f"girdi_{index:02d}"
    for index, col in enumerate(col_range(FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL), start=1):
        ws_first.cell(row=2, column=col).value = f"çıktı_{index:02d}"

    for row in range(FIRST_DATA_START_ROW, FIRST_DATA_START_ROW + input_rows):
        for index, col in enumerate(col_range(FIRST_INPUT_START_COL, FIRST_INPUT_END_COL), start=1):
            if (row + index) % 3 == 0:
                value = round(row * 10 + index * 0.25, 2)
            elif (row + index) % 3 == 1:
                value = f"R{row}_C{get_column_letter(col)}"
            else:
                value = row + index
            ws_first.cell(row=row, column=col).value = value

    style_header(ws_first, 2, FIRST_INPUT_START_COL, FIRST_INPUT_END_COL, "D9EAD3")
    style_header(ws_first, 2, FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL, "CFE2F3")
    set_column_widths(ws_first, FIRST_INPUT_START_COL, FIRST_INPUT_END_COL, 12)
    set_column_widths(ws_first, FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL, 17)
    ws_first.freeze_panes = "F3"

    ws_first_other = wb_first.create_sheet("EkSayfa")
    ws_first_other["A1"] = "Bu ek sayfa, sayfa seçiciyi test etmek için var."

    if small:
        wb_first.save(first_file)
    else:
        inflate_workbook_with_visible_rows(
            wb_first,
            first_file,
            MOCK_TARGET_SIZES_MB[first_file.name],
            "GörünürYük_İlk",
            "ilk",
        )
    wb_first.close()

    wb_second = Workbook()
    ws_second = wb_second.active
    ws_second.title = MOCK_SECOND_SHEET
    batch_end_row = SECOND_INPUT_START_ROW + batch_size - 1
    ws_second["A1"] = f"Örnek ikinci Excel. Inputlar B9:AP{batch_end_row} aralığına yapıştırılır."
    ws_second["A2"] = (
        "Demo inputu yapıştırır, gerçek Excel F9 hesaplamasını çalıştırır, "
        f"sonra AQ9:BM{batch_end_row} aralığını geri kopyalar."
    )

    for index, col in enumerate(col_range(SECOND_INPUT_START_COL, SECOND_INPUT_END_COL), start=1):
        ws_second.cell(row=8, column=col).value = f"girdi_{index:02d}"
    for index, col in enumerate(col_range(SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL), start=1):
        ws_second.cell(row=8, column=col).value = f"çıktı_{index:02d}"
        for row in range(SECOND_OUTPUT_ROW, SECOND_OUTPUT_ROW + batch_size):
            input_row_range = f"$B{row}:$AP{row}"
            if index % 4 == 1:
                formula = f'="gercek_f9_satir_"&ROW()&"_cikti_{index:02d}"'
            elif index % 4 == 2:
                formula = f"=COUNTA({input_row_range})+{index}"
            elif index % 4 == 3:
                formula = f"=SUM({input_row_range})+{index}"
            else:
                formula = f'=TEXT(SUM({input_row_range})+{index},"0.00")'
            ws_second.cell(row=row, column=col).value = formula

    style_header(ws_second, 8, SECOND_INPUT_START_COL, SECOND_INPUT_END_COL, "FCE5CD")
    style_header(ws_second, 8, SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL, "D9D2E9")
    set_column_widths(ws_second, SECOND_INPUT_START_COL, SECOND_INPUT_END_COL, 12)
    set_column_widths(ws_second, SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL, 17)
    ws_second.freeze_panes = "B9"

    ws_second_other = wb_second.create_sheet("DiğerİşlemciSayfası")
    ws_second_other["A1"] = "Bu ek sayfa, sayfa seçiciyi test etmek için var."

    if small:
        wb_second.save(second_file)
    else:
        inflate_workbook_with_visible_rows(
            wb_second,
            second_file,
            MOCK_TARGET_SIZES_MB[second_file.name],
            "GörünürYük_İkinci",
            "ikinci",
        )
    wb_second.close()

    return first_file, second_file


class BatchWorker(QThread):
    log_message = pyqtSignal(str)
    finished_ok = pyqtSignal(str)
    failed = pyqtSignal(str)

    def __init__(
        self,
        first_file: Path,
        first_sheet: str,
        second_file: Path,
        second_sheet: str,
        batch_size: int,
        parent=None,
    ) -> None:
        super().__init__(parent)
        self.first_file = first_file
        self.first_sheet = first_sheet
        self.second_file = second_file
        self.second_sheet = second_sheet
        self.batch_size = batch_size

    def run(self) -> None:
        try:
            process_workbooks(
                self.first_file,
                self.first_sheet,
                self.second_file,
                self.second_sheet,
                self.log_message.emit,
                batch_size=self.batch_size,
            )
        except Exception:
            self.failed.emit(traceback.format_exc())
            return

        self.finished_ok.emit("Toplu işlem bitti.")


class ExcelBatchWindow(QWidget):
    def __init__(
        self,
        first_file: Path | None = None,
        second_file: Path | None = None,
        small_mocks: bool = False,
        batch_size: int = BATCH_SIZE,
    ) -> None:
        super().__init__()
        self.worker: BatchWorker | None = None
        self.small_mocks = small_mocks

        self.setWindowTitle("Excel'den Excel'e Toplu İşlem")
        self.resize(980, 640)
        self.setMinimumSize(820, 560)

        self.first_file_edit = QLineEdit()
        self.first_file_edit.setReadOnly(True)
        self.first_file_edit.setPlaceholderText("İlk Excel seçilmedi")
        self.second_file_edit = QLineEdit()
        self.second_file_edit.setReadOnly(True)
        self.second_file_edit.setPlaceholderText("İkinci Excel seçilmedi")

        self.first_sheet_combo = QComboBox()
        self.second_sheet_combo = QComboBox()

        self.first_browse_button = QPushButton("Dosya seç")
        self.second_browse_button = QPushButton("Dosya seç")
        self.first_open_button = QPushButton("Aç")
        self.second_open_button = QPushButton("Aç")
        self.first_open_button.setEnabled(False)
        self.second_open_button.setEnabled(False)
        self.sample_profile_combo = QComboBox()
        self.sample_profile_combo.addItem("Küçük örnek (hızlı test)", "small")
        self.sample_profile_combo.addItem("Büyük örnek (30 MB / 17 MB)", "large")
        self.sample_profile_combo.setCurrentIndex(0)
        self.batch_size_spin = QSpinBox()
        self.batch_size_spin.setRange(1, MAX_BATCH_SIZE)
        self.batch_size_spin.setValue(batch_size)
        self.batch_size_spin.setSuffix(" satır")
        self.batch_size_spin.setMinimumWidth(120)

        self.reset_mock_button = QPushButton("Örnek dosya oluştur")
        self.help_button = QPushButton()
        self.help_button.setFixedSize(36, 36)
        self.help_button.setIcon(standard_help_icon(self))
        self.help_button.setIconSize(QSize(20, 20))
        self.help_button.setObjectName("helpButton")
        self.start_button = QPushButton("İşlemi başlat")
        self.start_button.setObjectName("primaryButton")
        self.start_button.setEnabled(False)

        self.title_label = QLabel("Excel Toplu Aktarım")
        self.title_label.setObjectName("titleLabel")
        self.subtitle_label = QLabel("İki Excel dosyasını seç, sayfaları kontrol et, grupları sırayla işle.")
        self.subtitle_label.setObjectName("subtitleLabel")
        self.status_label = QLabel("Dosyalar bekleniyor.")
        self.status_label.setObjectName("statusLabel")
        self.status_label.setWordWrap(True)
        self.log_label = QLabel("İşlem günlüğü")
        self.log_label.setObjectName("sectionLabel")

        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        self.log_box.setMinimumHeight(180)
        self.apply_tooltips()
        self.apply_styles()

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(18, 16, 18, 16)
        main_layout.setSpacing(12)

        header_layout = QHBoxLayout()
        header_text_layout = QVBoxLayout()
        header_text_layout.setSpacing(2)
        header_text_layout.addWidget(self.title_label)
        header_text_layout.addWidget(self.subtitle_label)
        header_layout.addLayout(header_text_layout, 1)
        header_layout.addWidget(self.help_button)
        main_layout.addLayout(header_layout)

        first_group = self.build_workbook_group(
            "1. Kaynak Excel",
            self.first_file_edit,
            self.first_browse_button,
            self.first_open_button,
            self.first_sheet_combo,
        )
        second_group = self.build_workbook_group(
            "2. İşlemci Excel",
            self.second_file_edit,
            self.second_browse_button,
            self.second_open_button,
            self.second_sheet_combo,
        )
        main_layout.addWidget(first_group)
        main_layout.addWidget(second_group)

        status_layout = QHBoxLayout()
        status_layout.addWidget(self.status_label, 1)
        main_layout.addLayout(status_layout)

        action_layout = QHBoxLayout()
        action_layout.addStretch(1)
        action_layout.addWidget(QLabel("Grup satırı:"))
        action_layout.addWidget(self.batch_size_spin)
        action_layout.addWidget(QLabel("Deneme dosyası:"))
        action_layout.addWidget(self.sample_profile_combo)
        action_layout.addWidget(self.reset_mock_button)
        action_layout.addWidget(self.start_button)
        main_layout.addLayout(action_layout)

        main_layout.addWidget(self.log_label)
        main_layout.addWidget(self.log_box, 1)

        self.first_browse_button.clicked.connect(self.select_first_file)
        self.second_browse_button.clicked.connect(self.select_second_file)
        self.first_open_button.clicked.connect(lambda: self.open_selected_file(self.first_file_edit))
        self.second_open_button.clicked.connect(lambda: self.open_selected_file(self.second_file_edit))
        self.reset_mock_button.clicked.connect(self.create_sample_files)
        self.help_button.clicked.connect(self.show_help)
        self.start_button.clicked.connect(self.start_batch)

        if first_file:
            self.set_file(first_file, self.first_file_edit, self.first_sheet_combo, preferred_sheet=MOCK_FIRST_SHEET)
        if second_file:
            self.set_file(second_file, self.second_file_edit, self.second_sheet_combo, preferred_sheet=MOCK_SECOND_SHEET)

        self.update_start_button()
    def log(self, message: str) -> None:
        self.log_box.append(message)

    def build_workbook_group(
        self,
        title: str,
        file_edit: QLineEdit,
        browse_button: QPushButton,
        open_button: QPushButton,
        sheet_combo: QComboBox,
    ) -> QGroupBox:
        group = QGroupBox(title)
        group_layout = QGridLayout(group)
        group_layout.setContentsMargins(14, 18, 14, 14)
        group_layout.setHorizontalSpacing(10)
        group_layout.setVerticalSpacing(10)

        file_edit.setMinimumWidth(440)
        browse_button.setMinimumWidth(110)
        open_button.setMinimumWidth(70)
        file_edit.setSizePolicy(QSizePolicy.Policy.Expanding if QT_MAJOR == 6 else QSizePolicy.Expanding, QSizePolicy.Policy.Fixed if QT_MAJOR == 6 else QSizePolicy.Fixed)

        group_layout.addWidget(QLabel("Dosya:"), 0, 0)
        group_layout.addWidget(file_edit, 0, 1)
        group_layout.addWidget(browse_button, 0, 2)
        group_layout.addWidget(open_button, 0, 3)
        group_layout.addWidget(QLabel("Sayfa:"), 1, 0)
        group_layout.addWidget(sheet_combo, 1, 1, 1, 3)
        group_layout.setColumnStretch(1, 1)
        return group

    def apply_styles(self) -> None:
        self.setStyleSheet(
            """
            QWidget {
                font-size: 10pt;
                background-color: #f4f6f8;
                color: #1f2933;
            }
            QLabel#titleLabel {
                font-size: 18pt;
                font-weight: 700;
                color: #111827;
            }
            QLabel#subtitleLabel,
            QLabel#statusLabel {
                color: #52606d;
            }
            QLabel#sectionLabel {
                font-weight: 600;
                color: #1f2933;
            }
            QGroupBox {
                border: 1px solid #d0d7de;
                border-radius: 6px;
                margin-top: 10px;
                font-weight: 600;
                background-color: #ffffff;
                color: #1f2933;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 4px;
                background-color: #f4f6f8;
                color: #1f2933;
            }
            QLabel {
                background: transparent;
                color: #1f2933;
            }
            QLineEdit,
            QComboBox,
            QSpinBox,
            QTextEdit {
                border: 1px solid #c9d1d9;
                border-radius: 4px;
                padding: 6px;
                background: #ffffff;
                color: #111827;
                selection-background-color: #0969da;
                selection-color: #ffffff;
            }
            QLineEdit:disabled,
            QComboBox:disabled,
            QSpinBox:disabled,
            QTextEdit:disabled {
                background: #f6f8fa;
                color: #57606a;
            }
            QComboBox QAbstractItemView {
                background: #ffffff;
                color: #111827;
                selection-background-color: #0969da;
                selection-color: #ffffff;
            }
            QPushButton {
                border: 1px solid #c9d1d9;
                border-radius: 4px;
                padding: 7px 12px;
                background: #f6f8fa;
                color: #1f2933;
            }
            QPushButton:hover {
                background: #eef2f6;
            }
            QPushButton:disabled {
                background: #e5e7eb;
                border-color: #d1d5db;
                color: #6b7280;
            }
            QPushButton#primaryButton {
                background: #0969da;
                border-color: #0969da;
                color: white;
                font-weight: 600;
                min-width: 150px;
            }
            QPushButton#primaryButton:hover {
                background: #0757b8;
            }
            QPushButton#primaryButton:disabled {
                background: #9ca3af;
                border-color: #9ca3af;
                color: #ffffff;
            }
            QPushButton#helpButton {
                padding: 0;
                border-radius: 18px;
                background: #ffffff;
                border-color: #cfd8e3;
            }
            QTextEdit {
                font-family: Consolas, "Courier New", monospace;
                background: #ffffff;
            }
            """
        )

    def apply_tooltips(self) -> None:
        self.first_file_edit.setToolTip("Kaynak Excel dosyası. Input satırları F:AT sütunlarından okunur.")
        self.second_file_edit.setToolTip("İşlemci Excel dosyası. Input B9:AP aralığına seçilen satır sayısı kadar yapıştırılır.")
        self.first_sheet_combo.setToolTip("İlk Excel içinde F3:AT... input verisini içeren sayfa.")
        self.second_sheet_combo.setToolTip("Gerçek F9 hesaplamasından önce aktif yapılacak ikinci Excel sayfası.")
        self.first_browse_button.setToolTip("İlk/kaynak Excel dosyasını seç.")
        self.second_browse_button.setToolTip("İkinci/işlemci Excel dosyasını seç.")
        self.first_open_button.setToolTip("Seçili ilk Excel dosyasını aç.")
        self.second_open_button.setToolTip("Seçili ikinci Excel dosyasını aç.")
        self.batch_size_spin.setToolTip("Her grupta ilk Excel'den kaç satır okunacağını seç.")
        self.sample_profile_combo.setToolTip("Deneme için oluşturulacak örnek Excel boyutunu seç.")
        self.reset_mock_button.setToolTip("Seçilen boyutta örnek Excel dosyaları oluştur ve dosya alanlarını bunlara değiştir.")
        self.help_button.setToolTip("Tam grup akışını ve hücre aralıklarını göster.")
        self.start_button.setToolTip("İki dosya ve iki sayfa seçildikten sonra toplu işlemi başlat.")
        self.log_box.setToolTip("Her grup ve kaydetme adımı için ilerleme günlüğü.")

    def show_help(self) -> None:
        batch_size = self.batch_size_spin.value()
        first_example_end = FIRST_DATA_START_ROW + batch_size - 1
        second_example_start = FIRST_DATA_START_ROW + batch_size
        second_example_end = second_example_start + batch_size - 1
        second_range_end = SECOND_INPUT_START_ROW + batch_size - 1
        QMessageBox.information(
            self,
            "Bu araç nasıl çalışır?",
            (
                "Grup akışı:\n\n"
                f"1. İlk Excel'den F3:AT{first_example_end} ile başlayarak en fazla {batch_size} satır okur.\n"
                f"2. Bu satırları seçili ikinci Excel sayfasında B9:AP{second_range_end} aralığına yapıştırır.\n"
                "3. Seçili ikinci sayfayı aktif yapar ve gerçek Excel F9 hesaplamasını çalıştırır.\n"
                f"4. Aynı sayıda output satırını AQ9:BM{second_range_end} aralığından okur.\n"
                "5. Bu satırları ilk Excel'de aynı kaynak satırlardaki AU:BQ aralığına yazar.\n\n"
                "Örnek:\n"
                f"F3:AT{first_example_end} -> B9:AP{second_range_end} -> F9 -> AQ9:BM{second_range_end} -> AU3:BQ{first_example_end}\n"
                f"F{second_example_start}:AT{second_example_end} -> B9:AP{second_range_end} -> F9 -> AQ9:BM{second_range_end} -> AU{second_example_start}:BQ{second_example_end}\n\n"
                f"Son grup {batch_size} satırdan az input içerirse sadece o kadar output satırı geri kopyalanır.\n"
                "Output satır sayısı her zaman input satır sayısına eşittir."
            ),
        )

    def select_excel_file(self) -> str:
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Excel dosyası seç",
            "",
            "Excel Dosyaları (*.xlsx *.xlsm);;Tüm Dosyalar (*)",
        )
        return file_path

    def set_file(
        self,
        file_path: Path,
        file_edit: QLineEdit,
        sheet_combo: QComboBox,
        preferred_sheet: str | None = None,
    ) -> bool:
        file_path = Path(file_path)
        try:
            names = sheet_names(file_path)
        except Exception:
            QMessageBox.critical(self, "Sayfalar okunamadı", traceback.format_exc())
            return False

        file_edit.setText(str(file_path))
        sheet_combo.clear()
        sheet_combo.addItems(names)

        if preferred_sheet and preferred_sheet in names:
            sheet_combo.setCurrentText(preferred_sheet)

        self.log(f"{file_path.name} dosyasından {len(names)} sayfa yüklendi: {', '.join(names)}")
        self.update_start_button()
        return True

    def open_selected_file(self, file_edit: QLineEdit) -> None:
        file_path = file_edit.text()
        if not file_path:
            QMessageBox.information(self, "Dosya seçilmedi", "Önce bir Excel dosyası seç.")
            return

        path = Path(file_path)
        if not path.exists():
            QMessageBox.warning(self, "Dosya bulunamadı", f"Dosya bulunamadı:\n{path}")
            return

        os.startfile(str(path))

    def select_first_file(self) -> None:
        selected = self.select_excel_file()
        if selected:
            self.set_file(Path(selected), self.first_file_edit, self.first_sheet_combo)

    def select_second_file(self) -> None:
        selected = self.select_excel_file()
        if selected:
            self.set_file(Path(selected), self.second_file_edit, self.second_sheet_combo)

    def create_sample_files(self) -> None:
        if self.first_file_edit.text() or self.second_file_edit.text():
            confirmed = question_yes_no(
                self,
                "Seçili dosyalar değişecek",
                (
                    "Örnek Excel dosyaları oluşturursan şu anda seçili olan dosya yolları "
                    "örnek dosyalarla değiştirilecek.\n\n"
                    "Devam etmek istiyor musun?"
                ),
            )
            if not confirmed:
                self.log("Örnek dosya oluşturma iptal edildi.")
                return

        small = self.sample_profile_combo.currentData() == "small"
        profile = "küçük" if small else "büyük"
        batch_size = self.batch_size_spin.value()
        self.status_label.setText(f"{profile.capitalize()} örnek dosyalar oluşturuluyor...")
        self.log(f"{profile.capitalize()} örnek dosyalar oluşturuluyor. Grup satır sayısı: {batch_size}.")
        QApplication.processEvents()

        first_file, second_file = create_mock_workbooks(
            Path(__file__).resolve().parent,
            small=small,
            batch_size=batch_size,
        )
        self.set_file(first_file, self.first_file_edit, self.first_sheet_combo, preferred_sheet=MOCK_FIRST_SHEET)
        self.set_file(second_file, self.second_file_edit, self.second_sheet_combo, preferred_sheet=MOCK_SECOND_SHEET)
        self.log(f"Örnek dosyalar oluşturuldu ({profile}).")
        self.status_label.setText("Örnek dosyalar seçildi.")

    def update_start_button(self) -> None:
        controls_available = self.worker is None
        self.first_open_button.setEnabled(controls_available and bool(self.first_file_edit.text()))
        self.second_open_button.setEnabled(controls_available and bool(self.second_file_edit.text()))

        ready = (
            bool(self.first_file_edit.text())
            and bool(self.second_file_edit.text())
            and self.first_sheet_combo.count() > 0
            and self.second_sheet_combo.count() > 0
            and self.worker is None
        )
        self.start_button.setEnabled(ready)

        if self.worker is not None:
            self.status_label.setText("İşlem çalışıyor...")
            return

        missing = []
        if not self.first_file_edit.text():
            missing.append("ilk Excel")
        if self.first_sheet_combo.count() == 0:
            missing.append("ilk sayfa")
        if not self.second_file_edit.text():
            missing.append("ikinci Excel")
        if self.second_sheet_combo.count() == 0:
            missing.append("ikinci sayfa")

        if missing:
            self.status_label.setText("Hazır değil. Eksik seçimler: " + ", ".join(missing) + ".")
        else:
            self.status_label.setText("Hazır.")

    def set_controls_enabled(self, enabled: bool) -> None:
        self.first_browse_button.setEnabled(enabled)
        self.second_browse_button.setEnabled(enabled)
        self.first_open_button.setEnabled(enabled and bool(self.first_file_edit.text()))
        self.second_open_button.setEnabled(enabled and bool(self.second_file_edit.text()))
        self.batch_size_spin.setEnabled(enabled)
        self.sample_profile_combo.setEnabled(enabled)
        self.reset_mock_button.setEnabled(enabled)
        self.first_sheet_combo.setEnabled(enabled)
        self.second_sheet_combo.setEnabled(enabled)
        self.update_start_button()

    def start_batch(self) -> None:
        batch_size = self.batch_size_spin.value()
        self.worker = BatchWorker(
            Path(self.first_file_edit.text()),
            self.first_sheet_combo.currentText(),
            Path(self.second_file_edit.text()),
            self.second_sheet_combo.currentText(),
            batch_size,
            self,
        )
        self.worker.log_message.connect(self.log)
        self.worker.finished_ok.connect(self.on_finished)
        self.worker.failed.connect(self.on_failed)

        self.log(f"Toplu işlem başlatılıyor. Grup satır sayısı: {batch_size}.")
        self.status_label.setText("İşlem çalışıyor...")
        self.set_controls_enabled(False)
        self.start_button.setEnabled(False)
        self.worker.start()

    def on_finished(self, message: str) -> None:
        self.log(message)
        QMessageBox.information(self, "Bitti", message)
        self.worker = None
        self.set_controls_enabled(True)
        self.status_label.setText("İşlem tamamlandı.")

    def on_failed(self, error_text: str) -> None:
        self.log(error_text)
        QMessageBox.critical(self, "Toplu işlem başarısız oldu", error_text)
        self.worker = None
        self.set_controls_enabled(True)
        self.status_label.setText("Hata oluştu.")


def parse_args(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Excel'den Excel'e toplu işlem arayüzü.")
    parser.add_argument("--make-mocks", action="store_true", help="Örnek Excel dosyalarını oluştur/sıfırla.")
    parser.add_argument("--small-mocks", action="store_true", help="Hızlı test için küçük örnek dosyalar oluştur.")
    parser.add_argument("--no-gui", action="store_true", help="Arayüzü başlatma.")
    parser.add_argument("--run-once", action="store_true", help="Arayüzsüz tek toplu işlem çalıştır.")
    parser.add_argument("--first", type=Path, help="Önceden seçilecek ilk Excel dosyası.")
    parser.add_argument("--second", type=Path, help="Önceden seçilecek ikinci Excel dosyası.")
    parser.add_argument("--first-sheet", default=MOCK_FIRST_SHEET, help="--run-once için ilk çalışma kitabı sayfası.")
    parser.add_argument("--second-sheet", default=MOCK_SECOND_SHEET, help="--run-once için ikinci çalışma kitabı sayfası.")
    parser.add_argument("--max-batches", type=int, help="Hızlı test için arayüzsüz çalışmayı bu grup sayısıyla sınırla.")
    parser.add_argument("--batch-size", type=int, default=BATCH_SIZE, help="Her grupta işlenecek satır sayısı.")
    parser.add_argument("--delay", type=float, default=None, help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.batch_size < 1:
        parser.error("--batch-size 1 veya daha büyük olmalı.")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    args = parse_args(argv or sys.argv[1:])
    script_dir = Path(__file__).resolve().parent

    first_file = args.first
    second_file = args.second

    if args.make_mocks:
        first_file, second_file = create_mock_workbooks(
            script_dir,
            small=args.small_mocks,
            batch_size=args.batch_size,
        )
        print(f"Oluşturuldu: {first_file}")
        print(f"Oluşturuldu: {second_file}")

    if args.run_once:
        if not first_file or not second_file:
            first_file = script_dir / "mock_first_input.xlsx"
            second_file = script_dir / "mock_second_processor.xlsx"
        process_workbooks(
            first_file,
            args.first_sheet,
            second_file,
            args.second_sheet,
            print,
            max_batches=args.max_batches,
            batch_size=args.batch_size,
        )

    if args.no_gui:
        return 0

    app = QApplication(sys.argv[:1])
    app.setStyle("Fusion")
    window = ExcelBatchWindow(
        first_file=first_file,
        second_file=second_file,
        small_mocks=args.small_mocks,
        batch_size=args.batch_size,
    )
    window.show()
    if QT_MAJOR == 6:
        return app.exec()
    return app.exec_()


if __name__ == "__main__":
    raise SystemExit(main())
