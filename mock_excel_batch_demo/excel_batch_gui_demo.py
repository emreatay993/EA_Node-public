from __future__ import annotations

import argparse
import hashlib
import math
import sys
import time
import traceback
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from PyQt6.QtCore import QThread, pyqtSignal
    from PyQt6.QtWidgets import (
        QApplication,
        QComboBox,
        QFileDialog,
        QGridLayout,
        QLabel,
        QLineEdit,
        QMessageBox,
        QPushButton,
        QTextEdit,
        QWidget,
    )

    QT_MAJOR = 6
except ImportError:  # pragma: no cover - fallback for machines with PyQt5 only.
    from PyQt5.QtCore import QThread, pyqtSignal
    from PyQt5.QtWidgets import (
        QApplication,
        QComboBox,
        QFileDialog,
        QGridLayout,
        QLabel,
        QLineEdit,
        QMessageBox,
        QPushButton,
        QTextEdit,
        QWidget,
    )

    QT_MAJOR = 5


BATCH_SIZE = 6
DEFAULT_F9_SIMULATION_SECONDS = 7.0
MOCK_FIRST_VISIBLE_INPUT_ROWS = 4000
MOCK_VISIBLE_LOAD_COLS = 20
MOCK_VISIBLE_LOAD_ESTIMATED_BYTES_PER_ROW = 930

MOCK_TARGET_SIZES_MB = {
    "mock_first_input.xlsx": 30,
    "mock_second_processor.xlsx": 17,
}

FIRST_DATA_START_ROW = 3
FIRST_INPUT_START_COL = 6  # F
FIRST_INPUT_END_COL = 46  # AT
FIRST_OUTPUT_START_COL = 47  # AU
FIRST_OUTPUT_END_COL = 69  # BQ

SECOND_INPUT_START_ROW = 9
SECOND_INPUT_START_COL = 2  # B
SECOND_INPUT_END_COL = 42  # AP
SECOND_OUTPUT_ROW = 9
SECOND_OUTPUT_START_COL = 43  # AQ
SECOND_OUTPUT_END_COL = 65  # BM


def col_range(start_col: int, end_col: int) -> range:
    return range(start_col, end_col + 1)


def iter_cells_2d(
    ws,
    row1: int,
    col1: int,
    row2: int,
    col2: int,
) -> Iterable[tuple[int, int]]:
    for row in range(row1, row2 + 1):
        for col in range(col1, col2 + 1):
            yield row, col


def clear_range(ws, row1: int, col1: int, row2: int, col2: int) -> None:
    for row, col in iter_cells_2d(ws, row1, col1, row2, col2):
        ws.cell(row=row, column=col).value = None


def read_values(ws, row1: int, col1: int, row2: int, col2: int) -> list[list[object]]:
    return [
        [ws.cell(row=row, column=col).value for col in col_range(col1, col2)]
        for row in range(row1, row2 + 1)
    ]


def write_values(ws, start_row: int, start_col: int, values: Sequence[Sequence[object]]) -> None:
    for row_offset, row_values in enumerate(values):
        for col_offset, value in enumerate(row_values):
            ws.cell(row=start_row + row_offset, column=start_col + col_offset).value = value


def sheet_names(file_path: Path) -> list[str]:
    wb = load_workbook(file_path, read_only=True, keep_vba=file_path.suffix.lower() == ".xlsm")
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def visible_load_value(workbook_key: str, row: int, col: int) -> str:
    digest = hashlib.sha256(f"{workbook_key}:{row}:{col}:visible-load".encode("ascii")).hexdigest()
    return f"{workbook_key}_visible_load_{row:06d}_{col:02d}_{digest}"


def ensure_visible_load_sheet(wb: Workbook, sheet_name: str):
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        ws.append([f"load_col_{index:02d}" for index in range(1, MOCK_VISIBLE_LOAD_COLS + 1)])
        style_header(ws, 1, 1, MOCK_VISIBLE_LOAD_COLS, "FFF2CC")
        set_column_widths(ws, 1, MOCK_VISIBLE_LOAD_COLS, 24)
        ws.freeze_panes = "A2"

    return ws


def append_visible_load_rows(ws, workbook_key: str, row_count: int) -> None:
    start_row = ws.max_row + 1
    for row in range(start_row, start_row + row_count):
        ws.append(
            [
                visible_load_value(workbook_key, row, col)
                for col in range(1, MOCK_VISIBLE_LOAD_COLS + 1)
            ]
        )


def inflate_workbook_with_visible_rows(
    wb: Workbook,
    file_path: Path,
    target_mb: int,
    sheet_name: str,
    workbook_key: str,
) -> None:
    target_bytes = target_mb * 1024 * 1024
    ws = ensure_visible_load_sheet(wb, sheet_name)

    wb.save(file_path)
    current_bytes = file_path.stat().st_size

    attempts = 0
    while current_bytes < target_bytes and attempts < 6:
        remaining_bytes = target_bytes - current_bytes
        rows_to_add = max(
            50,
            math.ceil(remaining_bytes / MOCK_VISIBLE_LOAD_ESTIMATED_BYTES_PER_ROW),
        )
        append_visible_load_rows(ws, workbook_key, rows_to_add)
        wb.save(file_path)
        current_bytes = file_path.stat().st_size
        attempts += 1


def find_last_used_row(ws, start_col: int, end_col: int, min_row: int) -> int:
    last_row = min_row - 1
    for row in range(min_row, ws.max_row + 1):
        for col in col_range(start_col, end_col):
            if ws.cell(row=row, column=col).value not in (None, ""):
                last_row = row
                break
    return last_row


def fake_output_for_batch(
    batch_index: int,
    source_start_row: int,
    source_end_row: int,
    source_values: Sequence[Sequence[object]],
) -> list[object]:
    numeric_values = [
        value
        for row in source_values
        for value in row
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    ]
    numeric_sum = round(sum(numeric_values), 2)
    row_count = source_end_row - source_start_row + 1

    output_count = SECOND_OUTPUT_END_COL - SECOND_OUTPUT_START_COL + 1
    output = []
    for index in range(output_count):
        if index % 4 == 0:
            output.append(f"batch_{batch_index:02d}_out_{index + 1:02d}")
        elif index % 4 == 1:
            output.append(source_start_row)
        elif index % 4 == 2:
            output.append(source_end_row)
        else:
            output.append(round(numeric_sum + row_count + index, 2))
    return output


def process_workbooks(
    first_file: Path,
    first_sheet: str,
    second_file: Path,
    second_sheet: str,
    log: Callable[[str], None],
    delay_seconds: float = DEFAULT_F9_SIMULATION_SECONDS,
) -> None:
    first_file = first_file.resolve()
    second_file = second_file.resolve()

    wb_first = load_workbook(first_file)
    wb_second = load_workbook(second_file)

    try:
        ws_first = wb_first[first_sheet]
        ws_second = wb_second[second_sheet]

        last_row = find_last_used_row(
            ws_first,
            FIRST_INPUT_START_COL,
            FIRST_INPUT_END_COL,
            FIRST_DATA_START_ROW,
        )

        if last_row < FIRST_DATA_START_ROW:
            raise ValueError("No input data found in F3:AT.")

        batch_index = 1
        for source_start_row in range(FIRST_DATA_START_ROW, last_row + 1, BATCH_SIZE):
            source_end_row = min(source_start_row + BATCH_SIZE - 1, last_row)
            row_count = source_end_row - source_start_row + 1

            log(f"Batch {batch_index}: reading first Excel rows {source_start_row}-{source_end_row}")
            source_values = read_values(
                ws_first,
                source_start_row,
                FIRST_INPUT_START_COL,
                source_end_row,
                FIRST_INPUT_END_COL,
            )

            log("  writing input to second Excel B9:AP14")
            clear_range(
                ws_second,
                SECOND_INPUT_START_ROW,
                SECOND_INPUT_START_COL,
                SECOND_INPUT_START_ROW + BATCH_SIZE - 1,
                SECOND_INPUT_END_COL,
            )
            clear_range(
                ws_second,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_START_COL,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_END_COL,
            )
            write_values(
                ws_second,
                SECOND_INPUT_START_ROW,
                SECOND_INPUT_START_COL,
                source_values,
            )

            log(f"  simulating F9 / long Excel script ({delay_seconds:g}s)")
            time.sleep(delay_seconds)

            second_output_values = fake_output_for_batch(
                batch_index,
                source_start_row,
                source_end_row,
                source_values,
            )
            write_values(
                ws_second,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_START_COL,
                [second_output_values],
            )

            log("  copying second Excel AQ9:BM9 back to first Excel AU:BQ")
            first_output_values = read_values(
                ws_second,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_START_COL,
                SECOND_OUTPUT_ROW,
                SECOND_OUTPUT_END_COL,
            )
            write_values(
                ws_first,
                source_start_row,
                FIRST_OUTPUT_START_COL,
                first_output_values,
            )

            batch_index += 1

        wb_first.save(first_file)
        wb_second.save(second_file)
        log(f"Saved first Excel: {first_file}")
        log(f"Saved second Excel: {second_file}")
    finally:
        wb_first.close()
        wb_second.close()


def style_header(ws, row: int, start_col: int, end_col: int, fill_color: str) -> None:
    fill = PatternFill("solid", fgColor=fill_color)
    for col in col_range(start_col, end_col):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(bold=True)


def set_column_widths(ws, start_col: int, end_col: int, width: float) -> None:
    for col in col_range(start_col, end_col):
        ws.column_dimensions[get_column_letter(col)].width = width


def create_mock_workbooks(output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    first_file = output_dir / "mock_first_input.xlsx"
    second_file = output_dir / "mock_second_processor.xlsx"

    wb_first = Workbook()
    ws_first = wb_first.active
    ws_first.title = "InputData"
    ws_first["A1"] = "Mock first Excel. Inputs are in F3:AT. Outputs will be written to AU:BQ."

    for index, col in enumerate(col_range(FIRST_INPUT_START_COL, FIRST_INPUT_END_COL), start=1):
        ws_first.cell(row=2, column=col).value = f"input_{index:02d}"
    for index, col in enumerate(col_range(FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL), start=1):
        ws_first.cell(row=2, column=col).value = f"output_{index:02d}"

    for row in range(FIRST_DATA_START_ROW, FIRST_DATA_START_ROW + MOCK_FIRST_VISIBLE_INPUT_ROWS):
        for index, col in enumerate(col_range(FIRST_INPUT_START_COL, FIRST_INPUT_END_COL), start=1):
            if (row + index) % 3 == 0:
                value = round(row * 10 + index * 0.25, 2)
            elif (row + index) % 3 == 1:
                value = f"R{row}_C{get_column_letter(col)}"
            else:
                value = row + index
            ws_first.cell(row=row, column=col).value = value

    style_header(ws_first, 2, FIRST_INPUT_START_COL, FIRST_INPUT_END_COL, "D9EAD3")
    style_header(ws_first, 2, FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL, "CFE2F3")
    set_column_widths(ws_first, FIRST_INPUT_START_COL, FIRST_INPUT_END_COL, 12)
    set_column_widths(ws_first, FIRST_OUTPUT_START_COL, FIRST_OUTPUT_END_COL, 17)
    ws_first.freeze_panes = "F3"

    ws_first_other = wb_first.create_sheet("ExtraSheet")
    ws_first_other["A1"] = "This extra sheet exists to test the sheet selector."

    inflate_workbook_with_visible_rows(
        wb_first,
        first_file,
        MOCK_TARGET_SIZES_MB[first_file.name],
        "VisibleLoad_First",
        "first",
    )
    wb_first.close()

    wb_second = Workbook()
    ws_second = wb_second.active
    ws_second.title = "Processor"
    ws_second["A1"] = "Mock second Excel. Inputs are pasted to B9:AP14."
    ws_second["A2"] = "The demo simulates F9, writes fake output to AQ9:BM9, then copies it back."

    for index, col in enumerate(col_range(SECOND_INPUT_START_COL, SECOND_INPUT_END_COL), start=1):
        ws_second.cell(row=8, column=col).value = f"input_{index:02d}"
    for index, col in enumerate(col_range(SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL), start=1):
        ws_second.cell(row=8, column=col).value = f"output_{index:02d}"

    style_header(ws_second, 8, SECOND_INPUT_START_COL, SECOND_INPUT_END_COL, "FCE5CD")
    style_header(ws_second, 8, SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL, "D9D2E9")
    set_column_widths(ws_second, SECOND_INPUT_START_COL, SECOND_INPUT_END_COL, 12)
    set_column_widths(ws_second, SECOND_OUTPUT_START_COL, SECOND_OUTPUT_END_COL, 17)
    ws_second.freeze_panes = "B9"

    ws_second_other = wb_second.create_sheet("OtherProcessorSheet")
    ws_second_other["A1"] = "This extra sheet exists to test the sheet selector."

    inflate_workbook_with_visible_rows(
        wb_second,
        second_file,
        MOCK_TARGET_SIZES_MB[second_file.name],
        "VisibleLoad_Second",
        "second",
    )
    wb_second.close()

    return first_file, second_file


class BatchWorker(QThread):
    log_message = pyqtSignal(str)
    finished_ok = pyqtSignal(str)
    failed = pyqtSignal(str)

    def __init__(
        self,
        first_file: Path,
        first_sheet: str,
        second_file: Path,
        second_sheet: str,
        parent=None,
    ) -> None:
        super().__init__(parent)
        self.first_file = first_file
        self.first_sheet = first_sheet
        self.second_file = second_file
        self.second_sheet = second_sheet

    def run(self) -> None:
        try:
            process_workbooks(
                self.first_file,
                self.first_sheet,
                self.second_file,
                self.second_sheet,
                self.log_message.emit,
            )
        except Exception:
            self.failed.emit(traceback.format_exc())
            return

        self.finished_ok.emit("Batch process finished.")


class ExcelBatchWindow(QWidget):
    def __init__(self, first_file: Path | None = None, second_file: Path | None = None) -> None:
        super().__init__()
        self.worker: BatchWorker | None = None

        self.setWindowTitle(f"Excel Batch Demo - PyQt{QT_MAJOR}")
        self.resize(880, 520)

        self.first_file_edit = QLineEdit()
        self.first_file_edit.setReadOnly(True)
        self.second_file_edit = QLineEdit()
        self.second_file_edit.setReadOnly(True)

        self.first_sheet_combo = QComboBox()
        self.second_sheet_combo = QComboBox()

        self.first_browse_button = QPushButton("Select")
        self.second_browse_button = QPushButton("Select")
        self.reset_mock_button = QPushButton("Reset mock files")
        self.start_button = QPushButton("Start batch")
        self.start_button.setEnabled(False)

        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)

        layout = QGridLayout(self)
        layout.addWidget(QLabel("First Excel:"), 0, 0)
        layout.addWidget(self.first_file_edit, 0, 1)
        layout.addWidget(self.first_browse_button, 0, 2)

        layout.addWidget(QLabel("First sheet:"), 1, 0)
        layout.addWidget(self.first_sheet_combo, 1, 1, 1, 2)

        layout.addWidget(QLabel("Second Excel:"), 2, 0)
        layout.addWidget(self.second_file_edit, 2, 1)
        layout.addWidget(self.second_browse_button, 2, 2)

        layout.addWidget(QLabel("Second sheet:"), 3, 0)
        layout.addWidget(self.second_sheet_combo, 3, 1, 1, 2)

        layout.addWidget(self.reset_mock_button, 4, 0)
        layout.addWidget(self.start_button, 4, 2)
        layout.addWidget(self.log_box, 5, 0, 1, 3)

        self.first_browse_button.clicked.connect(self.select_first_file)
        self.second_browse_button.clicked.connect(self.select_second_file)
        self.reset_mock_button.clicked.connect(self.reset_mocks)
        self.start_button.clicked.connect(self.start_batch)

        if first_file:
            self.set_file(first_file, self.first_file_edit, self.first_sheet_combo, preferred_sheet="InputData")
        if second_file:
            self.set_file(second_file, self.second_file_edit, self.second_sheet_combo, preferred_sheet="Processor")

        self.update_start_button()
        self.log(f"Running with PyQt{QT_MAJOR}.")

    def log(self, message: str) -> None:
        self.log_box.append(message)

    def select_excel_file(self) -> str:
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Excel file",
            "",
            "Excel Files (*.xlsx *.xlsm);;All Files (*)",
        )
        return file_path

    def set_file(
        self,
        file_path: Path,
        file_edit: QLineEdit,
        sheet_combo: QComboBox,
        preferred_sheet: str | None = None,
    ) -> bool:
        file_path = Path(file_path)
        try:
            names = sheet_names(file_path)
        except Exception:
            QMessageBox.critical(self, "Could not read sheets", traceback.format_exc())
            return False

        file_edit.setText(str(file_path))
        sheet_combo.clear()
        sheet_combo.addItems(names)

        if preferred_sheet and preferred_sheet in names:
            sheet_combo.setCurrentText(preferred_sheet)

        self.log(f"Loaded {len(names)} sheet(s) from {file_path.name}: {', '.join(names)}")
        self.update_start_button()
        return True

    def select_first_file(self) -> None:
        selected = self.select_excel_file()
        if selected:
            self.set_file(Path(selected), self.first_file_edit, self.first_sheet_combo)

    def select_second_file(self) -> None:
        selected = self.select_excel_file()
        if selected:
            self.set_file(Path(selected), self.second_file_edit, self.second_sheet_combo)

    def reset_mocks(self) -> None:
        first_file, second_file = create_mock_workbooks(Path(__file__).resolve().parent)
        self.set_file(first_file, self.first_file_edit, self.first_sheet_combo, preferred_sheet="InputData")
        self.set_file(second_file, self.second_file_edit, self.second_sheet_combo, preferred_sheet="Processor")
        self.log("Mock files reset.")

    def update_start_button(self) -> None:
        ready = (
            bool(self.first_file_edit.text())
            and bool(self.second_file_edit.text())
            and self.first_sheet_combo.count() > 0
            and self.second_sheet_combo.count() > 0
            and self.worker is None
        )
        self.start_button.setEnabled(ready)

    def set_controls_enabled(self, enabled: bool) -> None:
        self.first_browse_button.setEnabled(enabled)
        self.second_browse_button.setEnabled(enabled)
        self.reset_mock_button.setEnabled(enabled)
        self.first_sheet_combo.setEnabled(enabled)
        self.second_sheet_combo.setEnabled(enabled)
        self.update_start_button()

    def start_batch(self) -> None:
        self.worker = BatchWorker(
            Path(self.first_file_edit.text()),
            self.first_sheet_combo.currentText(),
            Path(self.second_file_edit.text()),
            self.second_sheet_combo.currentText(),
            self,
        )
        self.worker.log_message.connect(self.log)
        self.worker.finished_ok.connect(self.on_finished)
        self.worker.failed.connect(self.on_failed)

        self.log("Starting batch process.")
        self.set_controls_enabled(False)
        self.start_button.setEnabled(False)
        self.worker.start()

    def on_finished(self, message: str) -> None:
        self.log(message)
        QMessageBox.information(self, "Done", message)
        self.worker = None
        self.set_controls_enabled(True)

    def on_failed(self, error_text: str) -> None:
        self.log(error_text)
        QMessageBox.critical(self, "Batch failed", error_text)
        self.worker = None
        self.set_controls_enabled(True)


def parse_args(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Mock Excel batch processor GUI.")
    parser.add_argument("--make-mocks", action="store_true", help="Create/reset mock Excel files.")
    parser.add_argument("--no-gui", action="store_true", help="Do not launch the GUI.")
    parser.add_argument("--run-once", action="store_true", help="Run one non-GUI batch process.")
    parser.add_argument("--first", type=Path, help="First Excel file to preselect.")
    parser.add_argument("--second", type=Path, help="Second Excel file to preselect.")
    parser.add_argument("--first-sheet", default="InputData", help="First workbook sheet name for --run-once.")
    parser.add_argument("--second-sheet", default="Processor", help="Second workbook sheet name for --run-once.")
    parser.add_argument(
        "--delay",
        type=float,
        default=DEFAULT_F9_SIMULATION_SECONDS,
        help="Seconds to wait while simulating F9.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    script_dir = Path(__file__).resolve().parent

    first_file = args.first
    second_file = args.second

    if args.make_mocks:
        first_file, second_file = create_mock_workbooks(script_dir)
        print(f"Created {first_file}")
        print(f"Created {second_file}")

    if args.run_once:
        if not first_file or not second_file:
            first_file = script_dir / "mock_first_input.xlsx"
            second_file = script_dir / "mock_second_processor.xlsx"
        process_workbooks(
            first_file,
            args.first_sheet,
            second_file,
            args.second_sheet,
            print,
            delay_seconds=args.delay,
        )

    if args.no_gui:
        return 0

    app = QApplication(sys.argv[:1])
    window = ExcelBatchWindow(first_file=first_file, second_file=second_file)
    window.show()
    if QT_MAJOR == 6:
        return app.exec()
    return app.exec_()


if __name__ == "__main__":
    raise SystemExit(main())
